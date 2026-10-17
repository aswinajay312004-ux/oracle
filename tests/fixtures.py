"""Synthetic documents that reproduce the STRUCTURE of the client's files with made-up values.
Client workbooks are never committed; their layouts are covered here."""
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook

SIZES = ['0-3M', '3-6M', '6-12M', '12-18M']
ORDER = [[10, 20, 30, 40], [5, 10, 15, 20]]


def merch_bom(sample=(0, 0, 0, 0), extra_sewing_block=False):
    """Merch BOM: header labels, order/sample/summary matrices, fabric table, accessory blocks."""
    wb = Workbook(); ws = wb.active; ws.title = 'D99001'
    ws['B1'] = 'TESTBUYER - BOM'
    for r, (a, b, c, d) in enumerate([('BUYER', 'TESTBUYER', 'BOM DATE', datetime(2026, 9, 1)),
                                      ('SEASON', 'SS-27', 'FOB DATE', datetime(2027, 1, 15)),
                                      ('ITEM', 'T1001', 'FAB IH DATE', datetime(2026, 11, 1)),
                                      ('DESCRIPTION', 'TEST BODYSUIT', 'CONTRACT', 'TC100-1'),
                                      ('MERCHANDISER', 'TESTMERCH', 'DESTINATION', 'SEA')], start=2):
        ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'], ws[f'E{r}'] = a, b, c, d
    ws['B8'], ws['C8'], ws['D8'] = 'OCN', 'QTY(IN PACKS)', 'QTY(IN PIECES)'
    ws['C9'], ws['D9'] = 100, '=C9*2'
    def matrix(top, role, rows):
        ws[f'B{top}'], ws[f'C{top}'], ws[f'D{top}'] = 'Garment ', 'Description', 'CONTRACT'
        for i, s in enumerate(SIZES):
            ws.cell(top, 5 + i, s)
        ws.cell(top, 9, 'Total')
        for li, (label, color, qty) in enumerate(rows, start=1):
            r = top + li
            ws[f'B{r}'], ws[f'C{r}'] = label, color
            if li == 1:
                ws[f'D{r}'] = role
            for i, q in enumerate(qty):
                ws.cell(r, 5 + i, q)
            ws.cell(r, 9, f'=SUM(E{r}:H{r})')
        ws[f'B{top + len(rows) + 1}'] = 'Total'
    lines = [('GARMENT-1 BODY SUIT', 'WHITE AOP'), ('GARMENT-2 BODY SUIT', 'PINK SOLID')]
    matrix(11, 'TC100-1 SEA', [(l, c, q) for (l, c), q in zip(lines, ORDER)])
    matrix(15, 'SAMPLE', [(l, c, sample) for l, c in lines])
    matrix(19, 'SUMMARY', [(l, c, q) for (l, c), q in zip(lines, ORDER)])
    ws['B23'] = 'FABRIC REQUIREMENT'
    for i, h in enumerate(['Garment ', 'Colour Description', 'Base Colour', 'Fabric Quality', 'Gsm ', 'Content', 'Count', 'Dia', 'Order Qty', 'YY', 'Requirment', 'Rej 3%', 'Total Qty ']):
        ws.cell(24, 2 + i, h)
    for r, (g, color, q, gsm, cnt, qty, yy) in enumerate([('Garment - 1  BODY', 'WHITE AOP', 'Interlock', 200, "42'S", 100, 60),
                                                          ('Garment - 2 - BODY', 'PINK SOLID', 'Interlock', 200, "42'S", 50, 60),
                                                          ('Garment - 2 - BINDING', 'PINK SOLID', '1x1 Rib', 180, "32's", 50, 25)], start=25):
        for i, v in enumerate([g, color, color, q, gsm, '100% Cotton-Bio wash', cnt, '66 Dia', qty, yy]):
            ws.cell(r, 2 + i, v)
        ws.cell(r, 12, f'=J{r}*K{r}/1000'); ws.cell(r, 13, 1.03); ws.cell(r, 14, f'=L{r}*M{r}')
    ws['B30'] = 'CARE LABEL'
    for i, h in enumerate(['Garment ', 'CODE ', 'Size', 'Supplier', 'Order Qty', 'YY', 'Required Qty', 'Rejection', 'Req In Single']):
        ws.cell(31, 2 + i, h)
    for k, (s, q) in enumerate(zip(['0-3 Mths', '3-6 Mths', '6-12 Mths', '12-18 Mths'], [15, 30, 45, 60])):
        r = 32 + k
        if k == 0:
            ws[f'B{r}'], ws[f'C{r}'] = 'ALL GARMENTS', 'CARE-TEST'
        ws[f'D{r}'], ws[f'E{r}'], ws[f'F{r}'], ws[f'G{r}'] = s, 'ITL', q, 1
        ws[f'H{r}'] = q; ws[f'I{r}'] = 1.03; ws[f'J{r}'] = round(q * 1.03, 2)
    ws['D36'] = 'TOTAL'; ws['F36'] = 150
    ws['B38'] = 'TWILL TAPE 5MM'
    for i, h in enumerate(['Garment ', 'CODE ', 'Size', 'Supplier', 'Order Qty', 'YY', 'Required Qty', 'Rejection', 'Req In Mtr']):
        ws.cell(39, 2 + i, h)
    ws['B40'], ws['C40'], ws['D40'], ws['E40'], ws['F40'], ws['G40'], ws['H40'], ws['I40'], ws['J40'] = 'GARMENT-1', 'WHITE', 'ALL SIZE', 'SBL', 100, 0.06, 6, 1.03, 6.18
    ws['B41'], ws['C41'], ws['F41'], ws['H41'], ws['J41'] = 'GARMENT-2', 'PINK', 50, 3, 3.09
    ws['F42'], ws['J42'] = 150, 9.27  # unlabelled subtotal row
    ws['B44'] = 'A VERY LONG ACCESSORY NAME FOR A POLY BAG ITEM'
    for i, h in enumerate(['Garment ', 'CODE ', 'Size', 'Supplier', 'Order Qty', 'YY', 'Required Qty', 'Rejection', 'Req In Single']):
        ws.cell(45, 2 + i, h)
    ws['B46'], ws['C46'], ws['D46'], ws['F46'], ws['G46'], ws['H46'], ws['I46'], ws['J46'] = 'A VERY LONG ACCESSORY NAME FOR A POLY BAG ITEM', '40X30 CM', 'ALL SIZES', 150, 1, 150, 1.01, 151.5
    ws['D47'] = 'TOTAL'
    if extra_sewing_block:
        for top in (50, 55):
            ws[f'B{top - 1}'] = 'SEWING THREAD'
            for i, h in enumerate(['Garment', 'THREAD COLOR', 'SHADE', 'QUALITY', 'Order Qty', 'YY', 'Required Qty', 'Rejection', 'Req In Single']):
                ws.cell(top, 2 + i, h)
            ws[f'B{top + 1}'], ws[f'C{top + 1}'], ws[f'D{top + 1}'], ws[f'F{top + 1}'], ws[f'J{top + 1}'] = 'GARMENT-1', 'WHITE', 'COATS', 100, 2.5
    return save(wb)


def save(wb):
    buf = BytesIO(); wb.save(buf); return calculated(buf.getvalue())


def calculated(content):
    """openpyxl cannot calculate formulas; write the cached values Excel would have stored."""
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(content))
    import re
    def value(ws, ref, seen=()):
        v = ws[ref].value
        if not (isinstance(v, str) and v.startswith('=')):
            return v or 0
        expr = v[1:]
        m = re.fullmatch(r'SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)', expr)
        if m:
            from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
            return sum(value(ws, f'{gl(c)}{m.group(2)}') for c in range(ci(m.group(1)), ci(m.group(3)) + 1))
        return eval(re.sub(r'([A-Z]+\d+)', lambda x: str(value(ws, x.group(1))), expr))  # test-only arithmetic
    # Store formulas and cached values the way Excel does: keep formula XML, inject <v> values.
    cache = {(ws.title, c.coordinate): value(ws, c.coordinate) for ws in wb for row in ws for c in row
             if isinstance(c.value, str) and c.value.startswith('=')}
    from zipfile import ZipFile, ZIP_DEFLATED
    src = ZipFile(BytesIO(content)); out = BytesIO(); dst = ZipFile(out, 'w', ZIP_DEFLATED)
    for item in src.infolist():
        data = src.read(item.filename)
        if item.filename.startswith('xl/worksheets/sheet'):
            text = data.decode()
            title = wb.worksheets[int(re.search(r'sheet(\d+)', item.filename).group(1)) - 1].title
            for (sheet_title, ref), v in cache.items():
                if sheet_title == title:
                    text = re.sub(rf'(<c r="{ref}"[^>]*>)<f>(.*?)</f><v\s*/>', lambda m, v=v: f'{m.group(1)}<f>{m.group(2)}</f><v>{v}</v>', text)
            data = text.encode()
        dst.writestr(item, data)
    dst.close()
    return out.getvalue()


def flat_orders():
    wb = Workbook(); ws = wb.active; ws.title = 'Orders'
    ws.append(['Buyer', 'Style No', 'Colour', 'Size', 'Qty'])
    for row in [('BUYER X', '0008', 'BLUE', 'S', 4), ('BUYER X', '0008', 'BLUE', 'M', 6), ('BUYER X', '0008', 'RED', 'S', 3)]:
        ws.append(row)
    return save(wb)


def unknown_matrix():
    """A layout no built-in detector knows: labels in column A, sizes across from column C."""
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    ws['A1'] = 'Customer'; ws['B1'] = 'ZEN KIDS'
    ws['A2'] = 'Article'; ws['B2'] = 'ZK-55'
    ws['A3'] = 'Order no'; ws['B3'] = 'ZK/PO/9'
    ws['A5'] = 'Shade'; ws['B5'] = 'Remark'; ws['C5'] = 'S'; ws['D5'] = 'M'; ws['E5'] = 'L'; ws['F5'] = 'Total'
    ws['A6'] = 'NAVY'; ws['C6'] = 10; ws['D6'] = 20; ws['E6'] = 30; ws['F6'] = 60
    ws['A7'] = 'GREY'; ws['C7'] = 5; ws['D7'] = 5; ws['E7'] = 5; ws['F7'] = 15
    return save(wb)


def purchase_order_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas
    buf = BytesIO(); c = canvas.Canvas(buf, pagesize=landscape(A4)); y = 560
    for line in ['PURCHASE ORDER', 'Buyer: BABY SHOP', 'Style No: 0008', 'Description: BABY ROMPER', 'PO No: BS-7781', 'Ship Date: 2026-12-20']:
        c.drawString(40, y, line); y -= 18
    y -= 10
    for r in [['Style', 'Colour', 'Size', 'Qty'], ['0008', 'ABSTRACT BLUE', '0-3M', '40'], ['0008', 'ABSTRACT BLUE', '3-6M', '60'], ['0008', 'SOFT PINK', '0-3M', '25'], ['Total', '', '', '125']]:
        for i, v in enumerate(r):
            c.drawString(40 + i * 120, y, v)
        y -= 18
    c.save(); return buf.getvalue()
