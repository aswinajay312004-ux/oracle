"""Reading review: each test pins a way a buyer file used to be read silently wrong (a quantity, a code, a date or a
whole colour block) and what is read now — the right value, or a refusal / question / review item naming the cells.
A confident wrong answer is the one outcome none of these may give."""
from datetime import datetime
from io import BytesIO
import json
import zipfile
import pytest
from openpyxl import Workbook, load_workbook
from backend import ai, matcher, readers, workbooks
from backend.erp import planner, mapping, catalog
from fixtures import save, merch_bom

PW = 'test-password-123'


def book(rows, title='PO', at=None):
    wb = Workbook(); ws = wb.active; ws.title = title
    for r in rows:
        ws.append(list(r))
    for ref, v in (at or {}).items():
        ws[ref] = v
    return wb, ws


def sheet_of(content, name='x.xlsx'):
    return readers.read(content, name)[1][0]


def applied(content, name='x.xlsx'):
    """The orders a person gets by applying the matcher's suggestion (the layout screen's Apply)."""
    sheets = readers.read(content, name)[1]
    s = matcher.suggest(sheets[0])
    return [workbooks.finish(o) for o in workbooks.apply_template(sheets, sheets[0], s['definition'])], s


def lines_of(o):
    return [(line['color'], [(s['size'], s['quantity']) for s in line['sizes']]) for line in o['lines']]


def ready(o, **kw):
    return dict(o, buyer='ZEN KIDS', description='TEE', po_number=o.get('po_number') or 'P-1', order_ref=kw.pop('ref', 'R/1'),
                order_date='2026-10-01', ship_date=o.get('ship_date') or '2026-11-01', quantity_basis='sample', qty_unit='PCS',
                pieces_per_unit=1, **kw)


def plan(*orders):
    return planner.plan(catalog.SnapshotCatalog(), list(orders), 'f' * 64, mapping.effective())


def errors(p):
    return {i['path']: i['message'] for i in p['issues'] if i['severity'] == 'error'}


def reviews(o, field=None):
    return [i['message'] for i in o.get('review') or [] if field is None or i['field'] == field]


HEAD = [['Buyer', 'ZEN KIDS'], ['Style No', 'ZK-100'], ['PO No', 'PO-555'], []]


# 1 — a "Total" under each colour block used to end the size matrix: later colours were dropped
def test_every_colour_block_after_its_own_total_row_is_read_and_each_total_is_checked():
    rows = [['Colour', 'S', 'M', 'L', 'Total'], ['RED', 10, 20, 30, 60], ['Total', 10, 20, 30, 60], ['BLUE', 5, 6, 7, 18],
            ['Total', 5, 6, 7, 18], ['GREEN', 1, 2, 3, 6], ['Total', 1, 2, 3, 6], ['Grand Total', 16, 28, 40, 84]]
    (o,), s = applied(save(book(HEAD + rows)[0]))
    assert lines_of(o) == [('RED', [('S', 10), ('M', 20), ('L', 30)]), ('BLUE', [('S', 5), ('M', 6), ('L', 7)]),
                           ('GREEN', [('S', 1), ('M', 2), ('L', 3)])]
    assert not reviews(o) and plan(ready(o))['totals']['pieces'] == 84
    rows[4] = ['Total', 5, 6, 9, 20]                                       # BLUE's own total disagrees with its rows
    (o,), _ = applied(save(book(HEAD + rows)[0]))
    assert len(o['lines']) == 3 and reviews(o, 'lines') == [
        'Row 9 (“Total”) on “PO” says the total is 20, but row 8 above it adds up to 18 — check for missing, hidden or extra rows '
        'before saving.']
    rows = [['Colour', 'S', 'M', 'L'], ['RED', 1, 2, 3], ['Sub Total', 1, 2, 3], ['BLUE', 3, 4, 5], ['Grand Total', 4, 6, 8], ['OLD', 9, 9, 9]]
    (o,), _ = applied(save(book(HEAD + rows)[0]))                         # a grand total ends the table; rows under it are named
    assert [c for c, _ in lines_of(o)] == ['RED', 'BLUE'] and any('ends with the grand total in row 9' in m for m in reviews(o))


# 2 — total rows labelled "TTL", "Colour Total", "Order Total", "<style> Total", "SUM" used to become colours or sizes
@pytest.mark.parametrize('label', ['TTL', 'G. TTL', 'G.TTL', 'Colour Total', 'Order Total', 'ZK-100 Total', 'SUM', 'Total for RED'])
def test_total_labels_are_never_colours_or_sizes(label):
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L'], ['RED', 10, 20, 30], ['BLUE', 5, 6, 7], [label, 15, 26, 37]])[0]))
    assert [c for c, _ in lines_of(o)] == ['RED', 'BLUE'] and not reviews(o)
    wb, _ = book([['Buyer', 'ZEN KIDS'], [], ['Style', 'Colour', 'Size', 'Qty'], ['ZK-1', 'RED', 'S', 10], ['ZK-1', 'RED', 'M', 20],
                  ['ZK-1', 'RED', label, 30]])
    assert lines_of(workbooks.parse(save(wb))['orders'][0]) == [('RED', [('S', 10), ('M', 20)])]


def test_a_colour_that_starts_with_a_total_word_stays_a_colour():
    assert workbooks.total_kind('TOTAL ECLIPSE') == 'maybe' and not workbooks.is_total('TOTAL ECLIPSE')
    assert [workbooks.total_kind(x) for x in ('Sub-Total', 'G. TTL', 'Colour Total', 'Total Qty:', 'Summer')] == ['sub', 'grand', 'total', 'total', None]
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L'], ['RED', 10, 20, 30], ['TOTAL ECLIPSE', 3, 4, 5], ['NAVY', 1, 1, 1]])[0]))
    assert [c for c, _ in lines_of(o)] == ['RED', 'TOTAL ECLIPSE', 'NAVY'] and not reviews(o)
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L'], ['RED', 10, 20, 30], ['NAVY', 1, 1, 1], ['TOTAL RED NAVY', 11, 21, 31]])[0]))
    assert [c for c, _ in lines_of(o)] == ['RED', 'NAVY']                 # the sums of the rows above: a total row
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L'], ['RED', 10, 20, 30], ['TOTAL ECLIPSE', 10, 20, 30]])[0]))
    assert [c for c, _ in lines_of(o)] == ['RED', 'TOTAL ECLIPSE']        # one row above: could be either — read, and asked
    assert reviews(o, 'lines') == ['Row 7 (“TOTAL ECLIPSE”) on “PO” repeats the quantities of row 6; it was read as a colour, not a '
                                   'total row — check it.']


# 3 — a second size block on the same header row (Delivery 1 | Delivery 2) used to be dropped silently
def test_a_second_size_block_on_the_header_row_is_refused_with_its_cells():
    wb, ws = book(HEAD + [[None, 'Delivery 1', None, None, None, 'Delivery 2'], ['Colour', 'S', 'M', 'L', 'Total', 'S', 'M', 'L', 'Total'],
                          ['RED', 10, 20, 30, 60, 1, 2, 3, 6], ['BLUE', 5, 5, 5, 15, 7, 7, 7, 21]])
    sheet = sheet_of(save(wb))
    s = matcher.suggest(sheet)
    assert s['definition']['table'] is None
    reason = next(r['reason'] for r in s['rejected'] if r['field'] == 'size')
    assert reason.startswith('F6:H6 also hold sizes (S, M, L) with quantities under them — a second size block')
    d = {'fields': {}, 'table': {'type': 'matrix', 'header_row': 6, 'color_col': 1, 'first_size_col': 2, 'last_size_col': 4, 'total_col': 5}}
    ok, rejected = matcher.check(sheet, d)                                # a person's (or a template's) mapping too
    assert ok['table'] is None and rejected[0]['cell'] == 'F6:H6'


# 4 — sizes the pattern did not know (2T, 1X, 2XS, XXXXL) and headings next to the sizes used to be dropped
def test_toddler_plus_and_extended_sizes_are_sizes_and_a_heading_next_to_them_is_named():
    assert all(matcher.size_like(x) for x in ('2T', '4T', '1X', '3X', '2XS', 'XXXS', 'XXXXL', '5XL'))
    assert not any(matcher.size_like(x) for x in ('RED', 'NAVY', 'TOTAL', 'Qty', 'ZK-100', 'XYZ', 'TEAL'))
    for sizes in (['12M', '18M', '24M', '2T', '3T', '4T'], ['2XS', 'XS', 'S', 'XL', 'XXL', 'XXXXL'], ['S', 'M', 'L', '1X', '2X', '3X']):
        (o,), _ = applied(save(book(HEAD + [['Colour'] + sizes, ['RED', 1, 2, 3, 4, 5, 6]])[0]))
        assert [s['size'] for s in o['lines'][0]['sizes']] == sizes
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L', 'Extra'], ['RED', 1, 2, 3, 40], ['BLUE', 1, 1, 1, 9]])[0]))
    assert [s['size'] for s in o['lines'][0]['sizes']] == ['S', 'M', 'L'] and reviews(o, 'lines') == [
        'E5 “Extra”, right next to the size columns B–D on “PO”, has numbers under it but was not read as a size. If it is a '
        'size, map it with the other sizes; otherwise check that no quantity is missing.']
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L', 'Price'], ['RED', 1, 2, 3, 4.5]])[0]))
    assert not reviews(o)                                                 # a price column is not a forgotten size


# 5 — a saved layout read the next file's "Total" column as a size
def test_a_saved_layout_never_reads_a_total_column_as_a_size():
    first = sheet_of(save(book(HEAD + [['Colour', 'S', 'M', 'L', 'XL'], ['RED', 10, 20, 30, 40]])[0]))
    d = matcher.suggest(first)['definition']
    t = dict(id='t1', name='ZEN KIDS PO', version=1, signature=workbooks.layout_signature(first, d),
             labels=sorted(workbooks.layout_labels(first, d)), definition=d, anchors=matcher.anchors(first, d), unit=None, quantity_basis=None)
    nxt = workbooks.parse(save(book(HEAD + [['Colour', 'S', 'M', 'L', 'Total'], ['NAVY', 5, 6, 7, 18]])[0]), [t])
    assert nxt['templates_used'] == [] and nxt['orders'] == []
    assert 'E5 holds “Total”, a total or quantity column, inside the size columns B–E — it would be read as a size' in nxt['notes'][0]


# 6 — "2.000" (two thousand in a German export, or 2?) used to be read as 2 outside semicolon CSV
def test_dot_grouped_quantities_are_unclear_in_every_format():
    rows = [('Style', 'Colour', 'Size', 'Qty'), ('ZK-1', 'ROT', 'S', '2.000'), ('ZK-1', 'ROT', 'M', '1.250'), ('ZK-1', 'ROT', 'L', '7')]
    wb, _ = book(rows)
    files = {'order.txt': '\r\n'.join('\t'.join(r) for r in rows).encode('utf-16'),
             'order.csv': '\n'.join(','.join(f'"{x}"' for x in r) for r in rows).encode(),
             'export.xls': ('<table>' + ''.join('<tr>' + ''.join(f'<td>{x}</td>' for x in r) + '</tr>' for r in rows) + '</table>').encode(),
             'order.json': json.dumps([dict(zip(rows[0], r)) for r in rows[1:]]).encode(),
             'order.xlsx': save(wb)}
    for name, content in files.items():
        o = workbooks.parse(content, filename=name)['orders'][0]
        sizes = o['lines'][0]['sizes']
        assert [(s['size'], s['quantity'], s.get('unclear')) for s in sizes] == [('S', '2.000', '2.000'), ('M', '1.250', '1.250'), ('L', 7, None)], name
        assert any('written with a dot between thousands or a decimal point (for example “2.000”: 2000 or 2?)' in m for m in reviews(o, 'lines')), name
        bad = errors(plan(ready(o)))
        assert '“2.000” could mean 2000' in bad['orders.0.lines.0.sizes.0.quantity'], name
    assert workbooks.number('0.250') == 0.25 and readers.typed('0.250') == 0.25    # a consumption of 0.250 kg is not a grouped number


def test_merch_bom_and_contract_sizes_go_through_the_same_unclear_rule():
    wb = load_workbook(BytesIO(merch_bom())); ws = wb.active
    ws['E12'] = '2.000'; ws['I12'] = '10.000'                            # text quantities in the buyer-order block
    o = workbooks.parse(save(wb))['orders'][0]
    first = o['order_lines'][0]
    assert (first['sizes'][0]['quantity'], first['sizes'][0]['unclear'], first['declared_total']) == ('2.000', '2.000', '10.000')
    assert any('E12' in m and 'I12' in m and 'dot between thousands' in m for m in reviews(o, 'lines'))


# 7 — an unquoted “1,200” in a CSV row split into two columns and was read as 1
def test_unquoted_thousands_in_a_text_file_are_not_read_as_the_first_part():
    o = workbooks.parse(b'Style,Colour,Size,Qty\nZK-1,RED,S,1,200\nZK-1,RED,M,300\nZK-1,RED,L,2,500\n', filename='erp.csv')['orders'][0]
    assert lines_of(o) == [('RED', [('S', '1,200'), ('M', 300), ('L', '2,500')])]
    assert reviews(o, 'lines') == ['Rows 2, 4 on “Sheet1” have more values than headings (E2 “200” under no heading): a number written '
                                   'with a comma, such as “1,200”, may have been split in this text file. Type the quantities of these '
                                   'rows as whole numbers.']
    bad = errors(plan(ready(o)))
    assert '“1,200” could mean 1200' in bad['orders.0.lines.0.sizes.0.quantity'] and 'orders.0.lines.0.sizes.1.quantity' not in bad
    with pytest.raises(ValueError, match='more values than headings'):     # a size matrix: which size was split is unknown
        applied(b'Buyer,ZEN KIDS\nStyle No,ZK-1\n\nColour,S,M,L\nRED,10,1,200,30\nBLUE,1,2,3\n', 'm.csv')


# 8 — labels in one row with values below: the next label used to be read as the value
def test_a_strip_of_labels_reads_each_value_below_its_label():
    wb, _ = book([['PO No', 'Buyer Ref', 'Style', 'Ex-Fty'], ['PO-4471', 'BR-22', 'ZK-2041', '15/01/2027'], [],
                  ['Colour', 'Size', 'Qty'], ['RED', 'S', 10], ['RED', 'M', 20]])
    o = workbooks.parse(save(wb))['orders'][0]
    assert (o['style'], o['po_number'], o['provenance']['style']['cell']) == ('ZK-2041', 'PO-4471', 'C2')
    s = matcher.suggest(sheet_of(save(wb)))
    assert s['fields']['po_number']['cell'] == 'A2' and s['fields']['style']['cell'] == 'C2'
    wb, _ = book([['Buyer', 'ZEN KIDS', 'PO No', 'PO-1'], ['Fabric', 'Cotton', 'Wash', 'Enzyme'], [],
                  ['Style', 'Colour', 'Size', 'Qty'], ['ZK-1', 'RED', 'S', 10]])
    o = workbooks.parse(save(wb))['orders'][0]                            # label | value pairs: still the value to the right
    assert (o['buyer'], o['po_number']) == ('ZEN KIDS', 'PO-1')


# 9 — two labels of one field with different values: the first used to win silently, marked high confidence
def test_two_labels_of_one_field_with_different_values_are_asked():
    wb, _ = book([['Buyer', 'ZEN KIDS'], ['Item', 'Baby Romper'], ['Style No', 'ZK-2041'], ['Order No', 'SO-889'], ['PO No', 'PO-4471'],
                  [], ['Colour', 'Size', 'Qty'], ['RED', 'S', 10]])
    d = workbooks.parse(save(wb))
    assert d['orders'] == [] and ('A2 “Item” says “Baby Romper” and A3 “Style No” says “ZK-2041” — both name the style; choose the '
                                  'cell to use.') in d['notes'][0]
    s = matcher.suggest(d['sheets'][0])
    assert s['fields']['style']['confidence'] == 'low' and s['fields']['po_number']['confidence'] == 'low'
    assert 'style' not in s['definition']['fields'] and 'po_number' not in s['definition']['fields'] and 'order_ref' not in s['fields']
    assert any('A3 “Style No” also names Style #' in e for e in s['fields']['style']['evidence'])
    wb, _ = book([['Order No', 'SO-889'], ['PO No', 'PO-4471'], [], ['Style', 'Colour', 'Size', 'Qty'], ['ZK-1', 'RED', 'S', 10]])
    o = workbooks.parse(save(wb))['orders'][0]                            # style is a column: the order is read, the PO is asked
    assert o['po_number'] == '' and o['provenance']['po_number']['question'] == (
        'A1 “Order No” says “SO-889” and A2 “PO No” says “PO-4471” — which is the PO number? Enter it.')
    assert reviews(o, 'po_number') and 'orders.0.po_number' in errors(plan(dict(ready(o), po_number='')))


# 12 — a workbook only the backup reader can open loses leading zeros and hidden rows: now said on the fields
def test_what_the_backup_reader_cannot_see_is_a_review_item():
    wb, ws = book([['Buyer', 'ZEN KIDS'], ['PO No', 'PO-1'], [], ['Style', 'Colour', 'Size', 'Qty'], [123, 'RED', 'S', 10], [123, 'RED', 'M', 20]])
    for r in (5, 6):
        ws.cell(r, 1).number_format = '00000'
    src = zipfile.ZipFile(BytesIO(save(wb))); out = BytesIO()
    with zipfile.ZipFile(out, 'w') as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.endswith(('.xml', '.rels')):
                data = (data.replace(b'http://schemas.openxmlformats.org/spreadsheetml/2006/main', b'http://purl.oclc.org/ooxml/spreadsheetml/main')
                        .replace(b'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
                                 b'http://purl.oclc.org/ooxml/officeDocument/relationships'))
                if info.filename == 'xl/workbook.xml':
                    data = data.replace(b'<workbook ', b'<workbook conformance="strict" ', 1)
            dst.writestr(info.filename, data)
    d = workbooks.parse(out.getvalue(), filename='strict.xlsx')
    o = d['orders'][0]
    assert o['style'] == '123' and any('Strict Open XML' in n for n in d['notes'])
    assert reviews(o, 'style') == ['The style number “123” was read from “PO” without its cell format, so leading zeros (such as “00123”) '
                                   'may be missing — check it against the file.']
    assert any('Hidden rows and columns of “PO” cannot be seen' in m for m in reviews(o, 'lines'))


# 13 — "120 pcs" in a trims list under the order table decided the order's unit
def test_the_unit_is_read_only_from_the_rows_the_table_read():
    wb, _ = book(HEAD + [['Colour', 'S', 'M', 'L', 'Total'], ['RED', 10, 20, 30, 60], [], [], ['Trims'], ['Elastic', '120 pcs'], ['Tag', '75 pcs']])
    (o,), s = applied(save(wb))
    assert o['qty_unit'] == '' and s['unit']['value'] == '' and 'B11' not in str(o['provenance']['qty_unit'])
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L'], ['RED', '10 pcs', '20 pcs', '5 pcs']])[0]))
    assert o['qty_unit'] == 'PCS'                                          # quantities written with their unit still count


# 14 — Excel error values (#REF!, #N/A) were read as a style, PO or order reference
def test_excel_error_values_are_never_read_as_values():
    b = readers.Builder('PO', capabilities=readers.FULL_CAPS, reader='openpyxl')
    for r, (label, value, err) in enumerate([('Buyer', 'ZEN KIDS', False), ('PO No', '#N/A', True), ('Style No', 'ZK-1', False),
                                             ('Order Ref', '#DIV/0!', True)], 1):
        b.put(r, 1, label); b.put(r, 2, value, error=err)
    for r, row in enumerate([('Colour', 'Size', 'Qty'), ('RED', 'S', 10), ('#REF!', 'M', 20)], 6):
        for c, v in enumerate(row, 1):
            b.put(r, c, v, error=v == '#REF!')
    o = workbooks.parse_sheets([b.sheet()])['orders'][0]
    assert (o['po_number'], o['order_ref'], o['style']) == ('', '', 'ZK-1')
    assert o['provenance']['po_number']['question'] == 'B2 holds the Excel error “#N/A”, not a PO number — enter the PO number.'
    assert [line['color'] for line in o['lines']] == ['RED', ''] and any('A8 on “PO” holds Excel errors' in m for m in reviews(o, 'lines'))
    assert matcher.value_score('code', 'style', '#REF!')[0] == 0
    ok, rejected = matcher.check(b.sheet(), {'fields': {'po_number': 'B2'}, 'table': None})
    assert not ok['fields'] and 'holds the Excel error “#N/A”' in rejected[0]['reason']


# 15 — a "Ratio" row inside a size matrix was read as a colour
def test_ratio_and_pack_count_rows_are_not_colours():
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L', 'Total'], ['Ratio', 1, 2, 1, 4], ['RED', 100, 200, 100, 400],
                                        ['BLUE', 50, 100, 50, 200], ['No of Packs', 150, None, None, 150]])[0]))
    assert [c for c, _ in lines_of(o)] == ['RED', 'BLUE']
    msgs = reviews(o, 'qty_unit')
    assert msgs[0].startswith('Row 6 (“Ratio”) on “PO” is a size ratio (how a pack is made up), not quantities') and 'No of Packs' in msgs[1]
    assert any(i['path'] == 'orders.0.qty_unit' and i['severity'] == 'warning' for i in plan(ready(o))['issues'])


# 16 — a style column (merged down over its colours) used to put every line under one order without a style
def test_a_style_column_gives_one_order_per_style_and_merged_styles_carry_down():
    wb, ws = book([['Buyer', 'ZEN KIDS'], ['PO No', 'PO-77'], [], ['Style', 'Colour', 'Sizes', None, None, None, 'Total'],
                   [None, None, 'S', 'M', 'L', 'XL'], ['ZK-1', 'RED', 10, 20, 30, 40, 100], [None, 'BLUE', 1, 2, 3, 4, 10],
                   ['ZK-2', 'NAVY', 5, 5, 5, 5, 20]])
    ws.merge_cells('A4:A5'); ws.merge_cells('B4:B5'); ws.merge_cells('C4:F4'); ws.merge_cells('G4:G5'); ws.merge_cells('A6:A7')
    orders, _ = applied(save(wb))
    assert [(o['style'], o['provenance']['style']['cell'], [c for c, _ in lines_of(o)]) for o in orders] == \
        [('ZK-1', 'A6', ['RED', 'BLUE']), ('ZK-2', 'A8', ['NAVY'])]
    ws.unmerge_cells('A6:A7')                                              # not merged: which style row 7 belongs to is asked
    orders, _ = applied(save(wb))
    assert [o['style'] for o in orders] == ['ZK-1', '', 'ZK-2']
    assert orders[1]['provenance']['style']['question'] == 'Row 7 on “PO” has no style in column A, while other rows have one — enter the style.'


# 17 — the heuristic fallback took a row of quantities for the size headings
def test_the_heuristic_never_takes_a_quantity_row_for_the_size_headings():
    rows = [['Buyer', 'DENIM CO'], ['Style No', 'D-501'], [], ['Fabric Colour', 28, 30, 32, 34, 'Pcs'],
            ['INDIGO', 100, 200, 300, 400, 1000], ['BLACK', 10, 20, 30, 40, 100]]
    sheet = sheet_of(save(book(rows)[0]))
    assert ai.heuristic(sheet)['table']['header_row'] == 4
    s = ai.suggest(sheet, use_ai=False)
    assert s['definition']['table']['header_row'] == 4 and s['definition']['table']['color_col'] == 1
    assert s['table']['confidence'] == 'medium'                           # "Fabric Colour" is only like "colour": checked by a person
    rows[3] = ['', 28, 30, 32, 34]                                        # no headings beside the numbers: no table is guessed
    assert ai.heuristic(sheet_of(save(book(rows)[0]))) is None


# 18 — codes written with decimals lost their trailing zeros (5012.10 became 5012.1)
def test_codes_keep_the_text_the_file_shows():
    csv = b'Buyer,ZEN KIDS\nPO No,4500.10\n\nStyle,Colour,Size,Qty\n5012.10,RED,S,10.00\n5012.10,RED,M,20\n'
    o = workbooks.parse(csv, filename='codes.csv')['orders'][0]
    assert (o['style'], o['po_number'], lines_of(o)) == ('5012.10', '4500.10', [('RED', [('S', 10), ('M', 20)])])
    html = ('<table><tr><td>PO No</td><td x:num="4500.1">4500.10</td></tr><tr></tr><tr><td>Style</td><td>Size</td><td>Qty</td></tr>'
            '<tr><td>5012.10</td><td>S</td><td>10</td></tr></table>').encode()
    o = workbooks.parse(html, filename='export.xls')['orders'][0]
    assert (o['style'], o['po_number']) == ('5012.10', '4500.10')
    wb, ws = book([['PO No', 4500.1], [], ['Style', 'Size', 'Qty'], [5012.1, 'S', 10]])
    ws['B1'].number_format = '0.00'; ws['A4'].number_format = '0.00'
    o = workbooks.parse(save(wb))['orders'][0]
    assert (o['style'], o['po_number']) == ('5012.10', '4500.10')


# 19 — headings merged over two rows: zero lines were read, and the planner said "provide 1–100 garment lines"
def test_a_header_merged_over_two_rows_reads_its_lines_and_no_lines_is_a_plain_refusal():
    wb, ws = book(HEAD + [['Colour', 'S', 'M', 'L', 'Total'], [], ['RED', 10, 20, 30, 60], ['BLUE', 5, 5, 5, 15]])
    for c in 'ABCDE':
        ws.merge_cells(f'{c}5:{c}6')
    (o,), _ = applied(save(wb))
    assert lines_of(o) == [('RED', [('S', 10), ('M', 20), ('L', 30)]), ('BLUE', [('S', 5), ('M', 5), ('L', 5)])]
    sheet = sheet_of(save(book(HEAD + [['Colour', 'S', 'M', 'L'], ['Total', 0, 0, 0]])[0]))
    d = {'fields': {}, 'table': {'type': 'matrix', 'header_row': 5, 'color_col': 1, 'first_size_col': 2, 'last_size_col': 4}}
    with pytest.raises(ValueError, match='No order lines were found under the headings in row 5 of “PO”'):
        workbooks.apply_template([sheet], sheet, d)


# 20 — an unlabelled row of sums: a colour-less line (matrix) or "style is missing" for the whole table (flat)
def test_an_unlabelled_row_of_sums_is_a_total_row():
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L'], ['RED', 10, 20, 30], ['BLUE', 5, 5, 5], [None, 15, 25, 35]])[0]))
    assert [c for c, _ in lines_of(o)] == ['RED', 'BLUE'] and not reviews(o)
    (o,), _ = applied(save(book(HEAD + [['Colour', 'S', 'M', 'L'], ['RED', 10, 20, 30], [None, 10, 20, 30]])[0]))
    assert [c for c, _ in lines_of(o)] == ['RED'] and reviews(o, 'lines') == [
        'Row 7 on “PO” has no colour and repeats the quantities of row 6, so it was read as a total row, not a second line — check it.']
    wb, _ = book([['Buyer', 'ZEN KIDS'], [], ['Style', 'Colour', 'Size', 'Qty'], ['ZK-1', 'RED', 'S', 10], ['ZK-1', 'RED', 'M', 20], [None, None, None, 30]])
    o = workbooks.parse(save(wb))['orders'][0]
    assert lines_of(o) == [('RED', [('S', 10), ('M', 20)])] and not reviews(o)
    wb, _ = book([['Style', 'Colour', 'Size', 'Qty'], ['ZK-1', 'RED', 'S', 10], [None, None, None, 99]])
    assert workbooks.parse(save(wb))['orders'] == []                      # not the sum: still refused, never a line without a style


# 21 — a 1904 workbook (Mac Excel) with dates typed as plain numbers was read 4 years and a day early
def test_dates_typed_as_numbers_in_a_1904_workbook():
    from openpyxl.utils.datetime import CALENDAR_MAC_1904
    wb, _ = book([['Ship Date', (datetime(2027, 1, 15) - datetime(1904, 1, 1)).days], [], ['Style', 'Size', 'Qty'], ['ZK-1', 'S', 10]])
    wb.epoch = CALENDAR_MAC_1904
    d = workbooks.parse(save(wb))
    assert d['sheets'][0]['date1904'] is True and d['orders'][0]['ship_date'] == '2027-01-15'
    assert 'counts dates from 1904' in d['orders'][0]['provenance']['ship_date']['derivation']
    wb, _ = book([['Ship Date', 46402], [], ['Style', 'Size', 'Qty'], ['ZK-1', 'S', 10]])
    assert workbooks.parse(save(wb))['orders'][0]['ship_date'] == '2027-01-15'     # a 1900 workbook is unchanged


# 22 — "1E+400" crashed the upload (HTTP 500)
def test_numbers_beyond_any_float_are_not_numbers():
    assert workbooks.number('1E+400') is None and workbooks.qty_value(float('inf')) == (None, None)
    rows = b'Style,Colour,Size,Qty\nZK-1,RED,S,1E+400\nZK-1,RED,M,20\nZK-1,RED,L,30\nZK-1,RED,XL,40\nZK-1,RED,XXL,50\n'
    o = workbooks.parse(rows, filename='x.csv')['orders'][0]
    assert o['lines'][0]['sizes'][0]['quantity'] == '1E+400' and 'orders.0.lines.0.sizes.0.quantity' in errors(plan(ready(o)))


# Oracle review 5 — an ambiguous date read day first was only in the order's notes, not under "Please read before approving"
def test_an_ambiguous_day_first_date_is_a_warning_on_its_field():
    o = workbooks.finish(dict(workbooks.base_order(), ship_date='05/06/2026', order_date='01/02/2026'))
    assert (o['ship_date'], o['order_date']) == ('2026-06-05', '2026-02-01')
    assert reviews(o, 'ship_date') == ['Ship date: "05/06/2026" read as day/month/year (2026-06-05) — check it.']
    p = plan(dict(ready(o), ship_date=o['ship_date'], order_date=o['order_date']))
    assert {i['path'] for i in p['issues'] if i['severity'] == 'warning'} >= {'orders.0.ship_date', 'orders.0.order_date'}


# 10, 11 — POST /map dropped the review items and replaced the other sheets' orders
@pytest.fixture
def client(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from backend import local_store, main
    from fake_oracle import FakeOracle
    monkeypatch.setattr(local_store, 'DATA', tmp_path); monkeypatch.setattr(main, 'DATA', tmp_path)
    monkeypatch.setattr(ai, 'available', lambda: False)
    FakeOracle().install(monkeypatch)
    with TestClient(main.app) as c:
        assert c.post('/api/setup', json={'username': 'admin', 'password': PW}).status_code == 200
        yield c


def upload(c, wb):
    r = c.post('/api/uploads', files={'file': ('po.xlsx', save(wb))}); assert r.status_code == 200, r.text
    return r.json()


def map_sheet(c, d, sheet):
    s = c.post(f'/api/drafts/{d["id"]}/suggest', json={'sheet': sheet, 'use_ai': False}).json()
    r = c.post(f'/api/drafts/{d["id"]}/map', json={'version': d['version'], 'sheet': sheet, 'definition': s['definition'], 'qty_unit': 'PCS'})
    assert r.status_code == 200, r.text
    return r.json()


def test_a_mapped_layout_keeps_the_review_items_of_its_sheet(client):
    wb, ws = book([['Buyer', 'ZEN KIDS'], ['Style No', 'ZK-1'], ['PO No', 'PO-1'], ['Ship Date', '15/01/2027'], ['Revised Ship Date', '20/02/2027'],
                   [], ['Colour', 'S', 'M', 'L'], ['RED', 10, 20, 30], ['PINK', 1, 1, 1], ['BLUE', 5, 5, 5]])
    ws.row_dimensions[9].hidden = True
    d = map_sheet(client, upload(client, wb), 'PO')
    o = d['orders'][0]
    hidden = 'Hidden row 9 on “PO” holds data and was not read. Unhide it in Excel and upload again if it belongs to the order.'
    assert hidden in d['notes'] and hidden in reviews(o, 'lines') and o['sheet'] == 'PO'
    assert reviews(o, 'ship_date') == ['A5 “Revised Ship Date” says 2027-02-20, but the ship date was read from B4 (2027-01-15) — check '
                                       'which date is meant.']


def test_orders_on_several_sheets_are_kept_and_a_sheet_left_out_is_named(client):
    wb = Workbook(); a = wb.active; a.title = 'PO 1001'
    for row in [['Buyer', 'ZEN KIDS'], ['PO No', 'PO-1001'], [], ['Style', 'Colour', 'Size', 'Qty'], ['ZK-1', 'RED', 'S', 10], ['ZK-1', 'RED', 'M', 20]]:
        a.append(row)
    for name, style in (('PO 1002', 'ZK-2'), ('PO 1003', 'ZK-3')):
        b = wb.create_sheet(name)
        for row in [['Buyer', 'ZEN KIDS'], ['PO No', name], ['Style No', style], [], ['Colour', 'S', 'M', 'L'], ['NAVY', 100, 200, 300]]:
            b.append(row)
    d = upload(client, wb)
    left = [f'No order was read from sheet “{n}”, which holds data. If it belongs to this order, map that sheet on the layout screen.'
            for n in ('PO 1002', 'PO 1003')]
    assert [o['style'] for o in d['orders']] == ['ZK-1'] and d['notes'] == left and reviews(d['orders'][0], 'lines') == left
    assert d['suggestion']['sheet'] == 'PO 1002'                         # the layout screen opens on a sheet that was not read
    d = map_sheet(client, d, 'PO 1002')
    assert [(o['style'], o['sheet']) for o in d['orders']] == [('ZK-1', 'PO 1001'), ('ZK-2', 'PO 1002')]
    assert d['notes'] == left[1:] and all(reviews(o, 'lines') == left[1:] for o in d['orders'])
    first_id = d['orders'][0]['id']
    d = map_sheet(client, d, 'PO 1002')                                  # mapping the same sheet again replaces only its orders
    d = map_sheet(client, d, 'PO 1003')
    assert [o['style'] for o in d['orders']] == ['ZK-1', 'ZK-2', 'ZK-3'] and d['orders'][0]['id'] == first_id
    assert d['notes'] == [] and not any(reviews(o) for o in d['orders'])
