"""Synthetic buyer workbooks (TEST data only) in several layouts."""
from datetime import date
from io import BytesIO
from openpyxl import Workbook


def save(wb):
    out = BytesIO(); wb.save(out); return out.getvalue()


def label_layout(order_no='PK/26-27/S/1320', buyer='ABC TEXTILES', printing=None):
    """`Label | value` header with a vertical size table. Printing is an optional section."""
    wb = Workbook(); ws = wb.active; ws.title = 'Buyer Order'
    ws['A1'] = 'BUYER PURCHASE ORDER'
    ws['A3'] = 'Buyer'; ws['B3'] = buyer
    ws['A4'] = 'Order No'; ws['B4'] = order_no
    ws['A5'] = 'Style No'; ws['B5'] = '0008'
    ws['A6'] = 'Delivery Date'; ws['B6'] = date(2026, 12, 15); ws['B6'].number_format = 'dd-mm-yyyy'
    ws['D3'] = 'Season'; ws['E3'] = 'SS27'
    ws['A9'] = 'Size'; ws['B9'] = 'Quantity'; ws['C9'] = 'Colour'
    for r, (size, qty) in enumerate([('M', 2000), ('L', 3000)], 10):
        ws.cell(r, 1, size); ws.cell(r, 2, qty); ws.cell(r, 3, 'NAVY')
    if printing:
        ws['A14'] = 'Printing ink'; ws['B14'] = printing
    return save(wb)


def horizontal_layout():
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    ws['A1'] = 'Customer: DEMO BUYER'
    ws['A2'] = 'PO Number'; ws['B2'] = 'PO-7781'
    ws['A4'] = 'Sizes'; ws['B4'] = 'S'; ws['C4'] = 'M'; ws['D4'] = 'L'; ws['E4'] = 'Total'
    ws['B5'] = 10; ws['C5'] = 20; ws['D5'] = 30; ws['E5'] = '=SUM(B5:D5)'
    return save(wb)


def moved_layout():
    """Same labels, different rows: a saved template must not be applied blindly."""
    wb = Workbook(); ws = wb.active; ws.title = 'Buyer Order'
    ws['A2'] = 'Note'; ws['B2'] = 'Moved'
    ws['A8'] = 'Buyer'; ws['B8'] = 'ABC TEXTILES'
    return save(wb)


def merch_bom(ocn='TEST-OC-1', sample_qty=0, extra_header=False):
    """Merch BOM layout: header block, OCN under its label, order / sample / summary size blocks."""
    from openpyxl.worksheet.formula import ArrayFormula
    wb = Workbook(); ws = wb.active; ws.title = 'ST1'
    for cell, value in {'B2': 'BUYER', 'C2': 'TEST BUYER', 'D3': 'FOB DATE', 'E3': date(2027, 2, 18), 'B3': 'SEASON', 'C3': 'SS-27',
                        'B4': 'ITEM', 'C4': 'ST-001', 'B5': 'DESCRIPTION', 'C5': 'TEST BODYSUIT', 'E5': 'TEST-PO-1',
                        'B6': 'MERCHANDISER', 'C6': 'TEST MERCH', 'B8': 'OCN', 'C8': 'QTY(IN PACKS)', 'D8': 'QTY(IN PIECES)'}.items():
        ws[cell] = value
    if ocn:
        ws['B9'] = ocn
    sizes = ['S', 'M', 'L']
    for header, role, qty in ((11, 'CT-PO', [5, 7, 9]), (17, 'SAMPLE', [sample_qty] * 3), (23, 'SUMMARY', [5, 7, 9])):
        ws.cell(header, 2, 'Garment '); ws.cell(header, 3, 'Description'); ws.cell(header, 4, 'CONTRACT')
        for i, s in enumerate(sizes):
            ws.cell(header, 5 + i, s)
        ws.cell(header, 8, 'Total')
        ws.cell(header + 1, 2, 'GARMENT-1'); ws.cell(header + 1, 3, 'NAVY'); ws.cell(header + 1, 4, role)
        for i, q in enumerate(qty):
            ws.cell(header + 1, 5 + i, q)
        ws.cell(header + 1, 8, sum(qty))
        ws.cell(header + 2, 2, 'Total')
    ws['J20'] = ArrayFormula('J20', '=E18')
    if extra_header:
        ws['G2'] = 'SHIP MODE'; ws['H2'] = 'SEA'; ws['G3'] = 'Payment Terms'; ws['H3'] = 'LC 60 DAYS'
    ws['B29'] = 'FABRIC REQUIREMENT'
    for i, h in enumerate(['Garment', 'Colour', 'Fabric', 'Gsm'], 2):
        ws.cell(30, i, h)
    ws['B31'] = 'GARMENT-1'; ws['C31'] = 'NAVY'; ws['D31'] = 'Interlock'; ws['E31'] = 190
    ws['B34'] = 'CARE LABEL'
    for i, h in enumerate(['Garment', 'Code', 'Size', 'Supplier'], 2):
        ws.cell(35, i, h)
    ws['B36'] = 'ALL'; ws['C36'] = 'CARE-1'
    return save(wb)
