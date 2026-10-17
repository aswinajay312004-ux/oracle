"""Excel understanding engine (docs/EXCEL_UNDERSTANDING.md §2–7): fields API and alias customisation,
matcher scores / evidence / one-to-one assignment, validator rejections, quantity-unit evidence and the
planner's unit rules, buyer template versioning, anchor and validator refusals, dates in templates."""
from copy import deepcopy
from datetime import datetime
import json
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from backend import ai, fields as F, local_store, main, matcher, workbooks
from backend.erp import planner, mapping, catalog
from fake_oracle import FakeOracle
from fixtures import save, zen_kids_po, wrong_mapping_po, flat_qty_orders, flat_orders, merch_bom

PW = 'test-password-123'
ZEN = {'fields': {'buyer': 'B4', 'style': 'B5', 'description': 'B6', 'po_number': 'B7', 'ship_date': 'B8', 'season': 'B9'},
       'table': {'type': 'matrix', 'header_row': 11, 'color_col': 1, 'first_size_col': 3, 'last_size_col': 6, 'total_col': 7}}


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(local_store, 'DATA', tmp_path); monkeypatch.setattr(main, 'DATA', tmp_path)
    monkeypatch.setattr(ai, 'available', lambda: False)
    FakeOracle().install(monkeypatch)
    with TestClient(main.app) as c:
        assert c.post('/api/setup', json={'username': 'admin', 'password': PW}).status_code == 200
        yield c


def upload(c, content, name='file.xlsx'):
    r = c.post('/api/uploads', files={'file': (name, content)}); assert r.status_code == 200, r.text
    return r.json()


def sheet_of(content):
    return workbooks.parse(content)['sheets'][0]


def article_code_sheet():
    wb = Workbook(); ws = wb.active; ws.title = 'Lines'
    ws.append(['Article Code', 'Colour', 'Size', 'Qty'])
    for row in [('AC-1', 'RED', 'S', 5), ('AC-1', 'RED', 'M', 7)]:
        ws.append(row)
    return save(wb)


def ready_order(**kw):
    o = workbooks.base_order()
    o.update(buyer='B', style='S1', description='D', po_number='P', order_ref='R/1', order_date='2026-10-17', ship_date='2026-11-01',
             quantity_basis='sample', lines=[dict(label='S1', color='RED', declared_total=None, sizes=[dict(size='M', quantity=10, source={})])])
    o.update(kw)
    return o


def plan(*orders, **kw):
    return planner.plan(catalog.SnapshotCatalog(), list(orders), 'e' * 64, mapping.effective(), **kw)


def errors(p):
    return {i['path']: i['message'] for i in p['issues'] if i['severity'] == 'error'}


# --- §2 fields ------------------------------------------------------------------------------------
def test_default_fields_are_complete_and_aliases_belong_to_one_field_only():
    keys = [f['key'] for f in F.DEFAULTS]
    assert {'buyer', 'style', 'description', 'po_number', 'order_ref', 'order_date', 'ship_date', 'sample_type', 'exporter',
            'season', 'merch', 'department', 'gender', 'against', 'packing_type', 'components', 'size', 'quantity', 'color', 'qty_unit'} == set(keys)
    owners = {}
    for f in F.DEFAULTS:
        for a in f['aliases']:
            assert owners.setdefault(F.norm(a), f['key']) == f['key'], f'alias {a} belongs to {owners[F.norm(a)]} and {f["key"]}'
    assert all(f['required'] for f in F.DEFAULTS if f['locked']) and not F.DEFAULT.get('department')['required']
    # compatibility names are still importable
    assert workbooks.ALIASES['style'][:2] == ['style', 'style no'] and 'buyer' in workbooks.HEADER_FIELDS and ai.SIZE is matcher.SIZE


def test_fields_api_lists_changes_and_audits(env):
    c = env
    got = {f['key']: f for f in c.get('/api/fields').json()}
    assert set(got['buyer']) >= {'key', 'label', 'kind', 'where', 'aliases', 'default_aliases', 'locked', 'required'}
    assert got['style']['aliases'] == got['style']['default_aliases'] and got['qty_unit']['locked'] and got['size']['where'] == 'table'
    r = c.put('/api/fields/style', json={'aliases': got['style']['aliases'] + ['Article Code', ' article  code ']})
    assert r.status_code == 200 and r.json()['aliases'][-1] == 'Article Code' and r.json()['aliases'].count('Article Code') == 1
    assert 'Article Code' not in r.json()['default_aliases'] and r.json()['changed_by'] == 'admin'
    event = next(a for a in c.get('/api/audit').json() if a['action'] == 'field_changed')
    assert event['target'] == 'style' and event['detail']['added'] == ['Article Code']
    r = c.put('/api/fields/buyer', json={'aliases': ['buyer', 'Article Code']})
    assert r.status_code == 409 and 'Style' in r.json()['detail']          # an alias belongs to one field only
    assert c.put('/api/fields/buyer', json={'required': False}).status_code == 409   # locked field stays required
    assert c.put('/api/fields/nonsense', json={'required': True}).status_code == 404
    r = c.put('/api/fields/department', json={'required': True})
    assert r.status_code == 200 and r.json()['required'] is True
    c.post('/api/users', json={'username': 'op', 'password': PW, 'role': 'operator'})
    c.post('/api/logout'); c.post('/api/login', json={'username': 'op', 'password': PW})
    assert c.put('/api/fields/season', json={'required': True}).status_code == 403
    assert c.get('/api/fields').status_code == 200


def test_alias_customisation_changes_parsing_without_global_state(env):
    content = article_code_sheet()
    assert workbooks.parse(content)['orders'] == []                     # "Article Code" is not a default alias
    custom = F.Fields(F.effective({'style': {'aliases': F.DEFAULT.aliases('style') + ['Article Code']}}))
    o = workbooks.parse(content, fields=custom)['orders'][0]
    assert o['style'] == 'AC-1' and o['provenance']['style']['cell'] in ('A2', 'A3')
    assert workbooks.parse(content)['orders'] == []                     # defaults untouched afterwards
    no_qty = F.Fields(F.effective({'quantity': {'aliases': ['quantity']}}))
    assert workbooks.parse(flat_orders(), fields=no_qty)['orders'] == []  # a removed alias is no longer recognised
    c = env                                                             # and through the API: upload uses the saved aliases
    assert upload(c, content)['orders'] == []
    aliases = next(f for f in c.get('/api/fields').json() if f['key'] == 'style')['aliases']
    assert c.put('/api/fields/style', json={'aliases': aliases + ['Article Code']}).status_code == 200
    d = upload(c, content)
    assert d['format'] == 'Tabular orders' and d['orders'][0]['style'] == 'AC-1'


def test_admin_required_field_is_checked_by_the_planner(env):
    p = plan(ready_order(qty_unit='PCS'), required=['department'])
    assert errors(p).get('orders.0.department', '').startswith('Order 1: Department is required')
    assert 'orders.0.department' not in errors(plan(ready_order(qty_unit='PCS', department='KIDS'), required=['department']))
    c = env
    d = upload(c, flat_qty_orders('Pcs'))
    o = dict(d['orders'][0], buyer='B', description='D', po_number='P', order_ref='R/2', order_date='2026-10-17',
             ship_date='2026-11-01', quantity_basis='sample')
    d = c.put(f'/api/drafts/{d["id"]}', json={'version': d['version'], 'orders': [o]}).json()
    assert not c.post(f'/api/drafts/{d["id"]}/preview', json={'version': d['version']}).json()['plan']['errors']
    c.put('/api/fields/gender', json={'required': True})
    p = c.post(f'/api/drafts/{d["id"]}/preview', json={'version': d['version'] + 1}).json()['plan']
    assert list(errors(p)) == ['orders.0.gender']


# --- §3 matcher ------------------------------------------------------------------------------------
def test_name_score_exact_abbreviations_and_trigrams():
    assert matcher.name_score('Customer', F.DEFAULT.aliases('buyer')) == (1.0, 'customer')
    assert matcher.name_score('Del Dt', ['delivery date'])[0] == 1.0          # del → delivery, dt → date
    assert matcher.name_score('Style #', ['style number'])[0] == 1.0          # # → number
    assert matcher.name_score('Ex-Fty Date', ['ex factory date'])[0] == 1.0
    assert matcher.name_score('Cust. Name', ['customer name'])[0] == 1.0
    fuzzy, alias = matcher.name_score('Colour Desc', F.DEFAULT.aliases('color'))
    assert 0.6 < fuzzy < 1.0 or (fuzzy == 1.0 and alias == 'colour description')
    assert matcher.name_score('Remark', F.DEFAULT.aliases('merch'))[0] < matcher.FLOOR


def test_matcher_scores_explain_and_assign_one_to_one():
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    ws['A1'], ws['B1'] = 'Buyer', 'ZEN KIDS'
    ws['A2'], ws['B2'] = 'Customer', 'ZEN KIDS RETAIL'        # a second buyer label: only one wins
    ws['A3'], ws['B3'] = 'Del Dt', '20/12/2026'
    ws['A4'], ws['B4'] = 'Season', 2027                        # a number is not a season name
    for c, h in enumerate(['Colour', 'S', 'M', 'L'], 1):
        ws.cell(6, c, h)
    ws.append(['NAVY', 1, 2, 3])
    s = matcher.suggest(sheet_of(save(wb)), masters={'buyer': ['ZEN KIDS']})
    f = s['fields']
    assert f['buyer']['cell'] == 'B1' and f['buyer']['score'] == 1.0 and 'value found in the buyer master' in f['buyer']['evidence']
    assert 'B2' not in [p['cell'] for p in f.values()]                         # one field, one cell
    assert f['ship_date']['cell'] == 'B3' and f['ship_date']['confidence'] == 'high'
    assert any('2026-12-20' in e for e in f['ship_date']['evidence']) and any('Del Dt' in e for e in f['ship_date']['evidence'])
    assert 'season' not in f and {'field': 'season', 'cell': 'B4'}.items() <= next(r for r in s['rejected'] if r['field'] == 'season').items()
    assert s['definition']['fields'] == {k: p['cell'] for k, p in f.items() if p['confidence'] != 'low'}
    assert s['table']['type'] == 'matrix' and s['table']['header_row'] == 6 and s['table']['confidence'] == 'high'
    assert s['definition']['table']['color_col'] == 1 and s['unit']['value'] == '' and s['source'] == 'matcher'
    assert all(p['source'] == 'matcher' and p['evidence'] for p in f.values())


def test_low_confidence_is_shown_but_never_preselected():
    wb = Workbook(); ws = wb.active
    ws['A1'], ws['B1'] = 'Merchant Name', 'RAVI'        # like "merchant" but not an alias
    ws['A3'], ws['B3'], ws['C3'] = 'Style', 'Size', 'Qty'
    ws.append(['X1', 'M', 4])
    s = matcher.suggest(sheet_of(save(wb)))
    m = s['fields'].get('merch')
    assert m and m['confidence'] in ('medium', 'low')
    if m['confidence'] == 'low':
        assert 'merch' not in s['definition']['fields']
    assert matcher.confidence(0.85) == 'high' and matcher.confidence(0.6) == 'medium' and matcher.confidence(0.59) == 'low'


def test_validator_rejects_with_reasons_and_keeps_the_rest():
    sheet = sheet_of(wrong_mapping_po())
    ok, rejected = matcher.check(sheet, {'fields': {'buyer': 'B1', 'style': 'B2', 'sql': 'A1', 'po_number': 'ZZ'},
                                         'table': {'type': 'flat', 'header_row': 5, 'columns': {'size': 2, 'quantity': 3, 'color': 2}}})
    by = {r['field']: r['reason'] for r in rejected}
    assert ok['fields'] == {'buyer': 'B1', 'style': 'B2'} and 'not an order field' in by['sql'] and 'cell address' in by['po_number']
    assert 'already used for Size' in by['color'] and ok['table']['columns'] == {'size': 2, 'quantity': 3}
    ok, rejected = matcher.check(sheet, {'fields': {}, 'table': {'type': 'flat', 'header_row': 99, 'columns': {}}})
    assert ok['table'] is None and 'outside the sheet' in rejected[0]['reason']
    zen = sheet_of(zen_kids_po()); bad = deepcopy(ZEN); bad['table']['last_size_col'] = 7     # includes the Total column
    ok, rejected = matcher.check(zen, bad)
    assert ok['table'] is not None and not rejected                                            # 4 of 5 headings are sizes (≥ 60 %)
    bad['table'].update(first_size_col=1, last_size_col=2)
    ok, rejected = matcher.check(zen, bad)
    assert ok['table'] is None and rejected[0]['field'] == 'size' and 'do not look like sizes' in rejected[0]['reason']


def test_map_refusal_body_names_each_field(env):
    c = env
    d = upload(c, wrong_mapping_po())
    r = c.post(f'/api/drafts/{d["id"]}/map', json={'version': d['version'], 'sheet': 'PO',
               'definition': {'fields': {'style': 'B2', 'ship_date': 'H3'}, 'table': {'type': 'flat', 'header_row': 5, 'columns': {'size': 2, 'quantity': 4}}}})
    assert r.status_code == 409
    body = r.json()
    assert body['detail'].startswith('This layout does not fit the file.') and 'Ship Dt: H3 holds “NEXT”, which is not a date.' in body['detail']
    assert {(x['field'], x['cell'], x['source']) for x in body['rejected']} == {('ship_date', 'H3', 'person'), ('quantity', 'D5', 'person')}
    assert c.post(f'/api/drafts/{d["id"]}/map', json={'version': d['version'], 'sheet': 'PO', 'definition': ZEN, 'qty_unit': 'PACK',
                                                       'pieces_per_unit': 1}).status_code == 409


def test_ai_agreement_is_both_and_ai_only_fields_are_scored(env, monkeypatch):
    sheet = sheet_of(wrong_mapping_po())
    monkeypatch.setattr(ai, 'available', lambda: True)
    monkeypatch.setattr(ai, '_chat', lambda *a, **k: json.dumps({'fields': {'buyer': 'B1', 'merch': 'H2'}, 'table': {}, 'explanation': 'From the header.'}))
    s = ai.suggest(sheet)
    assert s['fields']['buyer']['source'] == 'both' and 'the local AI proposed the same cell' in s['fields']['buyer']['evidence']
    assert s['fields']['season']['cell'] == 'H2'                       # matcher's exact "Season" label beats the AI's guess
    assert any('local AI proposed H2' in e for e in s['fields']['season']['evidence']) or 'merch' not in s['definition']['fields']
    assert s['source'] == 'matcher + local AI' and s['explanation'].startswith('From the header.')
    monkeypatch.setattr(ai, '_chat', lambda *a, **k: 'not json')        # a broken model answer falls back to the matcher
    assert ai.suggest(sheet)['source'] == 'matcher'


def test_heuristic_is_validated(monkeypatch):
    sheet = sheet_of(zen_kids_po())
    d = ai.heuristic(sheet)
    assert d['table']['type'] == 'matrix' and d['table']['header_row'] == 11 and d['fields']['buyer'] == 'B4'
    ok, rejected = matcher.check(sheet, d)
    assert not rejected


# --- §5 quantity unit ------------------------------------------------------------------------------
@pytest.mark.parametrize('text,expected', [('Qty', ('', None)), ('Qty (pcs)', ('PCS', 1)), ('PIECES', ('PCS', 1)), ('Qty in packs', ('PACK', None)),
                                           ('3 PC PACK', ('PACK', 3)), ('Qty packs of 6', ('PACK', 6)), ('Ratio', ('PACK', None)),
                                           ('Qty pcs / packs', ('?', None)), ('PACKING TYPE', ('', None)), ('Sets', ('PACK', None))])
def test_unit_words(text, expected):
    assert workbooks.unit_words(text) == expected


def test_unit_evidence_on_orders_and_merch_rule():
    o = workbooks.parse(flat_qty_orders('Qty (3 pc pack)'))['orders'][0]
    assert (o['qty_unit'], o['pieces_per_unit']) == ('PACK', 3) and o['provenance']['qty_unit']['cell'] == 'D1'
    o = workbooks.parse(flat_qty_orders())['orders'][0]
    assert o['qty_unit'] == '' and o['provenance']['qty_unit']['question'] == 'The file says only “Qty” — are these pieces or packs?'
    m = workbooks.parse(merch_bom())['orders'][0]           # size total 150 ≠ QTY(IN PIECES) 200 → asked, with both numbers
    assert m['qty_unit'] == '' and m['pieces_per_unit'] is None and '150' in m['provenance']['qty_unit']['question']
    ok = workbooks.parse(merch_bom(sample=(1, 1, 1, 1)))
    assert ok['orders'][0]['provenance']['qty_unit']['cell'] == 'D9'


def test_planner_unit_rules_and_legacy_drafts():
    e = errors(plan(ready_order(provenance={'qty_unit': {'question': 'The file says only “Qty” — are these pieces or packs?'}})))
    assert e['orders.0.qty_unit'] == 'Order 1: the file says only “Qty” — are these pieces or packs?'
    assert errors(plan(ready_order(qty_unit='PACK')))['orders.0.pieces_per_unit'] == 'Order 1: how many pieces are in one pack?'
    assert 'at least 2' in errors(plan(ready_order(qty_unit='PACK', pieces_per_unit=1)))['orders.0.pieces_per_unit']
    assert 'must be 1' in errors(plan(ready_order(qty_unit='PCS', pieces_per_unit=6)))['orders.0.pieces_per_unit']
    assert 'Pieces or Packs' in errors(plan(ready_order(qty_unit='BOX')))['orders.0.qty_unit']
    assert plan(ready_order(qty_unit='PACK', pieces_per_unit=4))['totals']['pieces'] == 40
    p = plan(ready_order(qty_unit='PCS', pieces_per_unit=None))
    assert not p['errors'] and p['totals']['pieces'] == 10
    legacy = ready_order(pieces_per_unit=3); del legacy['qty_unit']        # saved before qty_unit existed
    p = plan(legacy)
    assert not p['errors'] and p['totals']['pieces'] == 30
    legacy['pieces_per_unit'] = None
    assert list(errors(plan(legacy))) == ['orders.0.pieces_per_unit']


# --- §6 templates ----------------------------------------------------------------------------------
def test_template_versions_by_name_or_signature_and_buyer(env):
    s = local_store
    a = s.save_template('Zen PO', 'sig-1', ['1:1:x'], ZEN, 'admin', buyer='ZEN KIDS', anchors={}, unit={'qty_unit': 'PCS', 'pieces_per_unit': 1})
    b = s.save_template('ZEN PO', 'sig-2', ['1:1:x'], ZEN, 'admin', buyer='OTHER', anchors={})          # same name (any case)
    c = s.save_template('Zen kids v3', 'sig-2', ['1:1:x'], ZEN, 'admin', buyer='other', anchors={})     # same signature + buyer
    other = s.save_template('Blue PO', 'sig-2', ['1:1:x'], ZEN, 'admin', buyer='BLUE', anchors={})      # same signature, other buyer
    assert (a['version'], b['version'], c['version'], other['version']) == (1, 2, 3, 1)
    assert a['family'] == b['family'] == c['family'] != other['family']
    rows = {t['id']: t for t in s.templates(active_only=False)}
    assert [rows[x['id']]['active'] for x in (a, b, c, other)] == [False, False, True, True]
    assert rows[a['id']]['unit'] == {'qty_unit': 'PCS', 'pieces_per_unit': 1} and rows[c['id']]['buyer'] == 'other'
    s.set_template_active(a['id'], True)                                  # reactivating v1 deactivates v3
    assert {t['id'] for t in s.templates()} == {a['id'], other['id']}
    listed = env.get('/api/templates').json()
    assert {'id', 'name', 'buyer', 'version', 'family', 'anchors', 'unit', 'quantity_basis', 'active', 'uses'} <= set(listed[0])


def zen_saved(c, **unit):
    d = upload(c, zen_kids_po(), 'zen-1.xlsx')
    r = c.post(f'/api/drafts/{d["id"]}/map', json=dict(version=d['version'], sheet='PO', definition=ZEN, save_template=True,
                                                         template_name='ZEN KIDS PO', **unit))
    assert r.status_code == 200, r.text
    return r.json()


def test_changed_anchor_label_refuses_the_template_and_records_why(env):
    c = env
    zen_saved(c, qty_unit='PCS')
    d = upload(c, zen_kids_po(style='ZK-9', ship_label='ETD'))
    assert 'ZEN KIDS PO' not in d['format'] and c.get('/api/templates').json()[0]['uses'] == 0
    assert any('“ZEN KIDS PO” (version 1) was not used' in n and 'A8 now says “ETD”, not “Ship Date”' in n for n in d['notes'])
    assert d['suggestion'] and d['suggestion']['fields']['style']['value'] == 'ZK-9'      # proposals from the new file only
    assert d['templates_refused'][0]['reasons'] == ['A8 now says “ETD”, not “Ship Date”']


def test_template_whose_cells_no_longer_validate_is_refused(env):
    c = env
    zen_saved(c, qty_unit='PCS')
    d = upload(c, zen_kids_po(ship='TBC'))                       # same layout, but the ship date cell holds text
    assert 'ZEN KIDS PO' not in d['format']
    assert any('B8 holds “TBC”, which is not a date' in n for n in d['notes'])


def test_template_signature_ignores_values_and_restores_unit_and_basis(env):
    c = env
    first = zen_saved(c, qty_unit='PACK', pieces_per_unit=4, quantity_basis='sample')
    t = c.get('/api/templates').json()[0]
    assert t['anchors']['ship_date'] == {'cell': 'B8', 'label_cell': 'A8', 'label': 'Ship Date'}
    assert t['anchors']['table.color_col']['label'] == 'Colour' and t['unit'] == {'qty_unit': 'PACK', 'pieces_per_unit': 4}
    assert first['template']['version'] == 1
    d = upload(c, zen_kids_po(style='ZK-2', po='ZK/PO/2', season='AW-27', lines=(('RED', '', (1, 2, 3, 4)),)))
    o = d['orders'][0]
    assert d['format'] == 'Template: ZEN KIDS PO' and d['template']['name'] == 'ZEN KIDS PO'
    assert (o['qty_unit'], o['pieces_per_unit'], o['quantity_basis']) == ('PACK', 4, 'sample')
    assert o['provenance']['qty_unit']['reason'] == 'template' and o['warnings'][0].startswith('Read with the saved layout “ZEN KIDS PO”')


def test_file_evidence_that_contradicts_the_template_is_asked(env):
    c = env
    zen_saved(c, qty_unit='PACK', pieces_per_unit=4)
    from openpyxl import load_workbook
    from io import BytesIO
    wb = load_workbook(BytesIO(zen_kids_po())); wb.active['A10'] = 'All quantities in pieces'
    d = upload(c, save(wb))
    o = d['orders'][0]
    assert d['format'] == 'Template: ZEN KIDS PO'                      # one new note does not change the layout (similarity ≥ 0.8)
    assert (o['qty_unit'], o['pieces_per_unit']) == ('', None)
    assert o['provenance']['qty_unit']['question'] == ('The file says pieces but the saved layout “ZEN KIDS PO” (version 1) says packs of 4'
                                                       ' — are these quantities pieces or packs?')
    wb = load_workbook(BytesIO(zen_kids_po())); wb.active['G11'] = 'Total Pcs'   # a changed table heading is an anchor:
    d = upload(c, save(wb))                                                          # refused, never applied
    assert 'ZEN KIDS PO' not in d['format'] and any('G11 now says “Total Pcs”, not “Total”' in n for n in d['notes'])


def test_unit_answered_on_review_is_learned_by_the_template(env):
    c = env
    zen_saved(c)                                                    # no unit chosen while mapping
    assert c.get('/api/templates').json()[0]['unit'] is None
    d = upload(c, zen_kids_po(style='ZK-5'))
    assert d['orders'][0]['qty_unit'] == ''
    o = dict(d['orders'][0], qty_unit='PACK', pieces_per_unit=2, quantity_basis='order_demo')
    assert c.put(f'/api/drafts/{d["id"]}', json={'version': d['version'], 'orders': [o]}).status_code == 200
    t = c.get('/api/templates').json()[0]
    assert t['unit'] == {'qty_unit': 'PACK', 'pieces_per_unit': 2} and t['quantity_basis'] == 'order_demo'
    assert any(a['action'] == 'template_unit_learned' for a in c.get('/api/audit').json())
    o = upload(c, zen_kids_po(style='ZK-6'))['orders'][0]
    assert (o['qty_unit'], o['pieces_per_unit'], o['quantity_basis']) == ('PACK', 2, 'order_demo')


# --- §7 dates in templates and header fields ------------------------------------------------------
@pytest.mark.parametrize('raw,iso,warned', [('20/12/2026', '2026-12-20', False), (46402, '2027-01-15', False),
                                            ('5 Mar 2027', '2027-03-05', False), ('03/04/2026', '2026-04-03', True)])
def test_template_dates_become_iso(raw, iso, warned):
    sheet = sheet_of(zen_kids_po(ship=raw))
    o = workbooks.finish(workbooks.apply_template([sheet], sheet, ZEN)[0])
    assert o['ship_date'] == iso and o['provenance']['ship_date']['cell'] == 'B8'
    assert any('day/month' in w for w in o['warnings']) == warned


def test_flat_table_with_style_in_the_header_is_read_automatically():
    d = workbooks.parse(wrong_mapping_po())
    o = d['orders'][0]
    assert d['format'] == 'Tabular orders' and o['style'] == 'AC-77' and o['provenance']['style']['cell'] == 'B2'
    assert (o['buyer'], o['po_number'], o['ship_date'], o['season']) == ('ACME KIDS', 'AC/PO/31', '2027-02-01', 'AW-27')
    assert [(l['color'], [(s['size'], s['quantity']) for s in l['sizes']]) for l in o['lines']] == [('NAVY', [('S', 120), ('M', 140)]), ('GREY', [('S', 110)])]


def test_number_fields_and_order_date_are_read_from_header_labels():
    wb = Workbook(); ws = wb.active
    ws['A1'], ws['B1'] = 'Style', 'Q-1'
    ws['A2'], ws['B2'] = 'Order Date', datetime(2026, 10, 1)
    ws['A3'], ws['B3'] = 'Components', '3'
    ws.append([]); ws.append(['Size', 'Qty'])
    ws.append(['M', 5])
    o = workbooks.parse(save(wb))['orders'][0]
    assert (o['style'], o['order_date'], o['components']) == ('Q-1', '2026-10-01', 3)


def test_numeric_size_headings_named_by_a_colour_heading_are_a_matrix_and_refs_are_normalised(monkeypatch):
    wb = Workbook(); ws = wb.active
    ws['A1'], ws['B1'] = 'Style No', 'J-28'
    ws.append([]); ws.append(['Colour', 28, 30, 32, 'Total']); ws.append(['INDIGO', 5, 6, 7, 18]); ws.append(['BLACK', 1, 2, 3, 6])
    sheet = sheet_of(save(wb))
    s = matcher.suggest(sheet)
    t = s['definition']['table']
    assert (t['type'], t['header_row'], t['first_size_col'], t['last_size_col'], t['total_col']) == ('matrix', 3, 2, 4, 5)
    assert s['table']['confidence'] == 'medium'                         # numeric sizes: shown, checked by a person
    ok, rejected = matcher.check(sheet, {'fields': {'style': 'b1'}, 'table': t})
    assert ok['fields'] == {'style': 'B1'} and not rejected
    monkeypatch.setattr(ai, 'available', lambda: True)
    monkeypatch.setattr(ai, '_chat', lambda *a, **k: json.dumps({'fields': ['B1'], 'table': 'flat'}))
    assert ai.suggest(sheet)['definition']['table']['header_row'] == 3   # a malformed AI answer is ignored, not fatal


def test_matcher_is_fast_on_a_large_sheet():
    import time
    wb = Workbook(); ws = wb.active
    for r in range(1, 401):
        ws.append([f'Note {r}', f'TEXT {r % 7}', r, f'ITEM-{r}', 'RED', 'BLUE', r * 2, 'remark'] + [f'X{c}' for c in range(12)])
    sheet = sheet_of(save(wb))
    start = time.perf_counter(); matcher.suggest(sheet)
    assert time.perf_counter() - start < 3


def test_master_names_for_matching_survive_one_oversized_master(monkeypatch):
    """§3 rule 4: the matcher reads only the buyer, colour and size masters, each on its own, so one master
    too large to list never hides the others (and the style / party masters are never loaded for it)."""
    from backend import oracle_store
    asked = []

    class Catalog:
        def __init__(self, conn): pass
        def columns(self, tables): return {t: {} for t in tables}
        def masters(self, spec):
            asked.append(spec['table'])
            if spec['table'] == 'COLORMAS':
                raise ValueError('COLORMAS: master catalogue exceeds 5000 rows; configure a scoped lookup.')
            return [dict(id=1, label=f'{spec["table"]} NAME', cancelled=False), dict(id=2, label='OLD', cancelled=True)]

    class Conn:
        def __enter__(self): return self
        def __exit__(self, *a): return False
    monkeypatch.setattr(oracle_store, 'OracleCatalog', Catalog)
    monkeypatch.setattr(oracle_store, 'connect', lambda: Conn())
    monkeypatch.setattr(oracle_store, 'current_mapping', lambda: mapping.effective())
    assert main.master_labels() == {'buyer': ['BUYMAS NAME'], 'color': [], 'size': ['SIZE_MAS NAME']}
    assert sorted(asked) == ['BUYMAS', 'COLORMAS', 'SIZE_MAS']
