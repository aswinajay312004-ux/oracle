"""Zero silent data loss and truthful save outcomes (docs/EXCEL_DATA_FIDELITY_SPEC.md, docs/TRANSACTION_AND_RECONCILIATION.md):
the original file is kept byte for byte, every sheet and number of it is accounted for, values Oracle will not receive
are listed and must be read before approving, and a save whose outcome is not known is never called saved or failed."""
from copy import deepcopy
import hashlib
import os
import stat
from io import BytesIO
import oracledb
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from backend import coverage, local_store, main, oracle_store
from fake_oracle import FakeOracle
from fixtures import merch_bom, flat_qty_orders, save

PW = 'test-password-123'


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(local_store, 'DATA', tmp_path); monkeypatch.setattr(main, 'DATA', tmp_path)
    monkeypatch.setattr(main, 'RECHECK_S', 0)
    oracle = FakeOracle(masters={'BUYMAS': [dict(id=7, label='TESTBUYER', cancelled=False)],
                                 'UNITMAS': [dict(id=8, label='PCS', cancelled=False)]}).install(monkeypatch)
    with TestClient(main.app) as c:
        assert c.post('/api/setup', json={'username': 'admin', 'password': PW}).status_code == 200
        yield c, oracle, tmp_path


def upload(c, content, name='orders.xlsx'):
    r = c.post('/api/uploads', files={'file': (name, content)}); assert r.status_code == 200, r.text
    return r.json()


def ready(c, d, ref='TEST/FID/1', **changes):
    o = deepcopy(d['orders'][0])
    o.update(order_ref=ref, order_date='2026-10-17', ship_date='2027-01-15', po_number='PO-1', buyer='TESTBUYER',
             description='TEST TEE', quantity_basis='sample', qty_unit='PCS', pieces_per_unit=1)
    o.update(changes)
    r = c.put(f'/api/drafts/{d["id"]}', json={'version': d['version'], 'orders': [o]}); assert r.status_code == 200, r.text
    r = c.post(f'/api/drafts/{d["id"]}/preview', json={'version': r.json()['version']}); assert r.status_code == 200, r.text
    p = r.json(); assert not p['plan']['errors'], p['plan']['errors']
    return p


def approval(p, seen=True):
    return {'version': p['version'], 'plan_hash': p['plan']['hash'],
            'exceptions_seen': (p['plan'].get('exceptions') or {}).get('digest') if seen else None}


def self_approval(c):
    pass                       # one access level: the person who uploads may save (no four-eyes rule any more)


# --- Layer A: the original file ----------------------------------------------------------------------------
def test_the_original_file_is_kept_byte_for_byte_read_only_and_served_only_while_its_hash_matches(env):
    c, _, data = env
    content = flat_qty_orders()
    d = upload(c, content)
    stored = data / (d['id'] + '.upload')
    assert stored.read_bytes() == content and d['filehash'] == hashlib.sha256(content).hexdigest() and d['size'] == len(content)
    assert not os.stat(stored).st_mode & (stat.S_IWUSR | stat.S_IWGRP | stat.S_IWOTH), 'the archived original is read-only'
    r = c.get(f'/api/drafts/{d["id"]}/original')
    assert r.status_code == 200 and r.content == content and r.headers['X-Content-SHA256'] == d['filehash']
    os.chmod(stored, 0o644); stored.write_bytes(content + b'tampered')
    assert c.get(f'/api/drafts/{d["id"]}/original').status_code == 409, 'a changed original is never served as the original'


def test_a_file_that_cannot_be_read_is_still_kept(env, monkeypatch):
    c, _, data = env
    def broken(*a, **k):
        raise ValueError('Multiple order blocks need explicit mapping.')
    monkeypatch.setattr(main.workbooks, 'parse', broken)
    content = flat_qty_orders()
    r = c.post('/api/uploads', files={'file': ('broken.xlsx', content)})
    assert r.status_code == 409 and 'original file is kept unchanged' in r.json()['detail']
    kept = [p for p in data.glob('*.upload')]
    assert [p.read_bytes() for p in kept] == [content]
    assert any(e['action'] == 'upload_unreadable' for e in c.get('/api/audit').json())


# --- Layer B coverage --------------------------------------------------------------------------------------
def two_sheets_one_unknown():
    wb = Workbook(); ws = wb.active; ws.title = 'Orders'
    ws.append(['Style', 'Colour', 'Size', 'Qty'])
    for row in [('FL-1', 'RED', 'S', 10), ('FL-1', 'RED', 'M', 20)]:
        ws.append(row)
    other = wb.create_sheet('Extra PO')
    other['A1'], other['B1'] = 'Article', 'ZK-9'
    for c, h in enumerate(['Shade', 'S', 'M', 'L'], 1):
        other.cell(3, c, h)
    other.append(['NAVY', 5, 6, 7])
    hidden = wb.create_sheet('Old'); hidden.sheet_state = 'hidden'; hidden.append(['Qty', 99])
    return save(wb)


def test_every_sheet_and_every_number_is_accounted_for(env):
    c, _, _ = env
    d = upload(c, two_sheets_one_unknown())
    cov = d['coverage']
    status = {s['name']: s['status'] for s in cov['sheets']}
    assert status == {'Orders': 'read', 'Extra PO': 'not read', 'Old': 'hidden, not read'}
    assert cov['totals']['numbers'] == cov['totals']['numbers_read'] + cov['totals']['numbers_not_read'] == 6
    assert cov['totals']['numbers_read'] == 2
    assert any('“Extra PO” was not read: 3 numbers' in e for e in cov['exclusions'])
    assert any('“Old” (hidden in Excel) was not read' in e for e in cov['exclusions'])


def test_a_second_size_table_on_the_same_sheet_is_named_not_dropped():
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    for r, row in enumerate([['Shade', 'S', 'M', 'L'], ['NAVY', 10, 20, 30], [], [], ['Shade', 'S', 'M', 'L'], ['RED', 1, 2, 3], ['PINK', 4, 5, 6]], 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    from backend import workbooks
    sheets = workbooks.inspect(save(wb), 'x.xlsx')
    orders = workbooks.apply_template(sheets, sheets[0], {'fields': {}, 'table': {'type': 'matrix', 'header_row': 1, 'color_col': 1,
                                                                                  'first_size_col': 2, 'last_size_col': 4}})
    cov = coverage.manifest(sheets, orders)
    assert cov['exclusions'] == ['Sheet “PO” rows 6–7, columns B–D: 6 numbers not read (under S, M, L at row 5).']
    assert cov['totals']['numbers_checked'] == 0          # 3 = 1 + 2 in a size row is not taken for a total


def test_totals_repeats_and_zero_tables_are_checked_not_reported_as_lost():
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    rows = [['Colour', 'S', 'M', 'Total'], ['NAVY', 10, 20, 30], ['RED', 5, 5, 10], ['Total', 15, 25, 40], [],
            ['Colour', 'S', 'M', 'Total'], ['NAVY', 10, 20, 30], ['RED', 5, 5, 10], [],
            ['Colour', 'S', 'M', 'Total'], ['NAVY', 0, 0, 0], [],
            ['Colour', 'S', 'M', 'Total'], ['GREEN', 7, 8, 15]]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    from backend import workbooks
    sheets = workbooks.inspect(save(wb), 'x.xlsx')
    orders = workbooks.apply_template(sheets, sheets[0], {'fields': {'style': 'A2'}, 'table': {
        'type': 'matrix', 'header_row': 1, 'color_col': 1, 'first_size_col': 2, 'last_size_col': 3, 'last_row': 3}})
    cov = coverage.manifest(sheets, orders)
    text = ' | '.join(cov['checked'])
    assert 'totals match' in text and 'rows 7–8, columns B–C: 4 numbers repeat rows 2–3 exactly' in text and 'all 0' in text
    # a table with other colours and numbers is still reported, whatever its totals
    assert cov['exclusions'] == ['Sheet “PO” row 14, columns B–C: 2 numbers not read (under S, M at row 13).']
    t = cov['totals']
    assert t['numbers'] == t['numbers_read'] + t['numbers_checked'] + t['numbers_not_read']


def test_a_matrix_longer_than_the_row_limit_says_so(monkeypatch):
    from backend import workbooks
    monkeypatch.setattr(workbooks, 'MATRIX_ROWS', 3)
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    ws.append(['Shade', 'S', 'M'])
    for i in range(6):
        ws.append([f'C{i}', i + 1, i + 2])
    sheets = workbooks.inspect(save(wb), 'x.xlsx')
    [o] = workbooks.apply_template(sheets, sheets[0], {'fields': {'style': 'A2'}, 'table': {'type': 'matrix', 'header_row': 1, 'color_col': 1,
                                                                                            'first_size_col': 2, 'last_size_col': 3}})
    assert len(o['lines']) == 3 and o['warnings'][0].startswith('Only the first 3 rows under the size headings')
    assert any('rows 5–7' in e for e in coverage.manifest(sheets, [o])['exclusions'])


def test_mapping_one_sheet_keeps_the_orders_read_from_the_other_sheets(env):
    c, _, _ = env
    d = upload(c, two_sheets_one_unknown())
    assert [o['sheet'] for o in d['orders']] == ['Orders']
    definition = {'fields': {'style': 'B1'}, 'table': {'type': 'matrix', 'header_row': 3, 'color_col': 1, 'first_size_col': 2, 'last_size_col': 4}}
    r = c.post(f'/api/drafts/{d["id"]}/map', json=dict(version=d['version'], sheet='Extra PO', definition=definition, qty_unit='PCS',
                                                        quantity_basis='sample'))
    assert r.status_code == 200, r.text
    d = r.json()
    assert [(o['sheet'], o['style']) for o in d['orders']] == [('Orders', 'FL-1'), ('Extra PO', 'ZK-9')]
    assert [s['status'] for s in d['coverage']['sheets']] == ['read', 'read', 'hidden, not read']


# --- what will not reach Oracle ----------------------------------------------------------------------------
def test_values_oracle_will_not_receive_are_listed_and_must_be_read_before_approving(env):
    c, oracle, _ = env
    self_approval(c)
    d = upload(c, flat_qty_orders())
    p = ready(c, d, season='SS-27', department='KIDS')
    ex = p['plan']['exceptions']
    got = {(x['label'], x['kept']) for x in ex['not_saved']}
    assert {('Season', 'poppys'), ('Department', 'unverified')} <= got, got
    assert ex['digest'] and not ex['exclusions']
    r = c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p, seen=False))
    assert r.status_code == 409 and r.json()['exceptions_unseen'] and not oracle.receipts, 'nothing is saved before the list was read'
    r = c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p))
    assert r.status_code == 200, r.text
    receipt = r.json()
    assert receipt['complete'] is False and {x['label'] for x in receipt['not_saved']} >= {'Season', 'Department'}
    assert c.get(f'/api/history/{d["id"]}').json()['not_saved'] == receipt['not_saved'], 'the Oracle receipt keeps the list'


def test_a_pack_conversion_is_shown_with_the_original_figures(env):
    c, _, _ = env
    p = ready(c, upload(c, flat_qty_orders()), qty_unit='PACK', pieces_per_unit=3)
    [conv] = p['plan']['conversions']
    assert (conv['source_quantity'], conv['pieces_per_unit'], conv['pieces']) == (2500, 3, 7500) and p['plan']['totals']['pieces'] == 7500
    assert conv['text'] == '2,500 packs × 3 pieces per pack = 7,500 pieces (entered on the review screen).'


# --- one access level (owner, 2026-10-21: one administrator login shared by everyone, no four-eyes rule) ----------
def test_whoever_uploads_and_changes_an_order_can_save_it(env):
    c, _, _ = env
    d = upload(c, flat_qty_orders())
    p = ready(c, d)
    r = c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p))
    assert r.status_code == 200, r.text


# --- save outcomes ---------------------------------------------------------------------------------------------
def lost_connection():
    return oracledb.OperationalError('DPY-4011: the database or network closed the connection')


def test_a_save_whose_outcome_is_unknown_is_neither_saved_nor_failed_until_oracle_answers(env, monkeypatch):
    c, oracle, _ = env
    self_approval(c)
    d = upload(c, flat_qty_orders()); p = ready(c, d)
    real_commit = oracle.commit

    def commit_then_lose(draft, *a, **k):
        real_commit(draft, *a, **k)                         # Oracle commits ...
        monkeypatch.setattr(oracle_store, 'connect', lambda: (_ for _ in ()).throw(lost_connection()))
        raise lost_connection()                             # ... and the connection dies before the answer arrives
    monkeypatch.setattr(oracle_store, 'commit', commit_then_lose)
    r = c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p))
    assert r.status_code == 503 and r.json()['outcome'] == 'unknown', r.text
    assert c.put(f'/api/drafts/{d["id"]}', json={'version': p['version'], 'orders': p['orders']}).status_code == 409, \
        'a draft whose save is unknown cannot be changed'
    assert c.post(f'/api/drafts/{d["id"]}/reconcile').status_code == 503
    monkeypatch.setattr(oracle_store, 'connect', lambda: __import__('fake_oracle').FakeConn())   # Oracle is back
    r = c.post(f'/api/drafts/{d["id"]}/reconcile')
    assert r.status_code == 200 and r.json()['result']['orders'][0]['order_ref'] == 'TEST/FID/1' and 'outcome' not in r.json()
    assert len(oracle.receipts) == 1, 'saved exactly once'
    actions = [e['action'] for e in c.get('/api/audit').json()]
    assert {'approve_started', 'approve_outcome_unknown', 'reconciled_saved'} <= set(actions)


def test_an_error_after_oracle_committed_is_reported_as_saved(env, monkeypatch):
    c, oracle, _ = env
    self_approval(c)
    d = upload(c, flat_qty_orders()); p = ready(c, d)
    real_commit = oracle.commit

    def commit_then_error(draft, *a, **k):
        real_commit(draft, *a, **k)
        raise lost_connection()
    monkeypatch.setattr(oracle_store, 'commit', commit_then_error)
    r = c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p))
    assert r.status_code == 200 and r.json()['already_committed'] is True
    assert c.get(f'/api/drafts/{d["id"]}').json()['result']


def test_a_refused_save_is_reported_as_not_saved_only_because_oracle_confirms_it(env, monkeypatch):
    c, oracle, _ = env
    self_approval(c)
    d = upload(c, flat_qty_orders()); p = ready(c, d)
    monkeypatch.setattr(oracle_store, 'commit', lambda *a, **k: (_ for _ in ()).throw(lost_connection()))
    r = c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p))
    assert r.status_code == 503 and r.json()['outcome'] == 'not_saved' and not oracle.receipts
    assert 'outcome' not in c.get(f'/api/drafts/{d["id"]}').json(), 'the draft can be approved again'


def test_a_local_failure_after_the_oracle_commit_still_says_saved_and_heals(env, monkeypatch):
    c, oracle, _ = env
    self_approval(c)
    d = upload(c, flat_qty_orders()); p = ready(c, d)
    real_save = local_store.save

    def disk_full(draft, actor, expected=None):
        if draft.get('result'):
            raise OSError('No space left on device')
        return real_save(draft, actor, expected)
    monkeypatch.setattr(local_store, 'save', disk_full)
    r = c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p))
    assert r.status_code == 200 and 'saved in Oracle' in r.json()['local_record'] and oracle.receipts
    monkeypatch.setattr(local_store, 'save', real_save)
    healed = c.get(f'/api/drafts/{d["id"]}').json()
    assert healed['result'] and 'outcome' not in healed, 'the draft is updated from Oracle when opened again'


def test_receipt_extras_say_complete_only_when_nothing_was_left_out():
    assert oracle_store.receipt_extras({'not_saved': [], 'hash': 'h'}, {'coverage': {'exclusions': []}})['complete'] is True
    assert oracle_store.receipt_extras({'not_saved': [{'label': 'Season'}]}, {})['complete'] is False
    assert oracle_store.receipt_extras({}, {'coverage': {'exclusions': ['Sheet x was not read']}})['complete'] is False
    assert oracle_store.receipt_extras({'reference_only': [{'table': 'ORDFAB'}]}, {})['complete'] is False


def test_merch_bom_lists_its_unsaved_summary_block_and_bom_rows(env):
    c, _, _ = env
    d = upload(c, merch_bom())
    # The summary block repeats the order block number for number: checked and listed, not reported as lost.
    assert any('rows 20–21' in e and 'repeat rows 12–13 exactly' in e for e in d['coverage']['checked']), d['coverage']
    assert not any('rows 20–21' in e for e in d['coverage']['exclusions'])
    _ = BytesIO


def test_verified_only_after_an_independent_read_back_with_control_totals(env):
    c, oracle, _ = env
    self_approval(c)
    d = upload(c, flat_qty_orders()); p = ready(c, d)
    assert c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p)).status_code == 200
    assert c.get('/api/drafts').json()[0]['status'] == 'Saved, not verified yet'
    g = c.get(f'/api/history/{d["id"]}/graph').json()
    assert g['verified'] and {x['check']: (x['expected'], x['found']) for x in g['controls']} == {
        'Pieces (sum of the size quantities)': (2500, 2500), 'Orders saved': (1, 1), 'Lines saved': (2, 2)}
    assert c.get('/api/drafts').json()[0]['status'] == 'Verified, with exceptions'   # the "not saved" list of this file
    size = next(t for t in oracle.rows if t == 'ORDSIZE')
    first = next(iter(oracle.rows[size])); oracle.rows[size][first]['RQTY'] += 1          # someone changed Oracle afterwards
    g = c.get(f'/api/history/{d["id"]}/graph').json()
    assert not g['verified'] and any(not x['ok'] for x in g['controls'])
    assert c.get('/api/drafts').json()[0]['status'] == 'Saved, differences found'
    assert any(e['action'] == 'verification_failed' for e in c.get('/api/audit').json())


def test_the_same_business_order_from_another_workbook_is_flagged_before_approval(env):
    c, _, _ = env
    self_approval(c)
    d = upload(c, flat_qty_orders()); p = ready(c, d, ref='TEST/DUP/1')
    assert c.post(f'/api/drafts/{d["id"]}/accept', json=approval(p)).status_code == 200
    again = ready(c, upload(c, flat_qty_orders(qty_header='Quantity'), 'resent.xlsx'), ref='TEST/DUP/2', buyer='testbuyer ', po_number='PO-1')
    assert any('were already saved as TEST/DUP/1' in w for w in again['plan']['warnings']), again['plan']['warnings']


# --- edge cases of real files (audit COMPLETE_EXCEL_PIPELINE_AUDIT B6, B18, B22) -----------------------------------
def matrix_sheet(rows, merged=()):
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    ws['A1'], ws['B1'] = 'Style No', 'MX-1'
    for c, h in enumerate(['Colour', 'S', 'M', 'L'], 1):
        ws.cell(3, c, h)
    for r, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            if v is not None:
                ws.cell(r, c, v)
    for rng in merged:
        ws.merge_cells(rng)
    from backend import workbooks
    sheets = workbooks.inspect(save(wb), 'x.xlsx')
    d = {'fields': {'style': 'B1'}, 'table': {'type': 'matrix', 'header_row': 3, 'color_col': 1, 'first_size_col': 2, 'last_size_col': 4}}
    return sheets, workbooks.apply_template(sheets, sheets[0], d)[0]


def test_a_merged_colour_cell_belongs_to_every_row_it_covers():
    _, o = matrix_sheet([('NAVY', 1, 2, 3), (None, 4, 5, 6)], merged=['A4:A5'])
    assert [l['color'] for l in o['lines']] == ['NAVY', 'NAVY'] and o['lines'][1]['provenance']['color']['cell'] == 'A4'


def test_a_blank_size_is_not_ordered_and_says_so_but_an_error_cell_is_never_a_quantity():
    from backend.erp import catalog, mapping, planner
    _, o = matrix_sheet([('NAVY', 10, None, 30), ('RED', 5, '#N/A', 7)])
    o.update(buyer='B', description='D', po_number='P', order_ref='R/BL', order_date='2026-10-17', ship_date='2026-11-01',
             quantity_basis='sample', qty_unit='PCS', pieces_per_unit=1, sample_type='DEVELOPMENT SAMPLE')
    p = planner.plan(catalog.SnapshotCatalog(), [o], 'h' * 64, mapping.effective())
    assert any('size M is blank in the file and treated as not ordered' in w for w in p['warnings'])
    assert [op['values']['RQTY'] for op in p['operations'] if op['table'] == 'ORDSIZE'][:2] == [10, 30], 'no size row for the blank M'
    assert any(i['path'] == 'orders.0.lines.1.sizes.1.quantity' and i['severity'] == 'error' for i in p['issues']), \
        'an Excel error in a quantity cell blocks the save on that cell'


def test_an_excel_error_in_a_header_field_is_asked_never_used_as_the_value():
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    ws['A1'], ws['B1'] = 'Buyer', '#N/A'
    ws['A2'], ws['B2'] = 'Style No', 'MX-2'
    for c, h in enumerate(['Colour', 'S'], 1):
        ws.cell(4, c, h)
    ws.append(['NAVY', 3])
    from backend import workbooks
    sheets = workbooks.inspect(save(wb), 'x.xlsx')
    [o] = workbooks.apply_template(sheets, sheets[0], {'fields': {'buyer': 'B1', 'style': 'B2'}, 'table': {
        'type': 'matrix', 'header_row': 4, 'color_col': 1, 'first_size_col': 2, 'last_size_col': 2}})
    assert o['buyer'] == '' and 'Excel error #N/A' in o['provenance']['buyer']['question']


def test_a_size_matrix_with_a_style_column_becomes_one_order_per_style():
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    for c, h in enumerate(['Style', 'Colour', 'S', 'M'], 1):
        ws.cell(1, c, h)
    for row in [('ST-1', 'NAVY', 1, 2), ('ST-1', 'RED', 3, 4), ('ST-2', 'NAVY', 5, 6)]:
        ws.append(row)
    from backend import workbooks
    sheets = workbooks.inspect(save(wb), 'x.xlsx')
    orders = workbooks.apply_template(sheets, sheets[0], {'fields': {}, 'table': {'type': 'matrix', 'header_row': 1, 'label_col': 1, 'color_col': 2,
                                                                                    'first_size_col': 3, 'last_size_col': 4}})
    assert [(o['style'], [l['color'] for l in o['lines']], o['provenance']['style']['cell']) for o in orders] == [
        ('ST-1', ['NAVY', 'RED'], 'A2'), ('ST-2', ['NAVY'], 'A4')]
    assert coverage.manifest(sheets, orders)['totals']['numbers_not_read'] == 0


def test_new_master_records_are_listed_for_the_approver(env):
    c, _, _ = env
    p = ready(c, upload(c, flat_qty_orders()), ref='TEST/NM/1')
    made = {(x['business'], x['label']) for x in p['plan']['exceptions']['new_masters']}
    assert ('Style number', 'FL-500') in made and ('Colour / combo', 'RED') in made
    assert not any(x['label'] == 'TESTBUYER' for x in p['plan']['exceptions']['new_masters']), 'an existing buyer is reused, not listed'


def test_unread_numbers_name_the_excel_block_they_are_in():
    """A trims block's calculation columns are reported under the block's title ("SEWING THREAD"), also for a following
    block that repeats its heading without a title of its own."""
    from backend import coverage
    rows = {(1, 2): 'SEWING THREAD', (1, 3): 'OCN 82', (2, 2): 'Garment', (2, 3): 'METER', (3, 2): 'G1', (3, 3): 5000,
            (4, 2): 'G2', (4, 3): 5000, (7, 2): 'Garment', (7, 3): 'METER', (8, 2): 'G3', (8, 3): 4000}
    cells = [dict(row=r, col=c, value=v, cell=f'{"ABCD"[c - 1]}{r}') for (r, c), v in rows.items()]
    by_rc = dict(rows)
    titles = {}
    first = coverage._describe({'name': 'S'}, dict(first=3, last=4, cols=[3], cells=[x for x in cells if x['row'] in (3, 4) and x['col'] == 3]), by_rc, titles)
    second = coverage._describe({'name': 'S'}, dict(first=8, last=8, cols=[3], cells=[x for x in cells if x['row'] == 8 and x['col'] == 3]), by_rc, titles)
    assert first['section'] == second['section'] == 'SEWING THREAD'
    assert 'not read in “SEWING THREAD” (under METER at row 2)' in first['text']
