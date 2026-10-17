"""Every accepted upload format yields the same order. Files written by other programs (Excel via the
calamine samples, LibreOffice, msoffcrypto-tool) are checked against our own fixtures so readers are not
only tested against themselves. Hostile and damaged inputs must end in a plain message, never a 500.

Most tests parse in-process (POPPYS_INLINE_PARSE=1) for speed; tests named test_isolated_* go through
the spawned worker process that production uses."""
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import json
import quopri
import struct
import sys
import time
import zipfile
import pytest
from backend import isolation, numfmt, readers, workbooks
from fixtures import merch_bom, flat_orders, as_xls, as_xml2003, as_html, as_csv, as_json_records

FORMATS = Path(__file__).parent / 'fixtures' / 'formats'


@pytest.fixture(autouse=True)
def inline_parse(request, monkeypatch):
    if request.node.name.startswith('test_isolated_'):
        monkeypatch.delenv(readers.INLINE_ENV, raising=False)
    else:
        monkeypatch.setenv(readers.INLINE_ENV, '1')


def _ws(v):
    return ' '.join(v.split()) if isinstance(v, str) else v


def summary(parsed, cells=True):
    return [(o['buyer'], o['style'], o['description'], o['po_number'], o['ship_date'], o['merch'], o['season'],
             [(l['label'], l['color'], [(x['size'], x['quantity']) + ((x['source']['cell'],) if cells else ()) for x in l['sizes']])
              for l in o['order_lines'] or o['lines']],
             [(_ws(f['component']), f['total']) for f in o['fabrics']], len(o['accessories']), (o.get('packs') or {}).get('pieces'))
            for o in parsed['orders']]


def test_calamine_reference_workbook_reads_identically_in_four_formats():
    def values(path):
        _, sheets = readers.read(path.read_bytes(), path.name)
        return {s['name']: sorted((c['row'], c['col'], c['value']) for c in s['cells']) for s in sheets if s['name'] in ('datatypes', 'issue2', 'issue5')}
    base = values(FORMATS / 'calamine_issues.xlsx')
    assert base['issue2'] == [(1, 1, 1), (1, 2, 'a'), (2, 1, 2), (2, 2, 'b'), (3, 1, 3), (3, 2, 'c')]
    for ext in ('xls', 'xlsb', 'ods'):
        assert values(FORMATS / f'calamine_issues.{ext}') == base, ext


@pytest.mark.parametrize('name,data,cells', [
    ('bom.xls', lambda: as_xls(merch_bom()), True),
    ('bom_lo.xls', lambda: (FORMATS / 'merch_bom_libreoffice.xls').read_bytes(), True),
    ('bom_lo.ods', lambda: (FORMATS / 'merch_bom_libreoffice.ods').read_bytes(), True),
    ('bom.xml', lambda: as_xml2003(merch_bom()), True),
    ('bom_export.xls', lambda: as_html(merch_bom()), True),
    ('bom_lo.html', lambda: (FORMATS / 'merch_bom_libreoffice.html').read_bytes(), False),
])
def test_merch_bom_gives_the_same_order_in_every_format(name, data, cells):
    expected = summary(workbooks.parse(merch_bom()), cells)
    got = workbooks.parse(data(), filename=name)
    assert got['format'] == 'Merch BOM'
    assert summary(got, cells) == expected


@pytest.mark.parametrize('name,data', [
    ('orders.csv', lambda: as_csv(flat_orders())),
    ('orders_semicolon.csv', lambda: as_csv(flat_orders(), ';')),
    ('orders.txt', lambda: as_csv(flat_orders(), '\t', 'utf-16')),          # Excel "Unicode text"
    ('orders.tsv', lambda: as_csv(flat_orders(), '\t', 'utf-8-sig')),
    ('orders.json', lambda: as_json_records(flat_orders())),
    ('orders.xls', lambda: as_html(flat_orders())),
])
def test_flat_orders_give_the_same_lines_in_every_text_format(name, data):
    expected = summary(workbooks.parse(flat_orders()), cells=False)
    got = workbooks.parse(data(), filename=name)
    assert summary(got, cells=False) == expected
    assert got['orders'][0]['style'] == '0008'          # leading-zero style code survives text formats


def test_json_document_and_sheet_map_shapes():
    doc = {'Buyer': 'ZEN KIDS', 'Style No': 'ZK-1', 'PO No': 'ZK/PO/1',
           'lines': [{'Style No': 'ZK-1', 'Colour': 'NAVY', 'Size': 'S', 'Qty': 10}, {'Style No': 'ZK-1', 'Colour': 'NAVY', 'Size': 'M', 'Qty': 12}]}
    o = workbooks.parse(json.dumps(doc).encode(), filename='po.json')['orders'][0]
    assert (o['buyer'], o['style'], o['po_number']) == ('ZEN KIDS', 'ZK-1', 'ZK/PO/1')
    assert [(s['size'], s['quantity']) for s in o['lines'][0]['sizes']] == [('S', 10), ('M', 12)]
    sheets = {'Orders': [['Style', 'Size', 'Qty'], ['A1', 'S', 5]], 'Notes': [{'note': 'x'}]}
    _, out = readers.read(json.dumps(sheets).encode(), 'm.json')
    assert [s['name'] for s in out] == ['Orders', 'Notes']
    own = {'sheets': [{'name': 'PO', 'cells': [{'cell': 'B2', 'value': 'NEXT'}, {'row': 3, 'col': 1, 'value': 7}]}]}
    _, out = readers.read(json.dumps(own).encode(), 'own.json')
    assert {c['cell']: c['value'] for c in out[0]['cells']} == {'B2': 'NEXT', 'A3': 7}


@pytest.mark.parametrize('name,content,reason', [
    ('empty.xlsx', b'', 'empty'),
    ('scan.pdf.xlsx', b'%PDF-1.7 binary\x00\x01', 'not an Excel workbook'),
    ('random.xls', bytes(range(256)) * 4, 'not an Excel workbook'),
    ('notes.xml', b'<?xml version="1.0"?><note>hi</note>', 'no spreadsheet table'),
    ('archive.xlsx', b'PK\x03\x04garbage', 'damaged'),
])
def test_non_spreadsheets_are_refused_with_a_reason(name, content, reason):
    with pytest.raises(readers.Unsupported, match=reason):
        readers.read(content, name)


def test_hostile_content_is_data_only():
    bomb = ('<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]>'
            '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet><Table><Row><Cell><Data>&b;</Data></Cell></Row></Table></Worksheet></Workbook>').encode()
    with pytest.raises(ValueError):
        readers.read(bomb, 'bomb.xml')                          # entity expansion refused (defusedxml)
    deep = ('[' * 40 + ']' * 40).encode()
    with pytest.raises(ValueError, match='nested too deeply'):
        readers.read(deep, 'deep.json')
    _, sheets = readers.read(b'Style,Size,Qty\r\n=cmd|calc,S,=1+1\r\n', 'inject.csv')
    values = {c['cell']: c['value'] for c in sheets[0]['cells']}
    assert values['A2'] == '=cmd|calc' and values['C2'] == '=1+1' and all(c['formula'] is None for c in sheets[0]['cells'])
    html = b'<table><tr><td>Buyer</td><td>NEXT<script>steal()</script></td></tr></table>'
    _, sheets = readers.read(html, 'x.xls')
    assert {c['cell']: c['value'] for c in sheets[0]['cells']} == {'A1': 'Buyer', 'B1': 'NEXT'}


def test_html_spans_become_merged_cells_like_excel():
    html = b'<table><tr><td colspan="2">TITLE</td><td>X</td></tr><tr><td rowspan="2">A</td><td>B</td><td>C</td></tr><tr><td>D</td><td>E</td></tr></table>'
    _, sheets = readers.read(html, 'span.xls')
    s = sheets[0]
    assert {c['cell']: c['value'] for c in s['cells']} == {'A1': 'TITLE', 'C1': 'X', 'A2': 'A', 'B2': 'B', 'C2': 'C', 'B3': 'D', 'C3': 'E'}
    assert s['merged'] == ['A1:B1', 'A2:A3']


def test_buyer_dates_become_iso_and_ambiguity_is_reported():
    cases = {'2027-01-15': ('2027-01-15', None), '2027-01-15 0:00:00': ('2027-01-15', None), '20/12/2026': ('2026-12-20', None),
             '20-Dec-26': ('2026-12-20', None), '5 Mar 2027': ('2027-03-05', None), '2026/12/20': ('2026-12-20', None),
             '12/20/2026': ('2026-12-20', '"12/20/2026" read as month/day/year'), 46402: ('2027-01-15', 'Excel date number 46402 read as a date'),
             'ASAP': (None, None), '31/02/2026': (None, None)}
    for raw, expected in cases.items():
        assert workbooks.iso_date(raw) == expected, raw
    assert workbooks.iso_date('03/04/2026') == ('2026-04-03', '"03/04/2026" read as day/month/year')


def test_typed_text_values_follow_excel_display_rules():
    assert [readers.typed(v) for v in ('1,500', '0008', '12', '-3', '2.50', '0.5', ' 7 ', 'N/A', '')] == \
        [1500, '0008', 12, -3, 2.5, 0.5, 7, 'N/A', None]


# --- helpers: small synthetic files in formats no library here writes ---------------------------------
ENDOFCHAIN, FREESECT, FATSECT, NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF


def cfb(tree):
    """A minimal valid OLE2 compound file (version 3, 512-byte sectors) holding `tree`:
    {name: bytes (a stream) or dict (a storage)}. Streams are padded to 4096 bytes (no mini stream)."""
    entries = [None]

    def add(children):
        sids = []
        for name, value in children.items():
            sid = len(entries); entries.append(None); sids.append(sid)
            if isinstance(value, dict):
                child = add(value)
                entries[sid] = dict(name=name, type=1, child=child, start=ENDOFCHAIN, size=0)
            else:
                entries[sid] = dict(name=name, type=2, child=NOSTREAM, data=value.ljust(4096, b'\0'))
        for a, b in zip(sids, sids[1:] + [NOSTREAM]):
            entries[a]['right'] = b
        return sids[0] if sids else NOSTREAM

    entries[0] = dict(name='Root Entry', type=5, child=add(tree), start=ENDOFCHAIN, size=0, right=NOSTREAM)
    dir_sectors = -(-len(entries) // 4)
    fat = [FATSECT] + [s + 1 for s in range(1, dir_sectors)] + [ENDOFCHAIN]
    blob, sector = b'', 1 + dir_sectors
    for e in entries:
        if 'data' in e:
            n = len(e['data']) // 512
            e['start'], e['size'] = sector, len(e['data'])
            fat += [sector + k + 1 for k in range(n - 1)] + [ENDOFCHAIN]
            sector += n; blob += e['data']
    fat += [FREESECT] * (128 - len(fat))

    def entry(e):
        name = e['name'].encode('utf-16-le') + b'\0\0'
        return (name.ljust(64, b'\0') + struct.pack('<HBB', len(name), e['type'], 1)
                + struct.pack('<III', NOSTREAM, e.get('right', NOSTREAM), e['child']) + b'\0' * 36
                + struct.pack('<III', e['start'], e['size'], 0))
    empty = b'\0' * 68 + struct.pack('<III', NOSTREAM, NOSTREAM, NOSTREAM) + b'\0' * 48
    directory = b''.join(entry(e) for e in entries) + empty * (dir_sectors * 4 - len(entries))
    header = (readers.OLE2 + b'\0' * 16 + struct.pack('<HHHHH', 0x3E, 3, 0xFFFE, 9, 6) + b'\0' * 6
              + struct.pack('<9I', 0, 1, 1, 0, 4096, ENDOFCHAIN, 0, ENDOFCHAIN, 0) + struct.pack('<I', 0) + struct.pack('<I', FREESECT) * 108)
    return header + struct.pack('<128I', *fat) + directory + blob


def biff2(cells):
    """A raw BIFF2 (Excel 2.x) worksheet stream, no OLE2 container: {(row, col): str | int | float}, 0-based."""
    rec = lambda op, data: struct.pack('<HH', op, len(data)) + data
    rows = max(r for r, _ in cells) + 1; cols = max(c for _, c in cells) + 1
    out = rec(0x0009, struct.pack('<HH', 0x0007, 0x0010)) + rec(0x0000, struct.pack('<HHHH', 0, rows, 0, cols))
    for (r, c), v in sorted(cells.items()):
        head = struct.pack('<HH', r, c) + b'\x00\x00\x00'
        if isinstance(v, str):
            b = v.encode('latin-1'); out += rec(0x0004, head + bytes([len(b)]) + b)
        elif isinstance(v, int) and 0 <= v < 65536:
            out += rec(0x0002, head + struct.pack('<H', v))
        else:
            out += rec(0x0003, head + struct.pack('<d', v))
    return out + rec(0x000A, b'')


def rezip(content, change):
    """The same zip with change(name, data) -> (name or None, data) applied to every member."""
    src = zipfile.ZipFile(BytesIO(content)); out = BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            name, data = change(info.filename, src.read(info.filename))
            if name:
                dst.writestr(zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0)), data,
                             compress_type=zipfile.ZIP_STORED if name == 'mimetype' else zipfile.ZIP_DEFLATED)
    return out.getvalue()


def zip_of(members, compression=zipfile.ZIP_DEFLATED):
    out = BytesIO()
    with zipfile.ZipFile(out, 'w', compression) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return out.getvalue()


def lie_about_size(content, member, declared):
    """Rewrite the local and central directory 'uncompressed size' of `member` (a zip bomb trick)."""
    data = bytearray(content)
    info = zipfile.ZipFile(BytesIO(content)).getinfo(member)
    struct.pack_into('<I', data, info.header_offset + 22, declared)
    at = data.find(b'PK\x01\x02')
    while at >= 0:
        n = struct.unpack_from('<H', data, at + 28)[0]
        if bytes(data[at + 46:at + 46 + n]) == member.encode():
            struct.pack_into('<I', data, at + 24, declared)
        at = data.find(b'PK\x01\x02', at + 4)
    return bytes(data)


def cells_of(sheet):
    return {c['cell']: c['value'] for c in sheet['cells']}


def read_one(content, name='x'):
    return readers.read(content, name)[1][0]


LAYOUT = {ext: (FORMATS / f'layout_libreoffice.{ext}').read_bytes() for ext in ('xlsx', 'xls', 'ods')}


# --- files written by LibreOffice: layout and number formats ------------------------------------------
@pytest.mark.parametrize('ext', ['xlsx', 'xls', 'ods'])
def test_libreoffice_layout_workbook_gives_merges_hidden_rows_columns_and_sheets(ext):
    kind, sheets = readers.read(LAYOUT[ext], f'layout.{ext}')
    po, secret = sheets
    assert (po['name'], po['hidden'], secret['name'], secret['hidden']) == ('PO', False, 'Secret', True)
    assert po['merged'] == ['A1:C1', 'B6:B7']
    hidden = {c['cell'] for c in po['cells'] if c['hidden']}
    assert hidden == {'E1', 'A3'}                                 # column E and row 3 are hidden
    values = cells_of(po)
    assert values['B5'] == 1500 and values['B6'] == 'NAVY'
    assert (values['B2'], values['C2'], values['D2']) == ('0008', '000-123', 'ST-0008')   # ODS: from its number styles
    b2 = next(c for c in po['cells'] if c['cell'] == 'B2')
    assert (b2['raw'], b2['fmt']) == (8, '0000')
    assert readers.limitations(po) == [] and po['capabilities']['number_formats'] is True


def test_number_format_renderer_handles_only_digit_codes():
    cases = {(8, '0000'): '0008', (123, '000-000'): '000-123', (123, '000\\-000'): '000-123', (8, '"ST-"0000'): 'ST-0008',
             (8, '[Red]0000;-0000'): '0008', (8, '[$₹-4009]0000'): '₹0008', (0, '000-000'): '000-000', (12345, '0000'): None,
             (8, '0'): None, (1500, '#,##0'): None, (8, 'General'): None, (8.5, '0000'): None, (-8, '0000'): None,
             (8, '@'): None, (8, '00/00'): None, (8, '0.00%'): None, (8, 'dd-mm-yy'): None, (True, '0000'): None,
             (10 ** 15, '0000'): None, (8, '"unclosed0000'): None}
    assert {k: numfmt.render(*k) for k in cases} == cases
    assert numfmt.is_text_format('@') and numfmt.is_text_format('"\\@"') and not numfmt.is_text_format('0000')


def test_zero_height_rows_and_zero_width_columns_count_as_hidden():
    def patch(name, data):
        if name == 'xl/worksheets/sheet1.xml':
            data = data.replace(b'<sheetData>', b'<cols><col min="2" max="3" width="0" customWidth="1"/></cols><sheetData>', 1)
            data = data.replace(b'<row r="3"', b'<row r="3" ht="0" customHeight="1"', 1)
        return name, data
    sheet = read_one(rezip(flat_orders(), patch), 'zero.xlsx')
    hidden = {c['cell'] for c in sheet['cells'] if c['hidden']}
    assert {'B1', 'C1', 'A3', 'E3'} <= hidden and 'A2' not in hidden and 'D1' not in hidden


def test_strict_open_xml_is_read_with_the_backup_reader():
    def strict(name, data):
        if name.endswith(('.xml', '.rels')):
            data = (data.replace(b'http://schemas.openxmlformats.org/spreadsheetml/2006/main', b'http://purl.oclc.org/ooxml/spreadsheetml/main')
                    .replace(b'http://schemas.openxmlformats.org/officeDocument/2006/relationships', b'http://purl.oclc.org/ooxml/officeDocument/relationships'))
            if name == 'xl/workbook.xml':
                data = data.replace(b'<workbook ', b'<workbook conformance="strict" ', 1)
        return name, data
    kind, sheets = readers.read(rezip(LAYOUT['xlsx'], strict), 'strict.xlsx')
    po = sheets[0]
    assert kind == 'ooxml' and po['reader'] == 'calamine' and cells_of(po)['A2'] == 'Style'
    assert po['merged'] == ['A1:C1', 'B6:B7'] and po['capabilities']['hidden_rows'] is False
    assert 'Strict Open XML' in po['notes'][0]


def test_damaged_xlsx_falls_back_to_calamine_instead_of_failing():
    broken = rezip(LAYOUT['xlsx'], lambda n, d: (n, b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                                                   b'<numFmts><numFmt numFmtId="x" formatCode="0"/></numFmts></styleSheet>') if n == 'xl/styles.xml' else (n, d))
    po = read_one(broken, 'broken.xlsx')
    assert po['reader'] == 'calamine' and cells_of(po)['A1'] == 'LAYOUT TEST' and 'backup reader' in po['notes'][0]


def test_openpyxl_finding_no_worksheets_falls_back(monkeypatch):
    monkeypatch.setattr(readers, '_openpyxl_sheets', lambda content, merges=None: [])
    assert read_one(LAYOUT['xlsx'], 'x.xlsx')['reader'] == 'calamine'


def test_unreadable_xlsx_gets_a_plain_message():
    broken = rezip(LAYOUT['xlsx'], lambda n, d: (n, b'<worksheet><sheetData><row r="x"') if n.startswith('xl/worksheets/') else (n, d))
    with pytest.raises(readers.Unreadable, match='Open it in Excel, save it again'):
        readers.read(broken, 'broken.xlsx')


# --- .xls fallback chain --------------------------------------------------------------------------------
def test_xls_corruption_is_retried_leniently(monkeypatch):
    import xlrd
    real, calls = xlrd.open_workbook, []

    def flaky(**kw):
        calls.append(kw['ignore_workbook_corruption'])
        if not kw['ignore_workbook_corruption']:
            raise xlrd.compdoc.CompDocError('Workbook corruption: seen[2] == 4')
        assert isinstance(kw['logfile'], readers.io.StringIO)     # xlrd never prints to the server console
        return real(**kw)
    monkeypatch.setattr(xlrd, 'open_workbook', flaky)
    po = read_one(LAYOUT['xls'], 'x.xls')
    assert calls == [False, True] and po['reader'] == 'xlrd' and cells_of(po)['B2'] == '0008'


def test_xls_that_xlrd_cannot_read_falls_back_to_calamine(monkeypatch):
    import xlrd

    def broken(**kw):
        raise AssertionError('xlrd internal check')
    monkeypatch.setattr(xlrd, 'open_workbook', broken)
    po = read_one(LAYOUT['xls'], 'x.xls')
    assert po['reader'] == 'calamine' and cells_of(po)['A1'] == 'LAYOUT TEST' and po['merged'] == ['A1:C1', 'B6:B7']


def test_password_protected_xls_is_refused_with_what_to_do(monkeypatch):
    import xlrd
    monkeypatch.setattr(xlrd, 'open_workbook', lambda **kw: (_ for _ in ()).throw(xlrd.XLRDError('Workbook is encrypted')))
    with pytest.raises(readers.Unsupported, match='password-protected.*remove the password'):
        readers.read(LAYOUT['xls'], 'x.xls')


def test_raw_biff2_workbook_is_read_by_xlrd():
    old = biff2({(0, 0): 'Style', (0, 1): 'Qty', (1, 0): 'B2-1', (1, 1): 12, (2, 0): 'B2-2', (2, 1): 7.5})
    assert readers.sniff(old, 'old.xls') == 'xls'
    sheet = read_one(old, 'old.xls')
    assert sheet['reader'] == 'xlrd' and cells_of(sheet) == {'A1': 'Style', 'B1': 'Qty', 'A2': 'B2-1', 'B2': 12, 'A3': 'B2-2', 'B3': 7.5}


# --- OLE2 compound files: protected workbooks and other Office documents -------------------------------
@pytest.mark.parametrize('name,content,message', [
    ('encrypted.xlsx', lambda: (FORMATS / 'encrypted_msoffcrypto.xlsx').read_bytes(), readers.MSG_PASSWORD),
    ('irm.xlsx', lambda: cfb({'\x06DataSpaces': {'TransformInfo': {'DRMEncryptedTransform': {'\x06Primary': b'x'}}}, 'EncryptedPackage': b'e'}), readers.MSG_IRM),
    ('irm_package.xlsx', lambda: cfb({'EncryptedPackage': b'e', '\x06DataSpaces': {'DataSpaceMap': b'm'}}), readers.MSG_IRM),
    ('irm_legacy.xls', lambda: cfb({'\x06DataSpaces': {'DataSpaceMap': b'm'}, '\tDRMContent': b'd', 'Workbook': b'w'}), readers.MSG_IRM),
    ('po.xls', lambda: (FORMATS / 'word_libreoffice.doc').read_bytes(), readers.MSG_WORD),
    ('slides.xls', lambda: cfb({'PowerPoint Document': b'p', 'Current User': b'u'}), readers.MSG_POWERPOINT),
    ('mail.xls', lambda: cfb({'__substg1.0_0037001F': b's', '__properties_version1.0': b'p'}), readers.MSG_OUTLOOK),
    ('other.xls', lambda: cfb({'Contents': b'c'}), readers.MSG_OFFICE),
])
def test_ole2_files_are_classified_with_exact_messages(name, content, message):
    with pytest.raises(readers.Unsupported) as err:
        readers.sniff(content(), name)
    assert str(err.value) == message and 'upload' in message.lower()        # every refusal says what to do


def test_damaged_xls_in_an_ole2_container_gets_a_plain_message():
    with pytest.raises(readers.Unreadable, match=r'\.xls workbook could not be read.*save it as \.xlsx'):
        readers.read(cfb({'Workbook': b'garbage that is no BIFF'}), 'bad.xls')


# --- zip containers: real sizes, not declared ones ---------------------------------------------------
def test_zip_member_that_lies_about_its_size_is_refused_before_calamine():
    xlsb = zip_of({'[Content_Types].xml': b'<Types/>', 'xl/workbook.bin': b'\0' * 64, 'xl/worksheets/sheet1.bin': b'\0' * 3_000_000})
    liar = lie_about_size(xlsb, 'xl/worksheets/sheet1.bin', 1000)
    assert readers.sniff(liar, 'x.xlsb') == 'xlsb'                 # declared sizes look harmless
    with pytest.raises(readers.Unsupported, match='larger than the file says'):
        readers.read(liar, 'x.xlsb')


def test_zip_bomb_inside_declared_limits_is_measured():
    honest = zip_of({'[Content_Types].xml': b'<Types/>', 'xl/workbook.xml': b'<workbook/>', 'xl/sharedStrings.xml': b'0' * (41 * 1024 * 1024)})
    with pytest.raises(readers.Unreadable, match='expands to more than 40 MB'):
        readers.read(honest, 'x.xlsx')
    total = zip_of({'xl/workbook.bin': b'', **{f'xl/worksheets/sheet{i}.bin': b'0' * (30 * 1024 * 1024) for i in range(3)}})
    with pytest.raises(readers.Unreadable, match='80 MB'):
        readers.read(total, 'x.xlsb')                              # declared total is refused while sniffing
    lying_total = total
    for i in range(3):
        lying_total = lie_about_size(lying_total, f'xl/worksheets/sheet{i}.bin', 1000)
    with pytest.raises(readers.Unsupported, match='larger than the file says'):
        readers.read(lying_total, 'x.xlsb')


def test_zip_with_too_many_parts_is_refused_from_its_directory_record():
    many = zip_of({'xl/workbook.bin': b'', **{f'xl/p{i}': b'' for i in range(readers.MAX_ZIP_ENTRIES + 1)}}, zipfile.ZIP_STORED)
    with pytest.raises(readers.Unreadable, match='internal parts'):
        readers.sniff(many, 'x.xlsb')


def test_unusual_zip_compression_is_refused_with_what_to_do():
    odd = zip_of({'xl/workbook.bin': b'x' * 100, '[Content_Types].xml': b'<Types/>'}, zipfile.ZIP_BZIP2)
    with pytest.raises(readers.Unsupported, match='compression method Excel never writes'):
        readers.read(odd, 'x.xlsb')


@pytest.mark.parametrize('members,message', [
    ({'word/document.xml': b'<w/>', '[Content_Types].xml': b'<Types/>'}, readers.MSG_WORD),
    ({'ppt/presentation.xml': b'<p/>', '[Content_Types].xml': b'<Types/>'}, readers.MSG_POWERPOINT),
    ({'Index/Document.iwa': b'x'}, readers.MSG_NUMBERS),
    ({'mimetype': b'application/vnd.oasis.opendocument.text', 'content.xml': b'<x/>'}, readers.MSG_ODT),
    ({'photos/a.jpg': b'x'}, readers.MSG_ZIP),
])
def test_other_zip_documents_are_named(members, message):
    with pytest.raises(readers.Unsupported) as err:
        readers.sniff(zip_of(members), 'x.xlsx')
    assert str(err.value) == message


def test_zip_member_names_are_matched_case_insensitively():
    upper = rezip(LAYOUT['xlsx'], lambda n, d: (n.upper() if n.startswith('xl/') else n, d))
    assert readers.sniff(upper, 'x.xlsx') == 'ooxml'
    assert cells_of(read_one(upper, 'x.xlsx'))['A1'] == 'LAYOUT TEST'   # openpyxl refuses; calamine reads it


# --- sniffing text formats ------------------------------------------------------------------------------
def test_html_that_declares_the_spreadsheet_namespace_is_still_html():
    page = b'<html xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><body><table><tr><td>Style</td><td>0008</td></tr></table></body></html>'
    assert readers.sniff(page, 'export.xls') == 'html'
    assert cells_of(read_one(page, 'export.xls')) == {'A1': 'Style', 'B1': '0008'}
    prefixed = ('<?xml version="1.0"?><!DOCTYPE x [<!ELEMENT x ANY>]><ss:Workbook xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
                '<ss:Worksheet ss:Name="A"><ss:Table><ss:Row><ss:Cell><ss:Data ss:Type="String">x</ss:Data></ss:Cell></ss:Row></ss:Table></ss:Worksheet></ss:Workbook>').encode()
    assert readers.sniff(prefixed) == 'xml2003'


@pytest.mark.parametrize('codec', ['utf-16-le', 'utf-16-be'])
def test_unicode_text_without_a_byte_order_mark_is_read(codec):
    text = 'Style\tColour\tQty\r\n0008\tNAVY\t1,50,000\r\nST-9\tಕೆಂಪು\t12\r\n'.encode(codec)
    assert readers.sniff(text, 'orders.txt') == 'csv'
    assert cells_of(read_one(text, 'orders.txt')) == {'A1': 'Style', 'B1': 'Colour', 'C1': 'Qty', 'A2': '0008', 'B2': 'NAVY',
                                                      'C2': 150000, 'A3': 'ST-9', 'B3': 'ಕೆಂಪು', 'C3': 12}


@pytest.mark.parametrize('content,message', [
    (b'%PDF-1.4 no nul bytes here', 'PDF file, not an Excel workbook'),
    (b'\x89PNG\r\n\x1a\n', 'picture, not an Excel workbook'),
    (b'{\\rtf1\\ansi order}', 'RTF'),
    (b'\x1f\x8b\x08\x00', 'gzip'),
    (bytes(range(1, 32)) * 50, 'not an Excel workbook'),
])
def test_known_non_spreadsheets_are_named(content, message):
    with pytest.raises(readers.Unsupported, match=message):
        readers.sniff(content, 'x.xlsx')


# --- MHTML (single-file web page) ------------------------------------------------------------------
def mhtml(parts, boundary='----=_NextPart_01DA.TEST'):
    """An MHTML document laid out the way Excel's 'Single File Web Page' is: quoted-printable parts."""
    out = [f'MIME-Version: 1.0\r\nX-Document-Type: Workbook\r\nContent-Type: multipart/related; boundary="{boundary}"\r\n\r\n'
           'This document is a Single File Web Page, also known as a Web Archive file.\r\n']
    for location, ctype, charset, body in parts:
        encoded = quopri.encodestring(body.encode(charset)).decode('ascii')
        out.append(f'\r\n--{boundary}\r\nContent-Location: file:///C:/PO/{location}\r\nContent-Transfer-Encoding: quoted-printable\r\n'
                   f'Content-Type: {ctype}; charset="{charset}"\r\n\r\n{encoded}\r\n')
    out.append(f'\r\n--{boundary}--\r\n')
    return ''.join(out).encode('ascii')


EXCEL_MHTML = [
    ('Book1.htm', 'text/html', 'us-ascii',
     '<html xmlns:x="urn:schemas-microsoft-com:office:excel"><head><xml><x:ExcelWorkbook><x:ExcelWorksheets>'
     '<x:ExcelWorksheet><x:Name>Order</x:Name><x:WorksheetSource HRef="Book1_files/sheet001.htm"/></x:ExcelWorksheet>'
     '<x:ExcelWorksheet><x:Name>Sizes &amp; Qty</x:Name><x:WorksheetSource HRef="Book1_files/sheet002.htm"/></x:ExcelWorksheet>'
     '</x:ExcelWorksheets></x:ExcelWorkbook></xml></head><frameset rows="*,39"><frame src="Book1_files/sheet001.htm">'
     '<frame src="Book1_files/tabstrip.htm"></frameset></html>'),
    ('Book1_files/stylesheet.css', 'text/css', 'windows-1252', '.xl65\n\t{mso-number-format:"\\@";}\n.xl66 {display:none}'),
    ('Book1_files/tabstrip.htm', 'text/html', 'windows-1252', '<html><body><table><tr><td><a href="sheet001.htm">Order</a></td></tr></table></body></html>'),
    ('Book1_files/sheet001.htm', 'text/html', 'windows-1252',
     '<html xmlns:x="urn:schemas-microsoft-com:office:excel"><head><link rel=Stylesheet href=stylesheet.css></head><body><table>'
     '<tr><td>Buyer</td><td>CAFÉ KIDS</td></tr><tr><td>Style No</td><td class=xl65>0008</td></tr>'
     '<tr><td>Qty</td><td x:num="1500">1,500</td></tr></table></body></html>'),
    ('Book1_files/sheet002.htm', 'text/html', 'windows-1252',
     '<html><body><table><tr><td>Size</td><td>Qty</td></tr><tr class=xl66><td>S</td><td x:num>4</td></tr><tr><td>M</td><td x:num>6</td></tr></table></body></html>'),
]


def test_excel_mhtml_gives_one_named_sheet_per_worksheet_part():
    content = mhtml(EXCEL_MHTML)
    assert readers.sniff(content, 'po.xls') == 'mhtml'
    kind, sheets = readers.read(content, 'po.xls')
    assert kind == 'mhtml' and [s['name'] for s in sheets] == ['Order', 'Sizes & Qty']          # no tab strip
    assert cells_of(sheets[0]) == {'A1': 'Buyer', 'B1': 'CAFÉ KIDS', 'A2': 'Style No', 'B2': '0008', 'A3': 'Qty', 'B3': 1500}
    assert {c['cell'] for c in sheets[1]['cells'] if c['hidden']} == {'A2', 'B2'}                # class with display:none


def test_browser_mhtml_and_saved_emails():
    page = mhtml([('page.htm', 'text/html', 'utf-8', '<html><body><h1>PO 7781</h1><table><tr><td>Style</td><td>Qty</td></tr>'
                                                     '<tr><td>ZK-1</td><td>10</td></tr></table></body></html>')])
    assert cells_of(read_one(page, 'po.mht')) == {'A1': 'PO 7781', 'A3': 'Style', 'B3': 'Qty', 'A4': 'ZK-1', 'B4': 10}
    from email.mime.application import MIMEApplication
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    mail = MIMEMultipart(); mail.attach(MIMEText('Please find the PO attached.'))
    attachment = MIMEApplication(LAYOUT['xlsx']); attachment.add_header('Content-Disposition', 'attachment', filename='po.xlsx')
    mail.attach(attachment)
    with pytest.raises(readers.Unsupported) as err:
        readers.read(mail.as_bytes(), 'po.eml.txt')
    assert str(err.value) == readers.MSG_EMAIL


# --- HTML exports: text formats, hidden rows/columns, charsets, span limits ---------------------------
def test_html_export_keeps_text_formatted_codes_and_hidden_rows_and_columns():
    page = ('<html xmlns:x="urn:schemas-microsoft-com:office:excel"><head><style>.txt {mso-number-format:"\\@";}</style></head><body><table>'
            '<col width=64><col width=0 style="display:none"><col width=64>'
            '<tr><td>Style</td><td>internal</td><td>Qty</td></tr>'
            '<tr><td class=txt>1500</td><td>x</td><td x:num="1500">1,500</td></tr>'
            '<tr><td style=\'mso-number-format:"\\@"\'>00123</td><td>y</td><td x:str>0042</td></tr>'
            '<tr><td x:num="8">0008</td><td>z</td><td x:num="0.15">15%</td></tr>'
            '<tr style="display:none"><td>hidden</td></tr>'
            '<tr><td>1.234,56</td><td>1.23457E+11</td><td x:fmla="=SUM(C2:C3)" x:num="1500">1,500</td></tr></table></body></html>').encode()
    sheet = read_one(page, 'export.xls')
    cells = {c['cell']: c for c in sheet['cells']}
    assert (cells['A2']['value'], cells['A2']['fmt'], cells['C2']['value']) == ('1500', '@', 1500)
    assert (cells['A3']['value'], cells['C3']['value']) == ('00123', '0042')
    assert (cells['A4']['value'], cells['A4']['raw'], cells['C4']['value']) == ('0008', 8, 0.15)
    assert {c for c, v in cells.items() if v['hidden']} == {'B1', 'B2', 'B3', 'B4', 'A5', 'B6'}
    assert (cells['A6']['flag'], cells['B6']['flag']) == ('decimal_comma', 'scientific')
    assert (cells['C6']['formula'], cells['C6']['value']) == ('=SUM(C2:C3)', 1500)


def test_html_meta_charset_is_honoured():
    page = '<html><head><meta http-equiv="Content-Type" content="text/html; charset=shift_jis"></head><body><table><tr><td>色</td><td>紺</td></tr></table></body></html>'
    assert cells_of(read_one(page.encode('shift_jis'), 'jp.xls')) == {'A1': '色', 'B1': '紺'}
    cp1252 = '<table><tr><td>Señor – 5 €</td></tr></table>'.encode('cp1252')
    assert cells_of(read_one(cp1252, 'x.htm')) == {'A1': 'Señor – 5 €'}


def test_html_spans_are_bounded():
    big = b'<table><tr><td colspan="200" rowspan="500">X</td><td>Y</td></tr></table>'
    sheet = read_one(big, 'span.xls')
    assert sheet['merged'] == ['A1:GR50']                             # 200 x 50 = the 10,000-cell cap
    many = b'<table>' + b'<tr>' + b'<td colspan="200" rowspan="500">X</td>' * 120 + b'</tr></table>'
    with pytest.raises(readers.Unreadable, match='too many merged cells'):
        readers.read(many, 'span.xls')


def test_css_parsing_is_linear_on_hostile_style_sheets():
    hostile = b'<style>' + b'a' * 400_000 + b'/*' * 50_000 + b'</style><table><tr><td>ok</td></tr></table>'
    started = time.monotonic()
    assert cells_of(read_one(hostile, 'x.htm')) == {'A1': 'ok'}
    assert time.monotonic() - started < 5


# --- Excel 2003 XML Spreadsheet -------------------------------------------------------------------
SPREADSHEETML = '''<?xml version="1.0"?>
<?mso-application progid="Excel.Sheet"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:x="urn:schemas-microsoft-com:office:excel" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
 <Styles>
  <Style ss:ID="Default" ss:Name="Normal"><Alignment ss:Vertical="Bottom"/></Style>
  <Style ss:ID="s62"><NumberFormat ss:Format="0000"/></Style>
  <Style ss:ID="s63"><NumberFormat ss:Format="&quot;ST-&quot;0000"/></Style>
  <Style ss:ID="s64" ss:Parent="s62"><Font ss:Bold="1"/></Style>
  <Style ss:ID="s65"><NumberFormat ss:Format="Fixed"/></Style>
 </Styles>
 <Worksheet ss:Name="PO">
  <Table>
   <Column ss:Index="3" ss:StyleID="s62"/>
   <Column ss:Index="7" ss:Hidden="1" ss:Span="99999999"/>
   <Row><Cell><Data ss:Type="String">Style</Data></Cell><Cell ss:StyleID="s62"><Data ss:Type="Number">8</Data></Cell><Cell><Data ss:Type="Number">9</Data></Cell><Cell ss:StyleID="s63"><Data ss:Type="Number">8</Data></Cell><Cell ss:StyleID="s64"><Data ss:Type="Number">7</Data></Cell><Cell ss:StyleID="s65"><Data ss:Type="Number">2.5</Data></Cell><Cell><Data ss:Type="String">hidden col</Data></Cell></Row>
   <Row ss:Hidden="1" ss:Span="2"><Cell><Data ss:Type="String">hidden</Data></Cell></Row>
   <Row><Cell><Data ss:Type="String">after span</Data></Cell></Row>
   <Row><Cell ss:MergeAcross="16383" ss:MergeDown="5"><Data ss:Type="String">absurd merge</Data></Cell></Row>
   <Row><Cell ss:MergeAcross="2"><Data ss:Type="String">merge</Data></Cell><Cell><Data ss:Type="String">next</Data></Cell><Cell ss:Index="99999"><Data ss:Type="String">beyond</Data></Cell></Row>
   <Row ss:Index="2000000"><Cell><Data ss:Type="String">beyond Excel</Data></Cell></Row>
  </Table>
 </Worksheet>
</Workbook>'''


def test_spreadsheetml_styles_spans_and_limits():
    sheet = read_one(SPREADSHEETML.encode(), 'po.xml')
    cells = {c['cell']: c for c in sheet['cells']}
    assert [cells[a]['value'] for a in ('B1', 'C1', 'D1', 'E1', 'F1')] == ['0008', '0009', 'ST-0008', '0007', 2.5]
    assert (cells['B1']['raw'], cells['F1']['fmt']) == (8, 'Fixed')
    assert cells['A2']['hidden'] and 'A5' in cells and cells['A5']['value'] == 'after span'   # rows 2-4 hidden, next row 5
    assert cells['G1']['hidden'] and 'A2000000' not in cells and not any(c['col'] > 16384 for c in sheet['cells'])
    assert sheet['merged'] == ['A7:C7'] and cells['D7']['value'] == 'next'
    assert any('merged areas larger' in n for n in sheet['notes']) and any('beyond Excel' in n for n in sheet['notes'])


def test_spreadsheetml_without_utf8_declaration_or_bom_is_decoded():
    utf16 = SPREADSHEETML.replace('<?xml version="1.0"?>', '<?xml version="1.0" encoding="UTF-16"?>').encode('utf-16-le')
    assert readers.sniff(utf16, 'po.xml') == 'xml2003'
    assert cells_of(read_one(utf16, 'po.xml'))['B1'] == '0008'
    latin = SPREADSHEETML.replace('>Style<', '>Señor<').encode('cp1252')
    assert cells_of(read_one(latin, 'po.xml'))['A1'] == 'Señor'


# --- ODS layout pass --------------------------------------------------------------------------------
def test_ods_layout_pass_counts_repeats_without_expanding_them():
    def bomb(name, data):
        if name == 'content.xml':
            end = data.index(b'</table:table>')
            data = (data[:end] + b'<table:table-row table:number-rows-repeated="1000000" table:visibility="collapse">'
                    b'<table:table-cell table:number-columns-repeated="16384"/></table:table-row>'
                    b'<table:table-row><table:table-cell table:number-columns-spanned="16384" table:number-rows-spanned="1048576"/></table:table-row>'
                    + data[end:])
        return name, data
    content = rezip(LAYOUT['ods'], bomb)
    started = time.monotonic()
    layout = readers._ods_layout(content)
    assert time.monotonic() - started < 5
    assert (8, 1000007) in layout['PO']['rows'] and (1000008, 1, 1048576, 16384) in layout['PO']['merges']
    po = read_one(content, 'bomb.ods')                 # the empty repeats add no cells; the absurd merge is ignored
    assert po['merged'] == ['A1:C1', 'B6:B7'] and any('merged areas larger' in n for n in po['notes'])


# --- CSV / text exports ---------------------------------------------------------------------------
def test_csv_sep_line_excel_text_wrappers_and_lakh_numbers():
    text = b'sep=;\r\nStyle;Address;Qty;Price;Code\r\nA1;"12, MG Road, Tiruppur, TN, India";1,50,000;12,34,567.50;="00123"\r\nA2;"4, Main St, Erode, TN";6;2,50;=1+1\r\n'
    sheet = read_one(text, 'x.csv')
    cells = {c['cell']: c for c in sheet['cells']}
    assert [cells[a]['value'] for a in ('A1', 'B2', 'C2', 'D2', 'E2')] == ['Style', '12, MG Road, Tiruppur, TN, India', 150000, 1234567.5, '00123']
    assert (cells['D3']['value'], cells['D3']['flag'], cells['E3']['value']) == ('2,50', 'decimal_comma', '=1+1')
    assert sheet['capabilities'] == readers.TEXT_CAPS


def test_csv_delimiter_is_chosen_by_consistency_not_character_count():
    # More commas than semicolons, but only ';' splits every row into the same columns.
    text = b'Style;Address;Qty\r\nA1;"12, MG Road, Tiruppur, TN, India";5\r\nA2;"4, Main St, Erode, TN";6\r\n'
    assert readers._delimiter(text.decode()) == ';'
    assert cells_of(read_one(text, 'x.csv'))['C3'] == 6
    assert readers._delimiter('Style|Qty\nA|1\nB|2\n') == '|'
    assert readers._delimiter('Single column\nA\nB\n') == ','


def test_csv_with_an_unclosed_quote_gets_a_plain_message():
    with pytest.raises(readers.Unreadable, match='quotation mark is not closed'):
        readers.read(b'Style,Note\r\nA1,"' + b'x' * 200_000 + b'\r\n', 'x.csv')


def test_csv_long_codes_and_damaged_codes_are_flagged_not_changed():
    sheet = read_one(b'EAN,Barcode,Code\r\n1234567890123456789,1.23457E+11,0000000000000000012\r\n', 'x.csv')
    cells = {c['cell']: c for c in sheet['cells']}
    assert (cells['A2']['value'], cells['A2']['flag']) == ('1234567890123456789', 'big_number')
    assert (cells['B2']['value'], cells['B2']['flag']) == ('1.23457E+11', 'scientific')
    assert cells['C2']['value'] == '0000000000000000012' and 'flag' not in cells['C2']


@pytest.mark.parametrize('content,expected', [
    ('Colour,Note\r\nCRÈME,Señor 5€\r\n'.encode('cp1252'), 'CRÈME'),
    ('Colour,Note\r\nCRÈME,Señor 5€\r\n'.encode('utf-8'), 'CRÈME'),
    ('Colour,Note\r\nCRÈME,Señor 5€ and ₹ 9\r\n'.encode('utf-8') + b'\x96 stray\r\n', 'CRÈME'),   # UTF-8 with one stray byte
])
def test_text_encodings_are_detected(content, expected):
    assert cells_of(read_one(content, 'x.csv'))['A2'] == expected


def test_typed_text_values_follow_indian_and_western_grouping():
    assert [readers.typed(v) for v in ('1,50,000', '12,34,567.50', '1,23,45,678', '1,234,567', '12,34', '1.234,56', '0,500',
                                       '1234567890123456', '123456789012345', '-2,50,000')] == \
        [150000, 1234567.5, 12345678, 1234567, '12,34', '1.234,56', '0,500', '1234567890123456', 123456789012345, -250000]


# --- JSON --------------------------------------------------------------------------------------------
@pytest.mark.parametrize('content,message', [
    (b'[{"Style": "A", "Qty": NaN}]', 'NaN'),
    (b'[{"Style": "A", "Qty": -Infinity}]', 'Infinity'),
    (b'[{"Style": "A", "Style": "B"}]', 'the key "Style" twice'),
    (b'[' * 100_000 + b']' * 100_000, 'nested too deeply'),
    (b'[{"Qty": ' + b'9' * 5000 + b'}]', 'more than 4,300 digits'),
    (b'{"Style": }', 'not valid JSON'),
])
def test_json_that_cannot_be_trusted_is_refused_plainly(content, message):
    with pytest.raises(readers.Unreadable, match=message):
        readers.read(content, 'x.json')


def test_json_decimals_ndjson_and_api_envelopes():
    sheet = read_one(b'[{"Price": 2.675, "Exact": 0.1000000000000000000001, "Qty": 1e3, "Big": 1E+999999999}]', 'x.json')
    assert cells_of(sheet) == {'A1': 'Price', 'B1': 'Exact', 'C1': 'Qty', 'D1': 'Big', 'A2': 2.675,
                               'B2': '0.1000000000000000000001', 'C2': 1000, 'D2': '1E+999999999'}
    json.dumps(sheet)                                                 # no Decimal reaches the cell model
    rows = {'A1': 'Style', 'B1': 'Qty', 'A2': 'A', 'B2': 1, 'A3': 'B', 'B3': 2}
    assert cells_of(read_one(b'{"Style": "A", "Qty": 1}\n{"Style": "B", "Qty": 2}\n', 'x.ndjson.json')) == rows
    assert cells_of(read_one(b'{"count": 2, "data": [{"Style": "A", "Qty": 1}, {"Style": "B", "Qty": 2}]}', 'x.json')) == rows
    assert cells_of(read_one(b'{"columns": ["Style", "Qty"], "rows": [["A", 1], ["B", 2]]}', 'x.json')) == rows
    doc = read_one(b'{"PO No": "X1", "rows": [{"Style": "A"}]}', 'x.json')     # a document, not an envelope
    assert cells_of(doc)['A1'] == 'PO No'


def test_clean_turns_decimals_into_consistent_values():
    assert [readers.clean(Decimal(v)) for v in ('2.675', '10', '1E+3', '1e20', '-12.50', '0.1000000000000000000001', 'NaN')] == \
        [2.675, 10, 1000, 1e20, -12.5, '0.1000000000000000000001', '#NUM!']
    assert readers.clean(float('inf')) == '#NUM!' and readers.clean('bad \ud800 text') == 'bad \ufffd text'


# --- merged-range limits and capability flags for every format --------------------------------------
def test_merged_ranges_are_clamped_and_absurd_ones_ignored():
    b = readers.Builder('S')
    b.merge(1, 1, 2, 3); b.merge(1, 1, 1, 1); b.merge(5, 5, 1, 1); b.merge(1, 1, readers.EXCEL_ROWS, readers.EXCEL_COLS)
    b.merge(1048570, 16380, 2_000_000, 99_999)
    sheet = b.sheet()
    assert sheet['merged'] == ['A1:C2', 'A1:E5', 'XEZ1048570:XFD1048576'] and sheet['notes'] == ['1 merged areas larger than 10,000 cells were ignored.']


FORMAT_SAMPLES = [
    ('ooxml', 'x.xlsx', lambda: LAYOUT['xlsx'], 'openpyxl'), ('xls', 'x.xls', lambda: LAYOUT['xls'], 'xlrd'),
    ('ods', 'x.ods', lambda: LAYOUT['ods'], 'calamine+layout'), ('xlsb', 'x.xlsb', lambda: (FORMATS / 'calamine_issues.xlsb').read_bytes(), 'calamine'),
    ('xml2003', 'x.xml', lambda: SPREADSHEETML.encode(), 'xml2003'), ('html', 'x.xls', lambda: as_html(flat_orders()), 'html'),
    ('mhtml', 'x.mht', lambda: mhtml(EXCEL_MHTML), 'html'), ('csv', 'x.csv', lambda: as_csv(flat_orders()), 'csv'),
    ('json', 'x.json', lambda: as_json_records(flat_orders()), 'json'),
]


@pytest.mark.parametrize('kind,name,content,reader', FORMAT_SAMPLES)
def test_every_sheet_says_what_its_format_could_provide(kind, name, content, reader):
    got_kind, sheets = readers.read(content(), name)
    assert got_kind == kind
    for s in sheets:
        assert set(s['capabilities']) == set(readers.CAPABILITY_KEYS) and set(s['capabilities'].values()) <= {True, False, None}
        assert s['reader'] == reader and isinstance(s['notes'], list)
        assert all(isinstance(n, str) and n.endswith('.') for n in readers.limitations(s))
    if kind == 'xlsb':
        caps = sheets[0]['capabilities']
        assert (caps['merged'], caps['hidden_rows'], caps['number_formats']) == (False, False, False)
        assert 'merged cells, hidden rows, hidden columns and number formats' in readers.limitations(sheets[0])[0]


# --- every failure is a plain message -----------------------------------------------------------------
@pytest.mark.parametrize('error', [AssertionError('xlrd check'), KeyError('sheet'), zipfile.BadZipFile('bad'), ValueError('library detail'),
                                   __import__('csv').Error('csv detail')])
def test_unexpected_reader_errors_become_plain_messages(monkeypatch, error):
    monkeypatch.setitem(readers.READERS, 'csv', lambda content: (_ for _ in ()).throw(error))
    with pytest.raises(readers.Unreadable) as err:
        readers.read(b'a,b\r\n1,2\r\n', 'x.csv')
    assert str(err.value) == 'This text file could not be read. Open it in Excel, save it again as an Excel Workbook (.xlsx) and upload.'


def test_memory_and_recursion_errors_become_plain_messages(monkeypatch):
    monkeypatch.setitem(readers.READERS, 'csv', lambda content: (_ for _ in ()).throw(MemoryError()))
    with pytest.raises(readers.Unreadable, match='more memory to read than is safe'):
        readers.read(b'a,b\r\n', 'x.csv')
    monkeypatch.setitem(readers.READERS, 'csv', lambda content: (_ for _ in ()).throw(RecursionError()))
    with pytest.raises(readers.Unreadable, match='nested too deeply'):
        readers.read(b'a,b\r\n', 'x.csv')


# --- process isolation --------------------------------------------------------------------------------
def test_isolated_read_gives_the_same_result_as_inline(monkeypatch):
    before = readers.stats['isolated']
    for name, content in (('layout.xlsx', LAYOUT['xlsx']), ('orders.csv', as_csv(flat_orders())), ('po.mht', mhtml(EXCEL_MHTML))):
        isolated = readers.read(content, name)
        monkeypatch.setenv(readers.INLINE_ENV, '1')
        assert readers.read(content, name) == tuple(isolated)
        monkeypatch.delenv(readers.INLINE_ENV)
    assert readers.stats['isolated'] == before + 3


def test_isolated_refusals_keep_their_class_and_message():
    with pytest.raises(readers.Unsupported) as err:
        readers.read((FORMATS / 'encrypted_msoffcrypto.xlsx').read_bytes(), 'enc.xlsx')
    assert str(err.value) == readers.MSG_PASSWORD
    with pytest.raises(readers.Unreadable, match='NaN'):
        readers.read(b'[{"Qty": NaN}]', 'x.json')


def test_isolated_worker_is_stopped_at_the_time_limit():
    started = time.monotonic()
    with pytest.raises(isolation.WorkerFailed) as err:
        isolation.run(isolation.diagnose, ('sleep', 60), timeout=1)
    assert err.value.reason == 'timeout' and time.monotonic() - started < 15


@pytest.mark.skipif(not sys.platform.startswith('linux'), reason='RLIMIT_AS is enforced on Linux; Windows uses a Job Object')
def test_isolated_worker_memory_is_capped():
    limits = isolation.run(isolation.diagnose, ('limits',))
    assert 0 < limits['address_space'] <= readers.PARSE_MEMORY and limits['cpu_seconds'] > 0
    with pytest.raises(isolation.WorkerFailed) as err:
        isolation.run(isolation.diagnose, ('allocate', 3 * 1024 ** 3), memory=512 * 1024 ** 2)
    assert err.value.reason == 'memory'
    assert isolation.run(isolation.diagnose, ('allocate', 1024 ** 2)) == 1024 ** 2


def test_isolated_worker_crash_is_reported():
    with pytest.raises(isolation.WorkerFailed) as err:
        isolation.run(isolation.diagnose, ('exit', 3))
    assert err.value.reason == 'crashed' and '3' in err.value.detail


def test_isolated_failures_become_plain_messages(monkeypatch):
    for reason in ('timeout', 'memory', 'crashed', 'busy', 'too_large', 'error'):
        monkeypatch.setattr(isolation, 'run', lambda *a, reason=reason, **k: (_ for _ in ()).throw(isolation.WorkerFailed(reason)))
        with pytest.raises(readers.Unreadable) as err:
            readers.read(b'a,b\r\n', 'x.csv')
        assert str(err.value) == readers.WORKER_MESSAGES.get(reason, readers.WORKER_MESSAGES['crashed'])


def test_isolated_spawn_failure_falls_back_to_inline_and_is_reported(monkeypatch):
    monkeypatch.setattr(isolation, 'run', lambda *a, **k: (_ for _ in ()).throw(isolation.WorkerFailed('spawn', 'no processes')))
    monkeypatch.setitem(readers.stats, 'inline_fallback', 0)
    assert cells_of(read_one(b'a,b\r\n1,2\r\n', 'x.csv')) == {'A1': 'a', 'B1': 'b', 'A2': 1, 'B2': 2}
    check = next(c for c in readers.selfcheck()['checks'] if c['name'] == 'Upload isolation')
    assert readers.stats['inline_fallback'] == 1 and check['ok'] is False and 'Restart POPPYS AI' in check['message']


def test_windows_job_object_structures_match_winnt_h():
    import ctypes
    basic, io_counters, extended = isolation.job_structures()
    if ctypes.sizeof(ctypes.c_void_p) == 8:
        assert (ctypes.sizeof(basic), ctypes.sizeof(io_counters), ctypes.sizeof(extended)) == (64, 48, 144)
        assert (basic.LimitFlags.offset, extended.ProcessMemoryLimit.offset) == (16, 112)


# --- platform self-check (health page) ------------------------------------------------------------
def test_isolated_selfcheck_reports_parsers_packages_and_isolation():
    import pyexpat
    result = readers.selfcheck(probe=True)
    checks = {c['name']: c for c in result['checks']}
    assert {'Python', 'XML parser (Expat)', 'openpyxl', 'xlrd', 'python-calamine', 'defusedxml', 'olefile', 'charset-normalizer',
            'Upload isolation'} <= set(checks)
    assert checks['XML parser (Expat)']['ok'] == (pyexpat.version_info >= (2, 7, 2))
    if not checks['XML parser (Expat)']['ok']:
        assert '2.7.2' in checks['XML parser (Expat)']['message'] and 'setup.ps1' in checks['XML parser (Expat)']['message']
    assert checks['Python']['ok'] == (sys.version_info[:3] >= (3, 13, 7))
    assert checks['xlrd']['version'] == '2.0.2' and checks['Upload isolation']['ok'] is True
    assert result['ok'] == all(c['ok'] for c in result['checks'])


def test_selfcheck_warns_when_uploads_are_parsed_inline():
    check = next(c for c in readers.selfcheck()['checks'] if c['name'] == 'Upload isolation')
    assert check['ok'] is False and readers.INLINE_ENV in check['message']


def test_formatting_down_to_the_last_row_does_not_make_a_sheet_too_large():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = 'PO'
    ws['A1'], ws['B1'], ws['A2'], ws['B2'] = 'Style', 'Qty', '0008', 5
    for r in range(3, 2000):
        ws.cell(r, 1).font = Font(bold=True)
    ws.cell(500_000, 40).font = Font(bold=True)          # one formatted, empty cell far away
    buf = BytesIO(); wb.save(buf)
    sheet = read_one(buf.getvalue(), 'styled.xlsx')
    assert sheet['reader'] == 'openpyxl' and (sheet['rows'], sheet['cols']) == (2, 2)
    wb = Workbook(); wb.active['A1'] = 1; wb.active['XFD1048576'] = 2; buf = BytesIO(); wb.save(buf)
    with pytest.raises(readers.Unreadable, match='spread over too wide an area'):
        readers.read(buf.getvalue(), 'wide.xlsx')


def test_very_long_digit_strings_stay_text():
    sheet = read_one(b'Code,Price\r\n' + b'7' * 5000 + b',1234567890123456789.5\r\n', 'x.csv')
    cells = {c['cell']: c for c in sheet['cells']}
    assert cells['A2']['value'] == '7' * 5000 and cells['A2']['flag'] == 'big_number'
    assert (cells['B2']['value'], cells['B2']['flag']) == ('1234567890123456789.5', 'big_number')


def test_isolated_worker_answer_of_the_wrong_shape_is_refused(monkeypatch):
    for answer in (['ooxml'], ['nonsense', []], ['csv', [{'name': 'x'}]], 'text'):
        monkeypatch.setattr(isolation, 'run', lambda *a, answer=answer, **k: answer)
        with pytest.raises(readers.Unreadable) as err:
            readers.read(b'a,b\r\n', 'x.csv')
        assert str(err.value) == readers.WORKER_MESSAGES['crashed']


def test_ods_number_styles_become_excel_codes():
    from xml.etree.ElementTree import fromstring
    N = 'xmlns:number="urn:oasis:names:tc:opendocument:xmlns:datastyle:1.0" xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"'
    def code(body):
        return readers._ods_code(fromstring(f'<number:number-style {N}>{body}</number:number-style>'))
    assert code('<number:number number:decimal-places="0" number:min-integer-digits="4"/>') == '0000'
    assert code('<number:text>ST-</number:text><number:number number:min-integer-digits="4"/>') == '"ST-"0000'
    assert code('<number:number number:min-integer-digits="6"><number:embedded-text number:position="3">-</number:embedded-text></number:number>') == '000"-"000'
    assert code('<number:number number:decimal-places="0" number:min-integer-digits="1" number:grouping="true"/>') is None
    assert code('<number:number number:decimal-places="2" number:min-integer-digits="1"/>') is None
    assert code('<number:text>no digits</number:text>') is None
    assert numfmt.render(123, '000"-"000') == '000-123'


def test_isolated_worker_that_cannot_start_falls_back_to_inline():
    """A server started from a script on stdin has a __main__ that spawn cannot re-import: uploads must
    still be read (in-process, logged), not all fail as "damaged"."""
    import os
    import subprocess
    root = Path(__file__).resolve().parents[1]
    script = ('import sys; sys.path.insert(0, %r)\n'
              'from backend import readers\n'
              'kind, sheets = readers.read(b"Style,Qty\\r\\nA,1\\r\\n", "x.csv")\n'
              'print(kind, sheets[0]["cells"][3]["value"], readers.stats["inline_fallback"])\n') % str(root)
    env = {k: v for k, v in os.environ.items() if k != readers.INLINE_ENV}
    out = subprocess.run([sys.executable, '-'], input=script, capture_output=True, text=True, timeout=120, env=env, cwd=root)
    assert out.stdout.split() == ['csv', '1', '1'], out.stderr[-2000:]
    assert 'Could not start an upload worker' in out.stderr


def test_worker_limits_can_be_tuned_within_bounds(monkeypatch):
    monkeypatch.setenv('POPPYS_PARSE_TIMEOUT', '5'); monkeypatch.setenv('POPPYS_PARSE_MEMORY_MB', 'lots')
    assert readers._env_int('POPPYS_PARSE_TIMEOUT', 60, 10, 600) == 10           # never below 10 s
    assert readers._env_int('POPPYS_PARSE_MEMORY_MB', 1536, 256, 16384) == 1536  # garbage keeps the default
    monkeypatch.setenv('POPPYS_PARSE_TIMEOUT', '120')
    assert readers._env_int('POPPYS_PARSE_TIMEOUT', 60, 10, 600) == 120


def test_native_reader_panics_become_plain_messages(monkeypatch):
    """pyo3's PanicException (raised when the Rust calamine library panics on a damaged file) derives
    from BaseException, so `except Exception` alone would turn it into a server error."""
    import python_calamine as pc
    import xlrd

    class PanicException(BaseException):
        pass

    def panic(*a, **k):
        raise PanicException('range end index 2337 out of range for slice of length 66')
    monkeypatch.setattr(pc.CalamineWorkbook, 'from_filelike', panic)
    with pytest.raises(readers.Unreadable, match='could not be opened'):
        readers.read((FORMATS / 'calamine_issues.xlsb').read_bytes(), 'x.xlsb')
    monkeypatch.setattr(xlrd, 'open_workbook', lambda **kw: (_ for _ in ()).throw(AssertionError('xlrd')))
    with pytest.raises(readers.Unreadable) as err:                 # xlrd fails, then calamine panics
        readers.read(LAYOUT['xls'], 'x.xls')
    assert str(err.value) == readers.MSG_XLS_DAMAGED
