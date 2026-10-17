"""Spreadsheet readers. Every accepted upload becomes the same cell model, so every layout detector,
template and validator works the same whatever format the buyer sent:

    [{name, hidden, rows, cols, merged: ["A1:C1"], cells: [{cell, row, col, value, formula, hidden, error}]}]

Formats (decided from the file's bytes, never from its name alone):
  .xlsx .xlsm .xltx .xltm   openpyxl (formulas, cached results, number formats, hidden rows/columns; macros never run)
  .xls (Excel 97-2003)      xlrd (values, dates, leading-zero formats, hidden rows/columns, merged cells)
  .xlsb, .ods               python-calamine (values, dates, sheet visibility, merged cells where the format stores them)
  Excel 2003 XML (.xml)     defusedxml (no entity expansion)
  HTML saved as .xls/.htm   html.parser (what ERP "Export to Excel" buttons often produce; laid out as Excel would)
  .csv .tsv .txt            csv with encoding + delimiter detection (UTF-8/UTF-16 "Unicode text"/Windows-1252)
  .json                     records, rows, sheet maps or a single nested document

Cell contents are data only: nothing is evaluated or executed. Size limits protect the laptop.
"""
from __future__ import annotations
import csv
from datetime import date, datetime, time, timedelta
from html.parser import HTMLParser
from io import BytesIO, StringIO
import json
import re
import zipfile

MAX_BYTES = 12 * 1024 * 1024
MAX_SHEETS = 30
MAX_CELLS = 250_000            # per sheet
MAX_TOTAL_CELLS = 1_000_000    # per file
MAX_ZIP_ENTRIES = 10_000
MAX_EXPANDED = 80 * 1024 * 1024
MAX_JSON_DEPTH = 20

OLE2 = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
ZIP = b'PK\x03\x04'
ODS_MIME = b'application/vnd.oasis.opendocument.spreadsheet'
SS_NS = 'urn:schemas-microsoft-com:office:spreadsheet'

ACCEPT = ('.xlsx', '.xlsm', '.xltx', '.xltm', '.xls', '.xlsb', '.ods', '.xml', '.htm', '.html', '.csv', '.tsv', '.txt', '.json')
LABELS = {
    'ooxml': 'Excel workbook', 'xls': 'Excel 97-2003 workbook', 'xlsb': 'Excel binary workbook',
    'ods': 'OpenDocument spreadsheet', 'xml2003': 'Excel 2003 XML spreadsheet', 'html': 'Excel web-page export',
    'csv': 'CSV / text export', 'json': 'JSON data',
}


class Unsupported(ValueError):
    """The upload is not a spreadsheet this product reads (shown to the user as-is)."""


def address(col, row):
    label = ''
    while col:
        col, rem = divmod(col - 1, 26); label = chr(65 + rem) + label
    return f'{label}{row}'


def clean(value):
    """One value convention for every reader: trimmed text, ISO dates, whole floats as ints."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.replace('\xa0', ' ').strip()
        return value or None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(timespec='seconds')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec='seconds')
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return int(value)
    return value


NUMBER = re.compile(r'^[-+]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$')


def typed(text):
    """Text from CSV/HTML/XML exports, typed the way Excel would show it. Codes with leading zeros
    (style 0008) stay text; '1,500' becomes 1500."""
    value = clean(text)
    if not isinstance(value, str) or not NUMBER.match(value):
        return value
    digits = value.lstrip('+-')
    if len(digits) > 1 and digits[0] == '0' and digits[1] != '.':
        return value
    number = float(value.replace(',', ''))
    return int(number) if number.is_integer() and '.' not in value else number


class Builder:
    """Collects cells for one sheet with limits and the shared cell format."""

    def __init__(self, name, hidden=False):
        self.name = str(name)[:120] or 'Sheet1'; self.hidden = hidden; self.cells = {}; self.merged = []
        self.hidden_rows = set(); self.hidden_cols = set()

    def put(self, row, col, value, formula=None, error=False):
        value = clean(value)
        if value is None:
            return
        if len(self.cells) >= MAX_CELLS:
            raise ValueError(f'{self.name}: sheet exceeds {MAX_CELLS:,} filled cells; trim it and upload again.')
        self.cells[(row, col)] = dict(cell=address(col, row), row=row, col=col, value=value, formula=formula, error=error)

    def sheet(self):
        rows = max((r for r, _ in self.cells), default=0); cols = max((c for _, c in self.cells), default=0)
        cells = [dict(c, hidden=c['row'] in self.hidden_rows or c['col'] in self.hidden_cols) for _, c in sorted(self.cells.items())]
        return dict(name=self.name, hidden=self.hidden, rows=rows, cols=cols, merged=self.merged, cells=cells)


# --- format detection ---------------------------------------------------------------------------
def _zip_kind(content):
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_ENTRIES or sum(z.file_size for z in infos) > MAX_EXPANDED:
                raise ValueError('Workbook expands to more than 80 MB; it cannot be read safely.')
            names = set(archive.namelist())
            if 'xl/workbook.bin' in names:
                return 'xlsb'
            if 'mimetype' in names and archive.read('mimetype').strip() == ODS_MIME:
                return 'ods'
            if '[Content_Types].xml' in names and any(n.startswith('xl/') for n in names):
                return 'ooxml'
    except zipfile.BadZipFile as exc:
        raise Unsupported('This workbook is damaged or incomplete. Open it in Excel, save it again and upload.') from exc
    raise Unsupported('This file is a zip archive, not an Excel workbook.')


def _decode(content):
    if content.startswith(b'\xef\xbb\xbf'):
        return content[3:].decode('utf-8', errors='replace')
    if content.startswith((b'\xff\xfe', b'\xfe\xff')):
        return content.decode('utf-16', errors='replace')
    try:
        return content.decode('utf-8')
    except UnicodeDecodeError:
        return content.decode('cp1252', errors='replace')


def sniff(content, filename=''):
    """The real format of the bytes: ooxml, xls, xlsb, ods, xml2003, html, csv or json."""
    if not content:
        raise Unsupported('The file is empty.')
    if content.startswith(ZIP):
        return _zip_kind(content)
    if content.startswith(OLE2):
        return 'xls'
    head = _decode(content[:4096]).lstrip('﻿ \t\r\n')
    lower = head.lower()
    if head[:1] in '{[':
        return 'json'
    if head.startswith('<'):
        if SS_NS in head or 'progid="excel.sheet"' in lower.replace(' ', ''):
            return 'xml2003'
        if '<table' in _decode(content).lower():
            return 'html'
        raise Unsupported('This XML/HTML file contains no spreadsheet table.')
    if b'\x00' in content[:4096] and not content.startswith((b'\xff\xfe', b'\xfe\xff')):
        raise Unsupported('This is not an Excel workbook. Upload the buyer file as .xlsx, .xls, .xlsb, .csv or .json.')
    if str(filename).lower().endswith(('.csv', '.tsv', '.txt')) or '\t' in head or ',' in head or ';' in head:
        return 'csv'
    raise Unsupported('This is not an Excel workbook. Upload the buyer file as .xlsx, .xls, .xlsb, .csv or .json.')


def read(content, filename=''):
    """(format, sheets) for an upload. Raises Unsupported or ValueError with a user-facing reason."""
    if len(content) > MAX_BYTES:
        raise ValueError('The file is larger than 12 MB.')
    kind = sniff(content, filename)
    sheets = READERS[kind](content)
    sheets = [s for s in sheets if s['cells'] or not s['hidden']]
    if not sheets:
        raise ValueError('The workbook has no sheets with data.')
    if len(sheets) > MAX_SHEETS:
        raise ValueError(f'The workbook has more than {MAX_SHEETS} sheets.')
    if sum(len(s['cells']) for s in sheets) > MAX_TOTAL_CELLS:
        raise ValueError(f'The workbook has more than {MAX_TOTAL_CELLS:,} filled cells.')
    return kind, sheets


# --- .xlsx .xlsm .xltx .xltm ---------------------------------------------------------------------------
def read_ooxml(content):
    from openpyxl import load_workbook
    _zip_kind(content)   # size checks before parsing XML
    try:
        raw = load_workbook(BytesIO(content), data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(content), data_only=True, keep_links=False)
    except Exception as exc:  # openpyxl raises many types for damaged files
        raise ValueError('This workbook could not be opened. Open it in Excel, save it again and upload.') from exc
    if len(raw.worksheets) > MAX_SHEETS:
        raise ValueError(f'The workbook has more than {MAX_SHEETS} sheets.')
    sheets = []
    for ws in raw.worksheets:
        if ws.max_row * ws.max_column > MAX_CELLS * 4:
            raise ValueError(f'{ws.title}: sheet is too large; trim unused formatting and upload again.')
        b = Builder(ws.title, ws.sheet_state != 'visible')
        b.hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
        from openpyxl.utils import column_index_from_string
        for key, d in ws.column_dimensions.items():
            if d.hidden:          # one dimension entry can cover a range of columns (min..max)
                first = d.min or column_index_from_string(key)
                b.hidden_cols.update(range(first, (d.max or first) + 1))
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                formula = cell.value if cell.data_type == 'f' else None
                value = _ooxml_value(cached[ws.title][cell.coordinate]) if formula else _ooxml_value(cell)
                if formula and value is None:   # keep uncalculated formulas visible to the parser
                    b.cells[(cell.row, cell.column)] = dict(cell=cell.coordinate, row=cell.row, col=cell.column, value=None,
                                                            formula=str(formula), error=False)
                    continue
                b.put(cell.row, cell.column, value, str(formula) if formula else None, cell.data_type == 'e')
        b.merged = [str(r) for r in ws.merged_cells.ranges]
        sheets.append(b.sheet())
    raw.close(); cached.close()
    return sheets


def _ooxml_value(cell):
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and re.fullmatch(r'0{2,}', cell.number_format or ''):
        return str(int(value)).zfill(len(cell.number_format))      # leading-zero codes such as style 0008
    return value


# --- .xls ---------------------------------------------------------------------------------------------
def read_xls(content):
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=content, formatting_info=True, on_demand=True)
    except xlrd.biffh.XLRDError as exc:
        text = str(exc).lower()
        if 'encrypt' in text or "can't find workbook" in text:
            raise ValueError('This workbook is password-protected (or not an Excel 97-2003 file). Remove the password in Excel, save and upload again.') from exc
        raise ValueError('This .xls workbook could not be read. Open it in Excel, save it as .xlsx and upload again.') from exc
    if book.nsheets > MAX_SHEETS:
        raise ValueError(f'The workbook has more than {MAX_SHEETS} sheets.')
    sheets = []
    try:
        for index in range(book.nsheets):
            ws = book.sheet_by_index(index)
            b = Builder(ws.name, ws.visibility != 0)
            b.hidden_rows = {r + 1 for r, info in ws.rowinfo_map.items() if info.hidden}
            b.hidden_cols = {c + 1 for c, info in ws.colinfo_map.items() if info.hidden}
            for r in range(ws.nrows):
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    b.put(r + 1, c + 1, _xls_value(book, cell, xlrd), error=cell.ctype == xlrd.XL_CELL_ERROR)
            b.merged = [f'{address(c0 + 1, r0 + 1)}:{address(c1, r1)}' for r0, r1, c0, c1 in ws.merged_cells]
            sheets.append(b.sheet())
            book.unload_sheet(index)
    finally:
        book.release_resources()
    return sheets


def _xls_value(book, cell, xlrd):
    t = cell.ctype
    if t == xlrd.XL_CELL_TEXT:
        return cell.value
    if t == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if t == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, '#ERROR')
    if t == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, book.datemode)
        except (ValueError, OverflowError):
            return cell.value
    value = cell.value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    xf = book.xf_list[cell.xf_index] if cell.xf_index is not None and cell.xf_index < len(book.xf_list) else None
    fmt = book.format_map.get(xf.format_key) if xf else None
    if isinstance(value, int) and fmt and re.fullmatch(r'0{2,}', fmt.format_str or ''):
        return str(value).zfill(len(fmt.format_str))
    return value


# --- .xlsb .ods (calamine) ---------------------------------------------------------------------------
def read_calamine(content):
    import python_calamine as pc
    try:
        book = pc.CalamineWorkbook.from_filelike(BytesIO(content))
    except pc.PasswordError as exc:
        raise ValueError('This workbook is password-protected. Remove the password, save and upload again.') from exc
    except Exception as exc:
        raise ValueError('This workbook could not be opened. Open it in Excel, save it again and upload.') from exc
    sheets = []
    try:
        for meta in book.sheets_metadata[:MAX_SHEETS + 1]:
            if 'work' not in str(meta.typ).lower():
                continue                              # chart sheets and dialog sheets hold no cells
            ws = book.get_sheet_by_name(meta.name)
            if (ws.total_height or 0) * (ws.total_width or 0) > MAX_CELLS * 8:
                raise ValueError(f'{meta.name}: sheet is too large; trim it and upload again.')
            b = Builder(meta.name, 'visible' not in str(meta.visible).lower() or 'hidden' in str(meta.visible).lower())
            for r, row in enumerate(ws.to_python(skip_empty_area=False), 1):
                for c, value in enumerate(row, 1):
                    if value not in ('', None):
                        b.put(r, c, value)
            for rng in ws.merged_cell_ranges or []:
                try:
                    (r0, c0), (r1, c1) = rng
                    b.merged.append(f'{address(c0 + 1, r0 + 1)}:{address(c1 + 1, r1 + 1)}')
                except (TypeError, ValueError):
                    continue
            sheets.append(b.sheet())
    finally:
        book.close()
    return sheets


def read_zip_calamine(content):
    _zip_kind(content)
    return read_calamine(content)


# --- Excel 2003 XML Spreadsheet ---------------------------------------------------------------------
def read_xml2003(content):
    from defusedxml import ElementTree as ET
    try:
        root = ET.fromstring(content)
    except Exception as exc:
        raise ValueError('This Excel 2003 XML file is damaged and could not be read.') from exc
    ss = '{%s}' % SS_NS; x = '{urn:schemas-microsoft-com:office:excel}'
    sheets = []
    for ws in root.iter(ss + 'Worksheet'):
        visible = ws.find(f'{x}WorksheetOptions/{x}Visible')
        b = Builder(ws.get(ss + 'Name') or f'Sheet{len(sheets) + 1}', visible is not None and 'hidden' in (visible.text or '').lower())
        table = ws.find(ss + 'Table')
        if table is None:
            sheets.append(b.sheet()); continue
        col = 0
        for column in table.findall(ss + 'Column'):
            col = int(column.get(ss + 'Index') or col + 1)
            span = int(column.get(ss + 'Span') or 0)
            if column.get(ss + 'Hidden') == '1':
                b.hidden_cols.update(range(col, col + span + 1))
            col += span
        r = 0
        for row in table.findall(ss + 'Row'):
            r = int(row.get(ss + 'Index') or r + 1)
            if row.get(ss + 'Hidden') == '1':
                b.hidden_rows.add(r)
            c = 0
            for cell in row.findall(ss + 'Cell'):
                c = int(cell.get(ss + 'Index') or c + 1)
                across, down = int(cell.get(ss + 'MergeAcross') or 0), int(cell.get(ss + 'MergeDown') or 0)
                data = cell.find(ss + 'Data')
                if data is not None:
                    text = ''.join(data.itertext()); kind = data.get(ss + 'Type', 'String')
                    if kind == 'Number':
                        value = typed(text) if NUMBER.match(text.strip() or 'x') else _float(text)
                    elif kind == 'DateTime':
                        value = _iso(text)
                    elif kind == 'Boolean':
                        value = text.strip() in ('1', 'true', 'TRUE')
                    else:
                        value = text
                    b.put(r, c, value, cell.get(ss + 'Formula'), kind == 'Error')
                if across or down:
                    b.merged.append(f'{address(c, r)}:{address(c + across, r + down)}')
                c += across
        sheets.append(b.sheet())
    if not sheets:
        raise Unsupported('This XML file is not an Excel spreadsheet.')
    return sheets


def _float(text):
    try:
        v = float(text)
        return int(v) if v.is_integer() else v
    except ValueError:
        return text


def _iso(text):
    try:
        return datetime.fromisoformat(text.strip()[:19])
    except ValueError:
        return text


# --- HTML tables saved as .xls / .htm ---------------------------------------------------------------
class _Tables(HTMLParser):
    """Lays HTML out the way Excel does when it opens an HTML 'xls': text blocks and tables go down
    column A onwards in document order; colspan/rowspan become merged cells."""
    BLOCK = {'p', 'div', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'li', 'caption', 'title'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.b = Builder('Sheet1'); self.row = 0; self.skip = 0; self.depth = 0
        self.text = []; self.cell = None; self.taken = set(); self.r = 0; self.c = 0; self.pending = []

    def _flush_text(self):
        t = ' '.join(' '.join(self.text).split())
        if t and self.depth == 0:
            self.row += 1; self.b.put(self.row, 1, t)
        self.text = []

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in ('script', 'style', 'head'):
            self.skip += 1; return
        if tag == 'table':
            if self.depth == 0:
                self._flush_text()
                if self.row:
                    self.row += 1          # one empty row between blocks, as Excel shows it
                self.r = self.row
            self.depth += 1
        elif self.depth == 1 and tag == 'tr':
            self.r += 1; self.c = 0
        elif self.depth == 1 and tag in ('td', 'th'):
            self.c += 1
            while (self.r, self.c) in self.taken:
                self.c += 1
            span_c = max(1, min(_int(a.get('colspan')), 200)); span_r = max(1, min(_int(a.get('rowspan')), 500))
            self.cell = dict(row=self.r, col=self.c, span_c=span_c, span_r=span_r, text=[],
                             num=a.get('x:num') if a.get('x:num') not in (None, '') else a.get('sdval'))
        elif tag == 'br' and self.cell is not None:
            self.cell['text'].append(' ')
        elif tag in self.BLOCK and self.depth == 0:
            self._flush_text()

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'head'):
            self.skip = max(0, self.skip - 1); return
        if tag in ('td', 'th') and self.depth == 1 and self.cell is not None:
            cell = self.cell; self.cell = None
            text = ' '.join(''.join(cell['text']).split())
            # Excel/LibreOffice store the raw number beside the shown text (x:num / sdval). Trust it only
            # when the shown text is itself a number: for dates the raw value is a serial day number.
            shown = re.sub(r'[\s,%₹$€£]', '', text)
            numeric = bool(re.fullmatch(r'[-+]?\d+(\.\d+)?', shown))
            value = _float(cell['num']) if cell['num'] not in (None, '') and numeric else typed(text)
            self.b.put(cell['row'], cell['col'], value)
            for rr in range(cell['row'], cell['row'] + cell['span_r']):
                for cc in range(cell['col'], cell['col'] + cell['span_c']):
                    self.taken.add((rr, cc))
            if cell['span_c'] > 1 or cell['span_r'] > 1:
                self.b.merged.append(f'{address(cell["col"], cell["row"])}:{address(cell["col"] + cell["span_c"] - 1, cell["row"] + cell["span_r"] - 1)}')
            self.c = cell['col'] + cell['span_c'] - 1
        elif tag == 'table' and self.depth:
            self.depth -= 1
            if self.depth == 0:
                self.row = max(self.r, max((r for r, _ in self.taken), default=self.r)); self.taken = set()
        elif tag in self.BLOCK and self.depth == 0:
            self._flush_text()

    def handle_data(self, data):
        if self.skip:
            return
        if self.cell is not None:
            self.cell['text'].append(data)
        elif self.depth == 0:
            self.text.append(data)


def _int(v):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return 1


def read_html(content):
    parser = _Tables()
    parser.feed(_decode(content)); parser.close(); parser._flush_text()
    sheet = parser.b.sheet()
    if not sheet['cells']:
        raise Unsupported('This HTML file has no table data.')
    return [sheet]


# --- CSV / TSV / Unicode text ------------------------------------------------------------------------
def read_csv(content):
    text = _decode(content)
    sample = text[:20000]
    lines = [line for line in sample.splitlines() if line.strip()][:20]
    counts = {d: sum(line.count(d) for line in lines) for d in ('\t', ',', ';', '|')}
    delimiter = max(counts, key=counts.get) if any(counts.values()) else ','
    b = Builder('Sheet1')
    for r, row in enumerate(csv.reader(StringIO(text), delimiter=delimiter), 1):
        if r > MAX_CELLS:
            raise ValueError('The text file has too many rows.')
        for c, value in enumerate(row, 1):
            b.put(r, c, typed(value))
    return [b.sheet()]


# --- JSON ---------------------------------------------------------------------------------------------
def read_json(content):
    try:
        data = json.loads(_decode(content))
    except ValueError as exc:
        raise ValueError(f'This JSON file is not valid JSON ({str(exc)[:80]}).') from exc
    _depth(data)
    if isinstance(data, dict) and isinstance(data.get('sheets'), list):
        return [_json_sheet(s) for s in data['sheets'][:MAX_SHEETS + 1] if isinstance(s, dict)]
    if isinstance(data, dict) and data and all(isinstance(v, list) and v and all(isinstance(x, (dict, list)) for x in v) for v in data.values()):
        return [_grid_from(name, rows) for name, rows in list(data.items())[:MAX_SHEETS + 1]]
    if isinstance(data, list):
        return [_grid_from('Data', data)]
    if isinstance(data, dict):
        b = Builder('Data'); _write_document(b, data, [0]); return [b.sheet()]
    raise Unsupported('This JSON file holds a single value, not order data.')


def _depth(value, level=0):
    if level > MAX_JSON_DEPTH:
        raise ValueError('This JSON file is nested too deeply.')
    if isinstance(value, dict):
        for v in value.values():
            _depth(v, level + 1)
    elif isinstance(value, list):
        for v in value:
            _depth(v, level + 1)


def _scalar(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)[:500]
    return v


def _json_sheet(spec):
    b = Builder(spec.get('name') or 'Sheet', bool(spec.get('hidden')))
    if isinstance(spec.get('rows'), list):
        for r, row in enumerate(spec['rows'], 1):
            for c, v in enumerate(row if isinstance(row, list) else [row], 1):
                b.put(r, c, _scalar(v))
    for cell in spec.get('cells') or []:
        if isinstance(cell, dict):
            row, col = cell.get('row'), cell.get('col')
            if (row is None or col is None) and isinstance(cell.get('cell'), str):
                m = re.fullmatch(r'([A-Za-z]{1,3})(\d{1,7})', cell['cell'].strip())
                if m:
                    col = 0
                    for ch in m.group(1).upper():
                        col = col * 26 + ord(ch) - 64
                    row = int(m.group(2))
            if isinstance(row, int) and isinstance(col, int) and row > 0 and col > 0:
                b.put(row, col, _scalar(cell.get('value')))
    return b.sheet()


def _grid_from(name, rows):
    b = Builder(name)
    if rows and all(isinstance(x, dict) for x in rows):
        keys = list(dict.fromkeys(k for x in rows for k in x))
        for c, k in enumerate(keys, 1):
            b.put(1, c, str(k))
        for r, x in enumerate(rows, 2):
            for c, k in enumerate(keys, 1):
                b.put(r, c, _scalar(x.get(k)))
    else:
        for r, row in enumerate(rows, 1):
            for c, v in enumerate(row if isinstance(row, list) else [row], 1):
                b.put(r, c, _scalar(v))
    return b.sheet()


def _write_document(b, doc, cursor, title=None):
    """A single nested JSON document laid out like a form: 'key | value' rows, then each list of
    records as a titled table — the same shapes the Excel detectors read."""
    if title:
        cursor[0] += 1; b.put(cursor[0], 1, title)
    scalars = [(k, v) for k, v in doc.items() if not isinstance(v, (dict, list))]
    for k, v in scalars:
        cursor[0] += 1; b.put(cursor[0], 1, str(k)); b.put(cursor[0], 2, v)
    for k, v in doc.items():
        if isinstance(v, dict):
            cursor[0] += 1; _write_document(b, v, cursor, str(k))
        elif isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
            cursor[0] += 2; b.put(cursor[0], 1, str(k))
            keys = list(dict.fromkeys(kk for x in v for kk in x))
            cursor[0] += 1
            for c, kk in enumerate(keys, 1):
                b.put(cursor[0], c, str(kk))
            for x in v:
                cursor[0] += 1
                for c, kk in enumerate(keys, 1):
                    b.put(cursor[0], c, _scalar(x.get(kk)))
        elif isinstance(v, list):
            cursor[0] += 1; b.put(cursor[0], 1, str(k))
            for c, x in enumerate(v, 2):
                b.put(cursor[0], c, _scalar(x))


READERS = {'ooxml': read_ooxml, 'xls': read_xls, 'xlsb': read_zip_calamine, 'ods': read_zip_calamine,
           'xml2003': read_xml2003, 'html': read_html, 'csv': read_csv, 'json': read_json}
