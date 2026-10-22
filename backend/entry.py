"""Order entry engine: one buyer Excel -> rows of the client's own Oracle tables.

The main table is ORDERENTRY (setting ENTRY_TABLE): one row per size, the order details repeating on every row,
exactly like the ERP's own order entry. More tables can be linked to it. A linked table gets one row per order,
per size, or per row of an Excel table (for example the FABRIC REQUIREMENT block), and carries the order's key
columns (for example IONO) so its rows belong to the same order.

How a column gets its value, first match wins:
  1. the person's choice for this upload: an Excel cell, a typed value, "leave blank" or "not applicable"
  2. for a linked table: the same column of the main table (IONO, STYLENO…)
  3. the size table (colour, size, quantity…) or the column of the linked Excel table
  4. the structural reader: buyer, style, OC number… each with its Excel cell
  5. an Excel label that means this column: built-in garment vocabulary, the column's Oracle comment and labels
     learned from the person's earlier mappings
Nothing is invented. A column without a value stays NULL, and that is never an error unless Oracle requires it.
One Excel cell fills one column: when two columns claim the same cell, the one whose labels name the Excel label
beside the cell keeps it.

Rows are written in one transaction, read back by ROWID inside it, and read back again on a fresh connection after
the commit. The client's tables get no extra columns: the receipt in PP_IMPORT_LOG keeps each row's ROWID."""
from __future__ import annotations
from datetime import datetime
from decimal import Decimal
import difflib
import json
import re
import oracledb
from openpyxl.utils.cell import range_boundaries
from .config import OWNER, setting
from .mapping import SIZE_TOKEN, anchor
from .oracle_store import IDENT, bind_value, blank, connect, convert, digest, normalize, plain, read_result, same
from .vocabulary import GLOBAL_LABELS, ORDER_FIELDS, ROW_FIELDS, condensed, group_of, label_set, meaning
from .workbooks import norm, source, text

MAIN = setting('ENTRY_TABLE', 'ORDERENTRY').upper()
MAX_ROWS = 5000
ISO_DATE = re.compile(r'^\d{4}-\d{2}-\d{2}$')
ROW_KINDS = {'order': 'One row per order', 'sizes': 'One row per size', 'lines': 'One row per colour line', 'section': 'One row per row of an Excel table'}
FIXED = ('cell', 'text', 'blank', 'na', 'field', 'link', 'auto_oc')   # same value on every row of the order
PER_ROW = ('row', 'header')                                 # a value per row


# ---------------------------------------------------------------- Oracle metadata

def columns(conn, table):
    """Live columns of one table: type, size, nullability, default, identity, comment and what the column means."""
    c = conn.cursor()
    c.execute('''select t.column_name, t.data_type, t.data_length, t.char_length, t.char_used, t.data_precision, t.data_scale, t.nullable,
            t.data_default, t.identity_column, t.virtual_column, m.comments
        from all_tab_cols t left join all_col_comments m on m.owner = t.owner and m.table_name = t.table_name and m.column_name = t.column_name
        where t.owner = :o and t.table_name = :t and t.hidden_column = 'NO' order by t.column_id''', o=OWNER, t=table)
    out = {}
    for name, kind, length, chars, used, precision, scale, nullable, default, identity, virtual, comment in c:
        if virtual == 'YES' or not IDENT.match(name):
            continue
        m = meaning(name, comment)
        out[name] = dict(type=kind, length=length, char_length=chars, char_used=used, precision=precision, scale=scale, nullable=nullable,
                         default=None if default is None else str(default).strip() or None, identity=identity == 'YES', comment=comment,
                         label=m['label'], level=m['level'], field=m['field'], labels=m['labels'], system=m['system'])
    if not out:
        raise KeyError(f'Table {table} not found in {OWNER}')
    return out


def keys(conn, table):
    c = conn.cursor()
    c.execute('''select c.constraint_type, c.constraint_name, cc.column_name from all_constraints c join all_cons_columns cc
        on cc.owner = c.owner and cc.constraint_name = c.constraint_name where c.owner = :o and c.table_name = :t and c.constraint_type in ('P','U')
        order by c.constraint_name, cc.position''', o=OWNER, t=table)
    out = dict(primary=[], unique={})
    for kind, name, col in c:
        if kind == 'P':
            out['primary'].append(col)
        else:
            out['unique'].setdefault(name, []).append(col)
    return out


def privileges(conn, table):
    c = conn.cursor()
    c.execute('''select distinct privilege from all_tab_privs where table_schema = :o and table_name = :t
        and (grantee in (user, 'PUBLIC') or grantee in (select role from session_roles))''', o=OWNER, t=table)
    return {r[0] for r in c}


def exists(conn, table=None):
    c = conn.cursor()
    c.execute('select count(*) from all_tables where owner = :o and table_name = :t', o=OWNER, t=table or MAIN)
    return c.fetchone()[0] > 0


def keygen(cols, pk):
    """A single numeric primary key without a default is numbered by POPPYS when saved (PP_RECORD_SEQ)."""
    if len(pk) != 1:
        return None
    m = cols.get(pk[0])
    return pk[0] if m and m['type'] == 'NUMBER' and not m['default'] and not m['identity'] else None


TABLE_LABELS = {'ORDERENTRY': 'Order entry', 'ORDMAS': 'Order header', 'ORDSTYLE': 'Order lines', 'ORDSIZE': 'Size quantities', 'BUYMAS': 'Buyers',
                'STYLENUM': 'Style numbers', 'STYLEMAS': 'Garment descriptions', 'SIZE_MAS': 'Sizes', 'COLORMAS': 'Colours', 'UNITMAS': 'Units',
                'MERCHMAS': 'Merchandisers', 'PARTYMAS': 'Parties', 'ORDSHIP': 'Shipments', 'ORDFAB': 'Order fabrics', 'ORDYARN': 'Order yarns',
                'ORDACC': 'Order accessories', 'ORDCOMP': 'Order components', 'ORDDOC': 'Order documents', 'ORDCMTFAB': 'CMT fabrics',
                'EXCEL_ORDERS': 'Excel orders (earlier POPPYS table)', 'FABRICS': 'Fabrics', 'FINANCE': 'Finance'}


def meaning_of_table(name):
    return TABLE_LABELS.get(name) or name.replace('_', ' ').title()


def describe(conn, table):
    cols = columns(conn, table); k = keys(conn, table); p = privileges(conn, table)
    return dict(table=table, columns=cols, primary=k['primary'], keygen=keygen(cols, k['primary']),
                can_insert='INSERT' in p, can_update='UPDATE' in p, can_select='SELECT' in p)


# ---------------------------------------------------------------- workbook helpers

class Book:
    """Indexed workbook: cells by (row, col), merged cells pointing at their top-left cell, header limits."""

    def __init__(self, sheets):
        self.sheets = [s for s in sheets if not s['hidden']]
        self.by_name = {s['name']: s for s in sheets}
        self.cells, self.merged, self.limit, self.address = {}, {}, {}, {}
        for s in sheets:
            cells = {(c['row'], c['col']): c for c in s['cells']}
            merged = {}
            for rng in s.get('merged', []):
                try:
                    c1, r1, c2, r2 = range_boundaries(rng)
                except (ValueError, TypeError):
                    continue
                if (r2 - r1 + 1) * (c2 - c1 + 1) > 5000:
                    continue
                for r in range(r1, r2 + 1):
                    for col in range(c1, c2 + 1):
                        if (r, col) != (r1, c1):
                            merged[(r, col)] = (r1, c1)
            self.cells[s['name']], self.merged[s['name']] = cells, merged
            self.address[s['name']] = {c['cell']: c for c in s['cells']}
            heads = [c['row'] for c in s['cells'] if norm(c['value']) in ('size', 'sizes', 'garment')]
            self.limit[s['name']] = min(heads) if heads else 60

    def cell(self, sheet, row, col, fill=True):
        cells = self.cells.get(sheet, {})
        if (row, col) in cells:
            return cells[(row, col)]
        top = self.merged.get(sheet, {}).get((row, col)) if fill else None
        return cells.get(top) if top else None

    def at(self, sheet, address):
        return self.address.get(sheet, {}).get(address)


def cell_source(sheet, cell, via=None):
    out = source(sheet, cell['cell'], cell['value'], cell.get('formula'))
    if via:
        out['via'] = via
    return out


def shown(value):
    return value.isoformat()[:10] if hasattr(value, 'isoformat') else value


def is_label(cell, known):
    return cell is not None and isinstance(cell['value'], str) and norm(cell['value']) in known


def value_beside(cells, c, known):
    """The value of a label cell: the first cell to its right (within 3 columns) unless that is another label,
    otherwise the cell below (header rows such as 'OCN | QTY(IN PACKS) | QTY(IN PIECES)' with values underneath)."""
    right = next((cells[(c['row'], x)] for x in range(c['col'] + 1, c['col'] + 4) if (c['row'], x) in cells), None)
    cand = right if right is not None and not is_label(right, known) else cells.get((c['row'] + 1, c['col']))
    if cand is None or blank(cand['value']) or is_label(cand, known) or cand.get('error'):
        return None
    return cand


def label_matches(book, wanted, empty=None):
    """{COLUMN: (cell, sheet, label text)} for columns whose Excel label is in the header area (above the size table).
    'Label: value' in one cell also counts. `empty` collects columns whose label is there but whose value cell is empty."""
    known = GLOBAL_LABELS.union(*wanted.values()) if wanted else set(GLOBAL_LABELS)
    out = {}
    for s in book.sheets:
        cells = book.cells[s['name']]; limit = book.limit[s['name']]
        for c in sorted(s['cells'], key=lambda x: (x['row'], x['col'])):
            if c['row'] >= limit or not isinstance(c['value'], str) or len(c['value']) > 80:
                continue
            txt = c['value']; head, sep, inline = txt.partition(':')
            for col, syns in wanted.items():
                if col in out:
                    continue
                if sep and inline.strip() and norm(head) in syns:
                    out[col] = (dict(c, value=inline.strip(), inline=True), s['name'], head.strip()); continue
                if norm(txt) not in syns:
                    continue
                cand = value_beside(cells, c, known)
                if cand is None:
                    if empty is not None:
                        empty.setdefault(col, dict(sheet=s['name'], cell=c['cell'], label=txt.strip()))
                    continue
                out[col] = (cand, s['name'], txt.strip())
    return out


def excel_pairs(book, known=None):
    """Label → value pairs of the header area (what a person would read as 'the order details')."""
    known = known or GLOBAL_LABELS
    out, value_cells = [], set()
    for s in book.sheets:
        cells = book.cells[s['name']]; limit = book.limit[s['name']]
        heads = sorted((c for c in s['cells'] if c['row'] < limit and isinstance(c['value'], str) and 0 < len(c['value'].strip()) <= 60),
                       key=lambda x: (x['row'], x['col']))
        for c in heads:  # known labels first: they decide which cells are values
            head, sep, inline = c['value'].partition(':')
            if sep and inline.strip() and len(head.strip()) <= 40:
                out.append(dict(label=head.strip(), value=inline.strip(), sheet=s['name'], cell=c['cell'], label_cell=c['cell'], inline=True)); continue
            if norm(c['value']) in known:
                v = value_beside(cells, c, known)
                if v is not None:
                    value_cells.add((s['name'], v['cell']))
                    out.append(dict(label=c['value'].strip(), value=shown(v['value']), sheet=s['name'], cell=v['cell'], label_cell=c['cell'], formula=v.get('formula')))
        for c in heads:  # other text followed by a number or a date on its right
            if norm(c['value']) in known or (s['name'], c['cell']) in value_cells or ':' in c['value']:
                continue
            right = next((cells[(c['row'], x)] for x in range(c['col'] + 1, c['col'] + 3) if (c['row'], x) in cells), None)
            looks = right is not None and not blank(right['value']) and (not isinstance(right['value'], str) or ISO_DATE.match(right['value']))
            if looks and (s['name'], right['cell']) not in value_cells:
                value_cells.add((s['name'], right['cell']))
                out.append(dict(label=c['value'].strip(), value=shown(right['value']), sheet=s['name'], cell=right['cell'], label_cell=c['cell'], formula=right.get('formula')))
    return out


def header_score(header, names):
    """How well an Excel table header names a column: 1 same words, ~0.9 near spelling, 0.8 most words shared, else 0."""
    h = condensed(header)
    if len(h) < 2:
        return 0
    hw = {w for w in norm(header).split() if w}
    best = 0
    for n in names:
        cn = condensed(n)
        if not cn:
            continue
        if h == cn:
            return 1.0
        if min(len(h), len(cn)) >= 5:
            r = difflib.SequenceMatcher(None, h, cn).ratio()
            if r >= 0.88:
                best = max(best, round(r, 2))
        nw = {w for w in norm(n).split() if w}
        if hw and nw and (nw <= hw or hw <= nw) and any(len(w) >= 3 for w in hw & nw):
            if len(hw & nw) / max(len(hw), len(nw)) >= 0.5:
                best = max(best, 0.8)
    return best


def find_sections(sheets, orders=None):
    """Excel tables in a workbook: a header row of 3+ labels with data rows under it, until an empty row.
    Titled tables have a short title row just above the header (FABRIC REQUIREMENT, CARE LABEL…); untitled ones
    (the size blocks) start after an empty row below the order header area."""
    out = []
    book = Book(sheets)
    for s in book.sheets:
        rows = {}
        for c in s['cells']:
            rows.setdefault(c['row'], []).append(c)
        for r in rows:
            rows[r].sort(key=lambda x: x['col'])
        limit = book.limit[s['name']]
        free = limit == 60 and not any(norm(c['value']) in ('size', 'sizes', 'garment') for c in s['cells'])

        def strings(r):
            return [c for c in rows.get(r, []) if isinstance(c['value'], str) and c['value'].strip() and not ISO_DATE.match(c['value'])]

        def header(r):
            cells = rows.get(r, []); st = strings(r)
            return len(st) >= 3 and len(st) >= 0.6 * len(cells) and (r + 1) in rows

        def title(r):
            cells = rows.get(r, [])
            return r > 1 and 1 <= len(cells) <= 2 and isinstance(cells[0]['value'], str) and cells[0]['value'].strip() and header(r + 1)

        taken = set()
        for r in sorted(rows):
            if r in taken:
                continue
            ttl = None
            if title(r):
                ttl, hr = rows[r][0], r + 1
            elif header(r) and (r - 1) not in rows and (r >= limit or free):
                hr = r
            else:
                continue
            end = hr; rr = hr + 1
            while rr in rows and not title(rr) and not (header(rr) and (rr - 1) not in rows):
                end = rr; rr += 1
            taken.update(range(r, end + 1))
            heads = [c for c in strings(hr)]
            cols = [c['col'] for c in rows[hr]]
            for rr2 in range(hr + 1, end + 1):
                cols += [c['col'] for c in rows.get(rr2, [])]
            first, last = min(cols), max(cols)
            data = [x for x in range(hr + 1, end + 1) if any(not blank(c['value']) for c in rows.get(x, []))
                    and norm(rows[x][0]['value']) not in ('total', 'grand total')]
            sid = f"{s['name']}!{ttl['cell']}" if ttl else f"{s['name']}!R{hr}"
            sizes = sum(1 for c in heads if SIZE_TOKEN.match(text(c['value'])) or norm(c['value']) in ('first size', 'total'))
            prev = out[-1] if out and out[-1]['sheet'] == s['name'] and 0 < hr - out[-1]['end_row'] <= 3 and not out[-1]['title'].startswith(('Size block', 'Table at')) else None
            if ttl:
                name = text(ttl['value'])
            elif any(norm(c['value']) == 'contract' for c in heads) or sizes >= 3:
                name = f'Size block (rows {hr}–{end})'
            elif prev:
                name = f"{prev['title'].split(' (continued')[0]} (continued, rows {hr}–{end})"
            else:
                name = f'Table at row {hr}'
            out.append(dict(id=sid, sheet=s['name'], title=name, titled=bool(ttl), title_cell=ttl['cell'] if ttl else None, title_row=r if ttl else None,
                            header_row=hr, first_col=first, last_col=last, end_row=end,
                            headers=[dict(col=h['col'], cell=h['cell'], text=text(h['value'])) for h in heads], rows=len(data)))
    return out


def section_rows(book, sec):
    """Data rows of an Excel table: non-empty rows below the header, without TOTAL rows. Merged cells read as Excel shows them."""
    out = []
    for r in range(sec['header_row'] + 1, sec['end_row'] + 1):
        own = [c for c in (book.cell(sec['sheet'], r, col, fill=False) for col in range(sec['first_col'], sec['last_col'] + 1))
               if c is not None and not blank(c['value'])]
        if not own:
            continue
        if norm(own[0]['value']) in ('total', 'grand total'):
            continue
        out.append((r, {col: book.cell(sec['sheet'], r, col) for col in range(sec['first_col'], sec['last_col'] + 1)}))
    return out


# ---------------------------------------------------------------- targets and bindings

def resolve_kind(target, cols):
    kind = target.get('rows') or 'auto'
    if kind in ROW_KINDS:
        return kind
    fields = {m['field'] for m in cols.values()}
    return 'sizes' if {'size', 'quantity'} <= fields else 'lines' if 'color' in fields else 'order'


def default_targets():
    return [dict(table=MAIN, rows='auto')]


def row_fields(kind):
    if kind == 'sizes':
        return set(ROW_FIELDS)
    if kind == 'lines':
        return {k for k, (lvl, _) in ROW_FIELDS.items() if lvl == 'line'}
    return set()


def factor_of(order):
    try:
        f = int(order.get('pieces_per_unit') or 1)
        return f if f >= 1 else 1
    except (TypeError, ValueError):
        return 1


def line_total(line, factor=1):
    qty = [s.get('quantity') for s in line.get('sizes', []) if not blank(s.get('quantity'))]
    nums = [q for q in qty if isinstance(q, (int, float)) and not isinstance(q, bool)]
    if not nums or len(nums) != len(qty):
        return None
    total = sum(nums) * factor
    return int(total) if float(total).is_integer() else total


def contexts(kind, order, book, section=None):
    """The rows one order produces for one table, each with a stable key."""
    factor = factor_of(order)
    if kind == 'section':
        if not section:
            return []
        return [dict(key=f'r{r}', excel_row=r, cells=vals, section=section, label=f'Excel row {r}') for r, vals in section_rows(book, section)]
    if kind in ('sizes', 'lines'):
        out = []
        for li, line in enumerate(order.get('lines', [])):
            total = line_total(line, factor)
            if kind == 'lines':
                out.append(dict(key=f"l{line.get('id', li)}", li=li, line=line, total=total, factor=factor, label=line.get('color') or f'Line {li + 1}'))
                continue
            for si, s in enumerate(line.get('sizes', [])):
                if blank(s.get('quantity')):
                    continue
                out.append(dict(key=f"s{line.get('id', li)}:{s.get('id', si)}", li=li, line=line, si=si, size=s, total=total, factor=factor,
                                label=f"{line.get('color') or 'Line ' + str(li + 1)} · {text(s.get('size'))}"))
        if out:
            return out
    return [dict(key='order', label='Order details')]


def wanted_labels(cols, learned):
    out = {}
    for name, m in cols.items():
        if m['system'] or m['identity']:
            continue
        labels = set(label_set(m['labels'])) | label_set((learned or {}).get(name, []))
        if labels:
            out[name] = labels
    return out


def auto_bindings(cols, kind, order, book, section=None, main=None, learned=None, empty=None):
    """Automatic source of every column that has one: {COLUMN: binding} with binding['auto'] True.
    `main` = (main table name, its columns, its chosen bindings) when this is a linked table."""
    out = {}
    usable = {n: m for n, m in cols.items() if not m['system'] and not m['identity']}
    if main:
        mtable, mcols, mchosen = main
        for name, m in usable.items():
            link = name if name in mcols and (mchosen.get(name) or {}).get('kind') in FIXED else next(
                (c for c, mm in mcols.items() if m['field'] in ORDER_FIELDS and mm['field'] == m['field'] and (mchosen.get(c) or {}).get('kind') in FIXED), None)
            if link:
                out[name] = dict(kind='link', column=link, auto=True, via=f'Same as {mtable}.{link} of this order')
    fields = row_fields(kind)
    for name, m in usable.items():
        if name in out:
            continue
        if m['field'] in fields:
            out[name] = dict(kind='row', field=m['field'], auto=True, via=ROW_FIELDS[m['field']][1])
    if kind == 'section' and section:
        taken = set()
        scored = sorted(((header_score(h['text'], [name.replace('_', ' '), m['label']] + m['labels'] + list((learned or {}).get(name, []))), name, h['col'], h)
                         for name, m in usable.items() if name not in out for h in section['headers']), key=lambda x: (-x[0], x[1], x[2]))
        for score, name, col, h in scored:
            if score < 0.8 or name in out or col in taken:
                continue
            out[name] = dict(kind='header', col=col, auto=True, score=score, via=f'Excel column "{h["text"]}"' + ('' if score == 1 else ' (similar name — please check)'))
            taken.add(col)
        return out  # a linked Excel table's columns come from its own headers, not from the order header area
    for name, m in usable.items():
        if name not in out and m['field'] in ORDER_FIELDS and not blank(order.get(m['field'])):
            out[name] = dict(kind='field', field=m['field'], auto=True, via=f"Read from Excel as {ORDER_FIELDS[m['field']]}")
    want = {n: v for n, v in wanted_labels(usable, learned).items() if n not in out}
    for name, (cell, sheet, label) in label_matches(book, want, empty).items():
        b = dict(kind='cell', sheet=sheet, cell=cell['cell'], auto=True, via=f'Excel label "{label}"')
        if cell.get('inline'):
            b['inline'] = cell['value']
        out[name] = b
    # One Excel cell fills one column: keep the column whose labels name the label beside the cell.
    by_cell = {}
    for name, b in out.items():
        where = cell_of(b, order)
        if where:
            by_cell.setdefault(where, []).append(name)
    for (sheet, address), names in by_cell.items():
        if len(names) < 2:
            continue
        s = book.by_name.get(sheet); c = book.at(sheet, address)
        near = anchor(s, c) if s and c else None
        exact = [n for n in names if near and near['text'] in label_set(usable[n]['labels'] + list((learned or {}).get(n, [])))]
        keep = exact[:1] or names[:1]
        for n in names:
            if n not in keep:
                out.pop(n)
    return out


def cell_of(b, order):
    if b['kind'] == 'cell' and not b.get('inline'):
        return (b['sheet'], b['cell'])
    if b['kind'] == 'field':
        p = (order.get('provenance') or {}).get(b['field']) or {}
        if p.get('cell') and not p.get('derivation'):
            return (p.get('sheet'), p['cell'])
    return None


# ---------------------------------------------------------------- values

def bound_value(b, order, book, ctx, main_values):
    """(value, source, origin) of one binding for one row; value None means blank."""
    kind = b['kind']
    if kind == 'cell':
        c = book.at(b['sheet'], b['cell'])
        if b.get('inline') is not None:
            return b['inline'], dict(source(b['sheet'], b['cell'], c['value'] if c else None, (c or {}).get('formula')), via=b.get('via')), 'excel'
        if c is None:
            return None, dict(sheet=b['sheet'], cell=b['cell'], via=b.get('via'), missing=True), 'excel'
        return shown(c['value']), cell_source(b['sheet'], c, b.get('via')), 'ai' if b.get('ai') and not b.get('confirmed') else 'excel'
    if kind == 'text':
        return b.get('value'), {'origin': 'Typed by you'}, 'manual'
    if kind == 'auto_oc':
        return b.get('value'), {'origin': 'Next number in your OC series'}, 'generated'
    if kind in ('blank', 'na'):
        return None, {'origin': 'Not applicable' if kind == 'na' else 'Left blank by you'}, 'manual'
    if kind == 'field':
        src = dict((order.get('provenance') or {}).get(b['field']) or {})
        if b.get('via'):
            src['via'] = b['via']
        value = order.get(b['field'])
        return (None if blank(value) else value), src, 'excel' if src.get('cell') else 'manual'
    if kind == 'link':
        v = main_values.get(b['column'])
        return (v[0] if v else None), dict((v[1] if v else {}), via=b.get('via')), 'link'
    if kind == 'row':
        f = b['field']; line = ctx.get('line') or {}; s = ctx.get('size') or {}
        if f == 'color':
            if line.get('typed_color'):
                return (text(line.get('color')) or None), {'origin': 'Typed by you', 'excel_value': line.get('excel_color')}, 'manual'
            return (text(line.get('color')) or None), (line.get('provenance') or {}).get('color') or {}, 'excel'
        if f == 'label':
            return (text(line.get('label')) or None), (line.get('provenance') or {}).get('label') or {}, 'excel'
        if f == 'line_no':
            return (ctx['li'] + 1 if 'li' in ctx else None), {'origin': 'Line number'}, 'derived'
        if f == 'line_total':
            return ctx.get('total'), {'origin': "Sum of this colour's sizes"}, 'derived'
        if f == 'size':
            if s.get('typed_size'):
                return (text(s.get('size')) or None), {'origin': 'Typed by you', 'excel_value': s.get('excel_size')}, 'manual'
            return (text(s.get('size')) or None), s.get('size_source') or {}, 'excel'
        if f == 'quantity':
            q = s.get('quantity')
            if blank(q):
                return None, s.get('source') or {}, 'excel'
            if s.get('typed_quantity'):
                return q, {'origin': 'Typed by you', 'excel_value': s.get('excel_quantity')}, 'manual'
            factor = ctx.get('factor') or 1
            if factor != 1 and isinstance(q, (int, float)) and not isinstance(q, bool):
                q = q * factor
            return q, dict(s.get('source') or {}, **({'via': f'× {factor} pieces per unit'} if factor != 1 else {})), 'excel'
        if f == 'size_no':
            return (ctx['si'] + 1 if 'si' in ctx else None), {'origin': 'Size position'}, 'derived'
        return None, {}, 'derived'
    if kind == 'header':
        c = (ctx.get('cells') or {}).get(b['col'])
        if c is None or blank(c['value']):
            return None, {}, 'excel'
        return shown(c['value']), cell_source(ctx['section']['sheet'], c, b.get('via')), 'excel'
    return None, {}, 'manual'


def friendly_error(meta, exc):
    msg = str(exc)
    if msg.startswith('longer than'):
        size = meta['char_length'] if meta.get('char_used') == 'C' else meta['length']
        return f'Too long for Oracle: at most {size} characters'
    return msg[:1].upper() + msg[1:]


def column_view(name, m, meta):
    size = m['char_length'] if m.get('char_used') == 'C' else m['length']
    return dict(name=name, type=m['type'], length=size, precision=m['precision'], scale=m['scale'], nullable=m['nullable'], default=m['default'],
                required=m['nullable'] == 'N' and not m['default'] and not m['identity'] and name != meta['keygen'],
                label=m['label'], system=bool(m['system'] or m['identity'] or name == meta['keygen']), level=m['level'])


# ---------------------------------------------------------------- plan

def build_plan(draft, metas, existing_refs=frozenset(), committed_file=None, learned=None, auto_oc=None):
    """Exact rows, column by column, that one approval writes. Pure: Oracle facts come in as arguments."""
    orders = draft.get('orders') or []
    targets = draft.get('targets') or default_targets()
    book = Book(draft['sheets'])
    sections = {s['id']: s for s in (draft.get('sections') or find_sections(draft['sheets'], orders))}
    main = targets[0]['table']
    errors, warnings, info, issues = [], [], [], []
    ops, views, tables = [], [], {}
    if not orders:
        errors.append('No order was read from this Excel. Open Edit mapping and click the cells that hold the order details.')
    if len(orders) > 50:
        errors.append('One upload can hold at most 50 orders.')
    for ti, t in enumerate(targets):
        meta = metas.get(t['table'])
        if meta is None:
            errors.append(f'Table {t["table"]} was not found in Oracle. Remove it from the mapping or create it first.')
            continue
        kind = resolve_kind(t, meta['columns'])
        sec = sections.get((t.get('section') or {}).get('id')) or (t.get('section') if (t.get('section') or {}).get('header_row') else None)
        tables[t['table']] = dict(main=ti == 0, rows_kind=kind, rows_label=ROW_KINDS[kind], link=t.get('link') or {},
                                  section=sec and {k: sec.get(k) for k in ('id', 'title', 'sheet', 'header_row', 'end_row', 'first_col', 'last_col', 'headers', 'rows')},
                                  can_insert=meta['can_insert'], can_update=meta['can_update'], keygen=meta['keygen'], primary=meta['primary'],
                                  columns=[column_view(n, m, meta) for n, m in meta['columns'].items()])
        if not meta['can_insert']:
            errors.append(f'POPPYS is not allowed to add rows to {t["table"]} yet. On this computer run: python scripts/upgrade_local.py')
        if kind == 'section' and sec is None:
            errors.append(f'{t["table"]}: choose the Excel table its rows come from (Edit mapping → {t["table"]}).')
    labelled_blank, used_cells, seen_refs = {}, set(), set()
    for oi, o in enumerate(orders):
        prefix = f'Order {oi + 1}: ' if len(orders) > 1 else ''
        view = dict(index=oi, tables={})
        main_values, main_chosen = {}, {}
        for ti, t in enumerate(targets):
            table = t['table']; meta = metas.get(table)
            if meta is None or (ti > 0 and main not in metas):
                continue
            cols = meta['columns']; kind = tables[table]['rows_kind']
            sec = sections.get((t.get('section') or {}).get('id')) or (t.get('section') if (t.get('section') or {}).get('header_row') else None)
            empty = {} if ti == 0 else None
            chosen = auto_bindings(cols, kind, o, book, sec, None if ti == 0 else (main, metas[main]['columns'], main_chosen), (learned or {}).get(table), empty)
            for col, linked in (t.get('link') or {}).items():
                if col in cols:
                    chosen[col] = dict(kind='link', column=linked, auto=False, via=f'Same as {main}.{linked} of this order')
            explicit = (o.get('bind') or {}).get(table) or {}
            if ti == 0 and not explicit and o.get('table_values'):   # drafts saved by the earlier single-table screen
                explicit = {c: (dict(kind='text', value=v.get('value')) if v.get('value') is not None else dict(kind='blank')) for c, v in o['table_values'].items()}
            for col, b in explicit.items():
                if col in cols and isinstance(b, dict) and b.get('kind'):
                    chosen[col] = dict(b, auto=False)
            if ti == 0 and 'IONO' in cols and (auto_oc or {}).get(oi) and (chosen.get('IONO') is None or chosen['IONO']['kind'] == 'auto_oc'
                                                                           or (chosen['IONO']['kind'] == 'field' and blank(o.get('order_ref')))):
                chosen['IONO'] = dict(kind='auto_oc', value=auto_oc[oi], auto=True, via='Next number in your OC series (given by POPPYS when saved)')
            if ti == 0:
                main_chosen = chosen
                labelled_blank.update({c: w for c, w in (empty or {}).items() if c not in labelled_blank})
            overrides = (o.get('overrides') or {}).get(table) or {}
            rows = []
            for ctx in contexts(kind, o, book, sec):
                values, infos = {}, {}
                for name, m in cols.items():
                    b = chosen.get(name); over = overrides.get(ctx['key'], {})
                    if name in over:
                        raw, src, origin = (None if blank(over[name]) else over[name]), {'origin': 'Typed by you'}, 'manual'
                    elif b is None:
                        continue
                    else:
                        raw, src, origin = bound_value(b, o, book, ctx, main_values)
                    if raw is None or blank(raw):
                        if b is not None and b['kind'] in ('blank', 'na'):
                            infos[name] = dict(origin='manual', source=src)
                        continue
                    if b is not None and b['kind'] in ('cell', 'field') and src.get('cell'):
                        used_cells.add((src.get('sheet'), src['cell']))
                    try:
                        values[name] = convert(m, raw)
                    except ValueError as exc:
                        msg = friendly_error(m, exc)
                        issues.append(dict(table=table, column=name, row=ctx['key'] if (b or {}).get('kind') in PER_ROW or name in over else None,
                                           order=oi, level='error', message=f'{msg} (value "{raw}")'))
                        errors.append(f'{prefix}{table}.{name} ({m["label"]}): {msg} — value "{raw}".')
                        continue
                    infos[name] = dict(origin=origin, source=src, raw=raw, label=m['label'], via=(b or {}).get('via'))
                if meta['keygen']:
                    values[meta['keygen']] = '$key'
                    infos[meta['keygen']] = dict(origin='generated', note='New number given when saved')
                rows.append(dict(key=ctx['key'], label=ctx.get('label'), values=values, infos=infos))
            per_row_cols = {n for n, b in chosen.items() if b and b['kind'] in PER_ROW}
            if kind == 'section':   # a row of an Excel table counts only when one of its own columns has a value
                rows = [r for r in rows if per_row_cols & set(r['values'])]
            real = [r for r in rows if set(r['values']) - {meta['keygen']}]
            if ti == 0:
                first = real[0] if real else dict(values={}, infos={})
                for name, b in chosen.items():
                    if b and b['kind'] in FIXED and name in first['values']:
                        main_values[name] = (first['values'][name], first['infos'][name].get('source') or {})
                if not real:
                    errors.append(f'{prefix}nothing to save in {table}: no column has a value. Click Edit mapping and choose the order cells.')
            elif not real and (kind != 'section' or sec):
                warnings.append(f'{prefix}{table}: no rows with values, so nothing will be added to it.')
            rows = real
            if len(rows) > MAX_ROWS:
                errors.append(f'{prefix}{table}: {len(rows):,} rows is more than one upload can hold ({MAX_ROWS:,}).')
            for r in rows:
                ops.append(dict(key=f'o{oi}t{ti}{r["key"]}', order=oi, table=table, row=r['key'], action='insert', label=r['label'],
                                values=r['values'], columns=r['infos']))
            row_sources = [{c: (i.get('source') or {}).get('cell') and f"{i['source'].get('sheet')}!{i['source']['cell']}" for c, i in r['infos'].items() if (i.get('source') or {}).get('cell')} for r in rows]
            view['tables'][table] = boxes_for(table, ti, cols, meta, chosen, rows, issues, oi, labelled_blank if ti == 0 else {}, sec, book, list(sections.values()))
            view['tables'][table]['row_values'] = [r['values'] for r in rows]
            view['tables'][table]['row_sources'] = row_sources
            if ti == 0:
                mv = view['tables'][table]
                ref = text(main_values['IONO'][0]) if 'IONO' in main_values else ('' if 'IONO' in cols else text(o.get('order_ref')))
                view.update(ref=ref or None, buyer=text((main_values.get('BUYERNAME') or (o.get('buyer'),))[0]) or None,
                            style=text((main_values.get('STYLENO') or (o.get('style'),))[0]) or None,
                            description=text((main_values.get('STYLEDESC') or (o.get('description'),))[0]) or None)
                qty_col = next((n for n, b in chosen.items() if b and b['kind'] == 'row' and b.get('field') == 'quantity'), None)
                view['pieces'] = sum(r['values'].get(qty_col) or 0 for r in rows) if qty_col else None
                if 'IONO' in cols:
                    key = normalize(ref)
                    if not key:
                        warnings.append(f'{prefix}there is no OC number (IONO), so a second upload of the same order cannot be recognised. Add it if you have it.')
                        if committed_file:
                            errors.append(f'{prefix}this exact workbook was already saved (import {committed_file[:8]}) and has no OC number to tell the orders apart.')
                    elif key in seen_refs:
                        errors.append(f'{prefix}OC number {ref} appears twice in this workbook.')
                    elif key in existing_refs:
                        errors.append(f'{prefix}OC number {ref} is already in {table}. Open Orders to see or edit it, or use a new OC number.')
                        mv['boxes']['IONO'].update(state='error', note=f'{ref} is already saved in {table}')
                    seen_refs.add(key)
                elif committed_file:
                    warnings.append(f'{prefix}this exact workbook was saved before (import {committed_file[:8]}).')
                oq = main_values.get('OQTY')
                if oq is not None and view['pieces']:
                    try:
                        if Decimal(str(oq[0])) != Decimal(str(view['pieces'])):
                            warnings.append(f'{prefix}the size quantities add up to {view["pieces"]:,}, but the total pieces in Excel (OQTY) is {oq[0]:,}.')
                    except Exception:
                        pass
                for li, line in enumerate(o.get('lines', [])):
                    total, declared = line_total(line), line.get('declared_total')
                    if total is not None and isinstance(declared, (int, float)) and not isinstance(declared, bool) and declared != total:
                        warnings.append(f'{prefix}{line.get("color") or "line " + str(li + 1)}: the sizes add up to {total:,} but the Excel line total is {declared:,}.')
        views.append(view)
    known = set(GLOBAL_LABELS).union(*wanted_labels(metas[main]['columns'], (learned or {}).get(main)).values()) if main in metas else None
    pairs = excel_pairs(book, known)
    unused = [p for p in pairs if (p['sheet'], p['cell']) not in used_cells and not p.get('inline')][:40]
    mv = views[0]['tables'].get(main) if views else None
    ncols = len(metas[main]['columns']) if main in metas else 0
    filled = mv['filled'] if mv else 0
    if mv:
        info.append(f'{ncols - filled} of {ncols} {main} columns are not in this Excel and stay blank. Nothing is invented.')
    plan = dict(mode='entry', table=main, targets=[dict(table=t['table'], rows=tables.get(t['table'], {}).get('rows_kind', t.get('rows')),
                                                       section=t.get('section'), link=t.get('link') or {}) for t in targets],
                tables=tables, orders=views, operations=ops, errors=list(dict.fromkeys(errors)), warnings=list(dict.fromkeys(warnings)),
                info=list(dict.fromkeys(info)), issues=issues, unused=unused, labelled_blank=sorted(labelled_blank),
                totals=dict(orders=len(orders), rows=len(ops), main_rows=sum(1 for op in ops if op['table'] == main),
                            pieces=sum(v.get('pieces') or 0 for v in views), columns=ncols, filled=filled, blank=ncols - filled,
                            tables={t: sum(1 for op in ops if op['table'] == t) for t in tables}),
                schema_fingerprint=digest({t: {n: {k: v for k, v in m.items() if k != 'labels'} for n, m in metas[t]['columns'].items()} for t in tables}))
    # Columns Oracle requires that nothing fills are blocking (boxes_for marked them).
    plan['errors'] += [e for e in (f"{i['table']}.{i['column']}: Oracle needs a value. Map an Excel cell or type it."
                                   for i in issues if i['message'] == 'Oracle needs a value here') if e not in plan['errors']]
    plan['hash'] = digest(plan)
    return plan


def label_of(book, sections, src):
    """The words in Excel that name a source cell: its table column header, the label beside it, or the nearest text."""
    if not src or not src.get('cell') or not src.get('sheet'):
        return None
    c = book.at(src['sheet'], src['cell'])
    if c is None:
        return None
    for t in sections:
        if t['sheet'] == src['sheet'] and t['first_col'] <= c['col'] <= t['last_col'] and t['header_row'] <= c['row'] <= t['end_row']:
            head = book.cell(t['sheet'], t['header_row'], c['col'])
            if c['row'] == t['header_row']:
                return f"{t['title']}: column header"
            return f"{t['title']}: {text(head['value']) if head else 'column ' + str(c['col'])}"
    lab = label_for(book, src['sheet'], src['cell'])
    if lab:
        return lab
    near = anchor(book.by_name.get(src['sheet']), c)
    return near and text(book.at(src['sheet'], near['cell'])['value'])


def boxes_for(table, ti, cols, meta, chosen, rows, issues, oi, labelled_blank, sec, book=None, sections=()):
    """One box per column: what it will hold, where that came from, and its state."""
    boxes = {}
    for name, m in cols.items():
        b = chosen.get(name)
        present = [r['values'][name] for r in rows if r['values'].get(name) is not None]
        per_row = bool(b and b['kind'] in PER_ROW) or len(set(map(str, present))) > 1
        info = next((r['infos'][name] for r in rows if name in r['infos']), None) or {}
        kind = (b or {}).get('kind')
        box = dict(binding={k: v for k, v in b.items() if k != 'inline'} if b else None, level='row' if per_row else 'order',
                   count=len(present), values=list(dict.fromkeys(str(v) for v in present))[:8], distinct=len(set(map(str, present))),
                   value=present[0] if present and not per_row else None, source=info.get('source'), origin=info.get('origin'),
                   group=group_of(name), excel_label=label_of(book, sections, info.get('source')) if book is not None else None)
        if kind == 'row':
            where = (box['excel_label'] or 'Size table').split(':')[0]
            box['excel_label'] = {'quantity': f'{where}: quantity under each size', 'size': f'{where}: size headings', 'color': f'{where}: colour of each line',
                                  'label': f'{where}: garment of each line', 'line_no': 'Line number (1, 2, 3…)', 'size_no': 'Size position (1, 2, 3…)',
                                  'line_total': "Sum of each colour's sizes"}.get(b.get('field'), box['excel_label'])
        summable = kind == 'header' or (kind == 'row' and b.get('field') == 'quantity')
        if per_row and summable and present and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in present):
            total = sum(present); box['total'] = int(total) if float(total).is_integer() else round(total, 4)
        mine = [i for i in issues if i['table'] == table and i['column'] == name and i.get('order') == oi]
        if mine:
            box.update(state='error', note=mine[0]['message'])
        elif name == meta['keygen']:
            box.update(state='key', note='New number given when saved', value=None)
        elif present:
            box['state'] = 'link' if kind == 'link' else 'auto' if kind == 'auto_oc' else {'ai': 'ai', 'manual': 'typed', 'derived': 'derived'}.get(box['origin'], 'excel')
            if per_row and len(present) < len(rows):
                box['note'] = f'{len(rows) - len(present)} of {len(rows)} rows are blank in Excel'
            if (b or {}).get('score', 1) < 1:
                box['note'] = 'Similar name — please check'
        elif kind == 'na':
            box['state'] = 'na'
        elif kind == 'blank':
            box['state'] = 'cleared'
        elif m['identity'] or m['default']:
            box.update(state='default', note='Oracle fills this itself' + (f' ({m["default"]})' if m['default'] and len(m['default']) < 40 else ''))
        elif m['system']:
            box.update(state='system', note='Filled by the ERP')
        elif name in labelled_blank:
            w = labelled_blank[name]; box.update(state='empty_label', note=f'Excel has "{w["label"]}" ({w["cell"]}) but its value cell is empty')
        else:
            box['state'] = 'blank'
        if m['nullable'] == 'N' and not m['default'] and not m['identity'] and name != meta['keygen'] and not present:
            box.update(state='error', note='Oracle needs a value here')
            issues.append(dict(table=table, column=name, row=None, order=oi, level='error', message='Oracle needs a value here'))
        boxes[name] = box
    return dict(rows=len(rows), row_keys=[r['key'] for r in rows], row_labels=[r['label'] for r in rows], boxes=boxes,
                filled=sum(1 for x in boxes.values() if x['count']), section=sec and {k: sec.get(k) for k in ('id', 'title', 'sheet', 'header_row', 'end_row')})


def learned_by_table(rows):
    out = {}
    for r in rows:
        out.setdefault(r['table_name'], {}).setdefault(r['column_name'], []).append(r['label'])
    return out


def metas_for(conn, draft):
    out = {}
    for t in draft.get('targets') or default_targets():
        try:
            out[t['table']] = describe(conn, t['table'])
        except KeyError:
            pass
    return out


def plan(conn, draft, learned=None):
    metas = metas_for(conn, draft)
    c = conn.cursor()
    c.execute('select JOBID from PP_IMPORT_LOG where FILEHASH=:h fetch first 1 rows only', h=draft['filehash'])
    row = c.fetchone()
    main = (draft.get('targets') or default_targets())[0]['table']
    auto_oc = {}
    if OC_AUTO and main in metas and 'IONO' in metas[main]['columns']:
        first = build_plan(draft, metas, learned=learned)
        missing = [v['index'] for v in first['orders'] if not v.get('ref') or (((draft['orders'][v['index']].get('bind') or {}).get(main) or {}).get('IONO') or {}).get('kind') == 'auto_oc']
        if missing:
            numbers = next_oc_numbers(conn, [draft['orders'][i] for i in missing], main)
            auto_oc = {i: numbers[k] for k, i in enumerate(missing)}
    probe = build_plan(draft, metas, learned=learned, auto_oc=auto_oc)
    existing = set()
    if main in metas and 'IONO' in metas[main]['columns']:
        for v in probe['orders']:
            ref = normalize(v.get('ref'))
            if not ref:
                continue
            c.execute(f'SELECT 1 FROM {OWNER}."{main}" WHERE upper(trim(IONO)) = :r FETCH FIRST 1 ROWS ONLY', r=ref)
            found = c.fetchone()
            c.execute('SELECT 1 FROM PP_ORDER_KEYS WHERE ORDER_REF = :r', r=ref[:30])
            if found or c.fetchone():
                existing.add(ref)
    return build_plan(draft, metas, existing, row[0] if row else None, learned, auto_oc)


# ---------------------------------------------------------------- commit

def commit(draft, expected_hash, approver, learned=None):
    """All tables in one transaction: insert, read back by ROWID, receipt in PP_IMPORT_LOG, commit; then verify again."""
    with connect() as conn:
        c = conn.cursor(); c.execute('SELECT ID FROM PP_WRITE_LOCK WHERE ID=1 FOR UPDATE WAIT 10')
        previous = read_result(c, draft['id'])
        if previous:
            return previous
        checked = plan(conn, draft, learned)
        if checked['errors']:
            raise ValueError('; '.join(checked['errors']))
        if checked['hash'] != expected_hash:
            raise ValueError('The Oracle table or the preview changed. Open the preview again and review it before saving.')
        metas = metas_for(conn, draft)
        main = checked['table']; rows_out = []; orders_out = []; order_keys = {}
        try:
            for v in checked['orders']:
                c.execute('SELECT PP_RECORD_SEQ.nextval FROM dual'); order_keys[v['index']] = c.fetchone()[0]
                ref = normalize(v.get('ref'))
                if ref and 'IONO' in metas[main]['columns']:
                    c.execute('INSERT INTO PP_ORDER_KEYS(ORDER_REF,JOBID) VALUES (:ref,:job)', ref=ref[:30], job=draft['id'])
                orders_out.append(dict(order_ref=v.get('ref') or '(no OC number)', ordmasid=order_keys[v['index']], table=main, buyer=v.get('buyer'),
                                       style=v.get('style'), pieces=v.get('pieces'), rows=v['tables'].get(main, {}).get('rows', 0)))
            for op in checked['operations']:
                table = op['table']; cols = metas[table]['columns']; vals = {}
                for col, value in op['values'].items():
                    if value == '$key':
                        c.execute('SELECT PP_RECORD_SEQ.nextval FROM dual'); value = c.fetchone()[0]
                    vals[col] = bind_value(cols[col], value)
                names = list(vals); binds = {f'v{i}': vals[n] for i, n in enumerate(names)}
                # Table and column names come only from Oracle's own metadata (checked identifiers); values are bound.
                if not IDENT.match(table) or not all(IDENT.match(n) for n in names):
                    raise ValueError('Unexpected table or column name.')
                quoted = ','.join(f'"{n}"' for n in names)
                c.execute(f'INSERT INTO {OWNER}."{table}" ({quoted}) VALUES ({",".join(":" + k for k in binds)})', binds)
                rid = c.lastrowid
                c.execute(f'SELECT {quoted} FROM {OWNER}."{table}" WHERE ROWID = :r', r=rid)
                got = c.fetchone()
                if not got or not all(same(cols[n], vals[n], g) for n, g in zip(names, got)):
                    raise ValueError(f'{table}: a saved row did not read back the same. Nothing was saved.')
                c.execute('SELECT PP_RECORD_SEQ.nextval FROM dual'); handle = c.fetchone()[0]
                rows_out.append(dict(table=table, key=op['key'], row=op['row'], id=handle, rowid=rid, order=order_keys[op['order']],
                                     values={n: plain(g) for n, g in zip(names, got)}))
            result = dict(job_id=draft['id'], mode='entry', table=main, tables=list(checked['tables']), orders=orders_out, rows=rows_out,
                          committed_at=datetime.now().astimezone().isoformat(), already_committed=False)
            c.setinputsizes(payload=oracledb.DB_TYPE_CLOB)
            c.execute('INSERT INTO PP_IMPORT_LOG(JOBID,FILEHASH,PLANHASH,APPROVER,RESULT_JSON) VALUES (:job,:filehash,:hash,:who,:payload)',
                      job=draft['id'], filehash=draft['filehash'], hash=expected_hash, who=approver, payload=json.dumps(result, ensure_ascii=False, default=str))
            conn.commit()
        except Exception:
            conn.rollback(); raise
    result['verified_rows'] = verify(result)
    result['verified'] = result['verified_rows'] == len(result['rows'])
    return result


def verify(receipt):
    """Fresh connection: every saved row read back by ROWID with the same values."""
    good = 0
    with connect() as conn:
        c = conn.cursor()
        for row in receipt['rows']:
            names = list(row['values']); quoted = ','.join(f'"{n}"' for n in names)
            c.execute(f'SELECT {quoted} FROM {OWNER}."{row["table"]}" WHERE ROWID = :r', r=row['rowid'])
            got = c.fetchone()
            if got and all(plain(g) == row['values'][n] for n, g in zip(names, got)):
                good += 1
    return good


# ---------------------------------------------------------------- saved orders

def fetch_rows(conn, table, rowids):
    """Current rows by ROWID: {rowid: {column: value}}; rows that no longer exist are absent."""
    out = {}
    if not IDENT.match(table):
        return out
    c = conn.cursor()
    for n in range(0, len(rowids), 500):
        part = rowids[n:n + 500]; binds = {f'r{i}': r for i, r in enumerate(part)}
        c.execute(f'SELECT ROWIDTOCHAR(t.ROWID) AS PP_ROWID, t.* FROM {OWNER}."{table}" t WHERE t.ROWID IN ({",".join(":" + k for k in binds)})', binds)
        names = [d[0] for d in c.description]
        for r in c:
            rec = dict(zip(names, map(plain, r)))
            out[rec.pop('PP_ROWID')] = rec
    return out


def order_rows(conn, receipt, order_key):
    """{table: [row dicts with _id (POPPYS row number) and _rowid]} for one saved order, in saved order."""
    mine = [r for r in receipt['rows'] if r.get('order') == order_key]
    if not mine:
        raise KeyError('Order not found')
    out = {}
    for table in dict.fromkeys(r['table'] for r in mine):
        rows = [r for r in mine if r['table'] == table]
        current = fetch_rows(conn, table, [r['rowid'] for r in rows])
        out[table] = [dict(current[r['rowid']], _id=r['id'], _rowid=r['rowid'], _row=r.get('row')) if r['rowid'] in current
                      else dict(_id=r['id'], _rowid=r['rowid'], _row=r.get('row'), _missing=True) for r in rows]
    return out


def summarize(rows_by_table, main):
    rows = [r for r in rows_by_table.get(main, []) if not r.get('_missing')]
    head = rows[0] if rows else {}
    sizes, lines = {}, {}
    for r in rows:
        if r.get('SIZES') is not None:
            sizes[r['SIZES']] = sizes.get(r['SIZES'], 0) + (r.get('ORDERQTY') or 0)
            line = lines.setdefault((r.get('ORDSTYLEROW'), r.get('COMBO')), dict(label=r.get('COMBO') or f'Line {r.get("ORDSTYLEROW")}', quantity=0, sizes=[]))
            line['quantity'] += r.get('ORDERQTY') or 0; line['sizes'].append(dict(size=r['SIZES'], quantity=r.get('ORDERQTY')))
    size_total = sum(sizes.values())
    return dict(order_ref=head.get('IONO'), buyer=head.get('BUYERNAME'), style=head.get('STYLENO'), description=head.get('STYLEDESC'),
                po=head.get('PONUM'), order_date=head.get('ORDERDATE'), ship_date=head.get('SHIPDT'),
                total=head.get('OQTY') if head.get('OQTY') is not None else size_total, size_total=size_total,
                reconciled=head.get('OQTY') is None or head.get('OQTY') == size_total,
                sizes=[dict(size=k, quantity=v) for k, v in sizes.items()], lines=list(lines.values()),
                records={t: len(v) for t, v in rows_by_table.items()}, missing=sum(1 for v in rows_by_table.values() for r in v if r.get('_missing')))


def preview_changes(conn, receipt, order_key, changes):
    """Edits of a saved order checked against the current rows. Nothing is written here.
    A change with id 'all' sets the column on every row of that table for this order."""
    rows_by_table = order_rows(conn, receipt, order_key)
    metas = {t: describe(conn, t) for t in rows_by_table}
    errors, notes, out, snapshot = [], [], [], {}
    proposed = {t: {r['_id']: dict(r) for r in rows if not r.get('_missing')} for t, rows in rows_by_table.items()}
    typed = {str(ch.get('column', '')).upper() for ch in changes}
    if not changes:
        errors.append('No changes to review.')
    for ch in changes[:5000]:
        table = str(ch.get('table') or receipt.get('table') or MAIN).upper(); col = str(ch.get('column') or '').upper(); value = ch.get('value')
        if table not in proposed:
            errors.append(f'{table}: not part of this order.'); continue
        meta = metas[table]; cols = meta['columns']
        if col not in cols or col == meta['keygen'] or cols[col]['identity']:
            errors.append(f'{col}: not editable.'); continue
        if not meta['can_update']:
            errors.append(f'POPPYS is not allowed to change {table} yet. On this computer run: python scripts/upgrade_local.py'); continue
        ids = list(proposed[table]) if ch.get('id') == 'all' else [ch.get('id')]
        if any(i not in proposed[table] for i in ids):
            errors.append(f'Row {ch.get("id")} does not belong to this order.'); continue
        if blank(value):
            new = None
        else:
            try:
                new = convert(cols[col], value)
                if isinstance(new, str) and cols[col]['type'] == 'NUMBER':
                    new = plain(Decimal(new))
            except ValueError as exc:
                errors.append(f'{cols[col]["label"]}: {friendly_error(cols[col], exc)}.'); continue
        if cols[col]['nullable'] == 'N' and new is None:
            errors.append(f'{cols[col]["label"]}: Oracle needs a value here.'); continue
        for i in ids:
            proposed[table][i][col] = new
    main = receipt.get('table') or MAIN
    if main in proposed and 'ORDERQTY' in metas[main]['columns']:
        before_rows = {r['_id']: r for r in rows_by_table[main] if not r.get('_missing')}
        before = sum(r.get('ORDERQTY') or 0 for r in before_rows.values()); after = sum(p.get('ORDERQTY') or 0 for p in proposed[main].values())
        if before != after and 'OQTY' not in typed and 'OQTY' in metas[main]['columns'] and any(r.get('OQTY') is not None for r in before_rows.values()):
            if all(r.get('OQTY') == before for r in before_rows.values()):
                for p in proposed[main].values():
                    p['OQTY'] = after
                notes.append(f'Total pieces (OQTY) follows the sizes: {before:,} → {after:,}.')
            else:
                notes.append('OQTY was not the sum of the sizes before, so it is left as it is.')
        if 'COMBOQTY' in metas[main]['columns'] and 'COMBOQTY' not in typed:
            groups = {}
            for rid, r in before_rows.items():
                groups.setdefault((r.get('ORDSTYLEROW'), r.get('COMBO')), []).append(rid)
            for (line, combo), ids in groups.items():
                b = sum(before_rows[i].get('ORDERQTY') or 0 for i in ids); a = sum(proposed[main][i].get('ORDERQTY') or 0 for i in ids)
                if a != b and all(before_rows[i].get('COMBOQTY') == b for i in ids):
                    for i in ids:
                        proposed[main][i]['COMBOQTY'] = a
                    notes.append(f'Colour quantity (COMBOQTY) of {combo or "line " + str(line)} follows its sizes: {b:,} → {a:,}.')
    for table, rows in rows_by_table.items():
        cols = metas[table]['columns']
        for r in rows:
            if r.get('_missing'):
                continue
            p = proposed[table][r['_id']]
            for col in cols:
                old, new = r.get(col), p.get(col)
                if old != new:
                    snapshot.setdefault(f'{table}:{r["_id"]}', {})[col] = old
                    out.append(dict(table=table, id=r['_id'], rowid=r['_rowid'], column=col, old=old, new=new, label=cols[col]['label'],
                                    derived=col in ('OQTY', 'COMBOQTY') and col not in typed))
    if not out and not errors:
        errors.append('The values are the same as in Oracle. Nothing to change.')
    result = dict(mode='entry', ordmasid=order_key, job_id=receipt['job_id'], changes=out, errors=list(dict.fromkeys(errors)), notes=notes, snapshot=snapshot)
    result['hash'] = digest(result)
    return result


def apply_changes(changeset_id, preview, approver):
    """One transaction: lock each row by ROWID, confirm nobody changed it since the preview, update, read back, audit."""
    if preview['errors']:
        raise ValueError('Resolve the highlighted issues first.')
    with connect() as conn:
        c = conn.cursor(); c.execute('SELECT ID FROM PP_WRITE_LOCK WHERE ID=1 FOR UPDATE WAIT 10')
        c.execute('SELECT count(*) FROM PP_CHANGE_LOG WHERE CHANGESET=:id', id=changeset_id)
        if c.fetchone()[0]:
            return dict(already_applied=True)
        grouped = {}
        for ch in preview['changes']:
            grouped.setdefault((ch['table'], ch['id'], ch['rowid']), []).append(ch)
        metas = {t: columns(conn, t) for t in {k[0] for k in grouped}}
        try:
            for (table, rid, rowid), items in grouped.items():
                if not IDENT.match(table) or not all(IDENT.match(i['column']) for i in items):
                    raise ValueError('Unexpected table or column name.')
                cols = [i['column'] for i in items]; quoted = ','.join(f'"{x}"' for x in cols)
                c.execute(f'SELECT {quoted} FROM {OWNER}."{table}" WHERE ROWID = :r FOR UPDATE WAIT 10', r=rowid)
                got = c.fetchone()
                if not got or any(plain(g) != i['old'] for g, i in zip(got, items)):
                    raise ValueError('This order was changed in Oracle after your preview. Reload the order and review again.')
                binds = {f'v{n}': bind_value(metas[table][i['column']], i['new']) for n, i in enumerate(items)}
                assignments = ','.join('"%s"=:v%d' % (i['column'], n) for n, i in enumerate(items))
                c.execute(f'UPDATE {OWNER}."{table}" SET {assignments} WHERE ROWID = :rowid_', dict(binds, rowid_=rowid))
                c.execute(f'SELECT {quoted} FROM {OWNER}."{table}" WHERE ROWID = :r', r=rowid)
                back = c.fetchone()
                if any(plain(g) != i['new'] for g, i in zip(back, items)):
                    raise ValueError(f'{table}: updated values did not read back the same. Nothing was changed.')
                for i in items:
                    c.execute('''INSERT INTO PP_CHANGE_LOG(CHANGESET,ORDMASID,TABLE_NAME,ROW_ID,COLUMN_NAME,OLD_VALUE,NEW_VALUE,APPROVER)
                        VALUES (:cs,:o,:t,:r,:c,:old,:new,:who)''', cs=changeset_id, o=preview['ordmasid'], t=table, r=rid, c=i['column'],
                              old=None if i['old'] is None else str(i['old'])[:4000], new=None if i['new'] is None else str(i['new'])[:4000], who=approver)
            conn.commit()
        except Exception:
            conn.rollback(); raise
    return dict(already_applied=False, applied=len(preview['changes']), verified=True)


def label_for(book, sheet, address):
    """The Excel label whose value is this cell (the inverse of label matching), or None. Used to learn from the
    person's own mapping: next time a workbook has the same label, the column is filled by itself."""
    target = book.at(sheet, address)
    if target is None:
        return None
    cells = book.cells.get(sheet, {})
    r, col = target['row'], target['col']
    candidates = [cells.get((r, x)) for x in range(col - 1, max(col - 4, 0), -1)] + [cells.get((r - 1, col))]
    for cand in candidates:
        if cand is None or not isinstance(cand['value'], str) or not cand['value'].strip() or len(cand['value']) > 40:
            continue
        known = set(GLOBAL_LABELS) | {norm(cand['value'])}
        v = value_beside(cells, cand, known)
        if v is not None and v['cell'] == address and any(ch.isalpha() for ch in cand['value']) and sum(ch.isdigit() for ch in cand['value']) <= 2:
            return cand['value'].strip()
    return None


def make_section(sheets, sheet, header_row, title=None):
    """An Excel table chosen by its header row (for tables without a title row)."""
    book = Book(sheets)
    if sheet not in book.cells:
        raise ValueError('Choose a worksheet.')
    heads = sorted((c for (r, _), c in book.cells[sheet].items() if r == header_row and isinstance(c['value'], str) and c['value'].strip()), key=lambda x: x['col'])
    if len(heads) < 2:
        raise ValueError(f'Row {header_row} does not look like a table header (it needs at least two column names).')
    rows = sorted({r for (r, _) in book.cells[sheet]})
    end, gap = header_row, 0
    for r in range(header_row + 1, (rows[-1] if rows else header_row) + 1):
        if any((r, c) in book.cells[sheet] for c in range(heads[0]['col'], heads[-1]['col'] + 1)):
            end, gap = r, 0
        else:
            gap += 1
            if gap >= 2:
                break
    sec = dict(id=f'{sheet}!R{header_row}', sheet=sheet, title=title or f'Rows below row {header_row}', title_cell=None, title_row=None,
               header_row=header_row, first_col=heads[0]['col'], last_col=heads[-1]['col'], end_row=end,
               headers=[dict(col=h['col'], cell=h['cell'], text=text(h['value'])) for h in heads])
    sec['rows'] = len(section_rows(book, sec))
    return sec


OC_FORMAT = setting('OC_FORMAT', 'PK/{fy}/{type}/{n}')
OC_AUTO = setting('OC_AUTO', '1') not in ('0', 'no', 'false', 'off')


def financial_year(day=None):
    d = day or datetime.now().date()
    start = d.year if d.month >= 4 else d.year - 1
    return f'{start % 100:02d}-{(start + 1) % 100:02d}'


def oc_series(order):
    """Prefix of the OC number series for an order, e.g. PK/26-27/E/ (E for bulk orders, S for samples)."""
    kind = setting('OC_TYPE_SAMPLE', 'S') if order.get('quantity_basis') == 'sample' else setting('OC_TYPE_ORDER', 'E')
    return OC_FORMAT.replace('{fy}', financial_year()).replace('{type}', kind).split('{n}')[0]


def next_oc_numbers(conn, orders, main):
    """{order index: next free OC number} for orders whose Excel has no OC number, continuing the client's own series
    (the highest number with the same prefix in the order table, in ORDMAS.DOCID and in POPPYS's own register).
    Read again inside the save transaction, under the POPPYS write lock, so two saves never get the same number."""
    out, last = {}, {}
    c = conn.cursor()
    for i, o in enumerate(orders):
        prefix = oc_series(o)
        if prefix not in last:
            found = []
            for sql in (f'SELECT "IONO" FROM {OWNER}."{main}" WHERE "IONO" LIKE :p', f'SELECT DOCID FROM {OWNER}.ORDMAS WHERE DOCID LIKE :p',
                        'SELECT ORDER_REF FROM PP_ORDER_KEYS WHERE ORDER_REF LIKE :p'):
                try:
                    c.execute(sql + ' FETCH FIRST 100000 ROWS ONLY', p=prefix.upper() + '%')
                    found += [r[0] for r in c]
                except oracledb.Error:
                    continue
            nums = [str(v)[len(prefix):].strip() for v in found if v and str(v).upper().startswith(prefix.upper())]
            nums = [n for n in nums if n.isdigit()]
            width = max((len(n) for n in nums), default=0)
            last[prefix] = (max((int(n) for n in nums), default=int(setting('OC_START', '1')) - 1), width)
        n, width = last[prefix]
        last[prefix] = (n + 1, width)
        out[i] = f'{prefix}{str(n + 1).zfill(width)}'
    return out


def writable_tables(conn):
    c = conn.cursor()
    c.execute('''select distinct table_name from all_tab_privs where table_schema = :o and privilege = 'INSERT'
        and (grantee in (user, 'PUBLIC') or grantee in (select role from session_roles))''', o=OWNER)
    return sorted(t for (t,) in c if IDENT.match(t) and not t.startswith(('PP_', 'TEST_')))


def suggest_links(conn, draft, cache=None):
    """Which other Oracle tables this upload's Excel tables belong to. A table is suggested for an Excel table when
    POPPYS may add rows to it, it has the order's key column (IONO) to link the rows, and at least three of the
    Excel headers name its columns (and at least half of them). Size blocks are the main table's own rows."""
    main = (draft.get('targets') or default_targets())[0]['table']
    used = {t['table'] for t in draft.get('targets') or []}
    out = []
    for table in writable_tables(conn):
        if table in used or table == main:
            continue
        cols = (cache or {}).get(table) or columns(conn, table)
        if cache is not None:
            cache[table] = cols
        if 'IONO' not in cols:
            continue
        names = {n: [n.replace('_', ' '), m['label']] + m['labels'] for n, m in cols.items() if not m['system'] and not m['identity'] and n != 'IONO'}
        for sec in draft.get('sections') or []:
            if sec['title'].startswith('Size block') or not sec.get('rows'):
                continue
            hits = {}
            for h in sec['headers']:
                best = max(((header_score(h['text'], v), n) for n, v in names.items()), default=(0, None))
                if best[0] >= 0.8 and best[1] not in hits.values():
                    hits[h['text']] = best[1]
            if len(hits) >= 3 and len(hits) >= len(sec['headers']) / 2:
                out.append(dict(section=sec['id'], title=sec['title'], table=table, matched=len(hits), headers=len(sec['headers']),
                                columns=hits, link={'IONO': 'IONO'}, rows=sec['rows']))
    out.sort(key=lambda x: (-x['matched'] / x['headers'], -x['matched']))
    best, taken = [], set()
    for s in out:
        if s['section'] in {b['section'] for b in best} or s['table'] in taken:
            continue
        best.append(s); taken.add(s['table'])
    return best


def current_summaries(conn, receipts):
    """Today's values of POPPYS-saved orders, read by ROWID: {order key: summary}."""
    wanted = {}
    for rec in receipts:
        for r in rec.get('rows', []):
            if r.get('table') == rec.get('table'):
                wanted.setdefault(r['table'], []).append(r['rowid'])
    current = {t: fetch_rows(conn, t, ids) for t, ids in wanted.items()}
    out = {}
    for rec in receipts:
        main = rec.get('table')
        for o in rec.get('orders', []):
            rows = [current.get(main, {}).get(r['rowid']) for r in rec['rows'] if r.get('order') == o['ordmasid'] and r['table'] == main]
            have = [x for x in rows if x is not None]
            head = have[0] if have else {}
            qty = sum(x.get('ORDERQTY') or 0 for x in have) if have and 'ORDERQTY' in head else head.get('OQTY')
            out[o['ordmasid']] = dict(order_ref=head.get('IONO') or o.get('order_ref'), buyer=head.get('BUYERNAME', o.get('buyer')),
                                      style=head.get('STYLENO', o.get('style')), quantity=qty, order_date=head.get('ORDERDATE'), ship_date=head.get('SHIPDT'),
                                      rows=len(rows), missing=len(rows) - len(have))
    return out


# ---------------------------------------------------------------- templates (a saved mapping for one Excel layout)

def template_payload(draft):
    """Targets plus the person's own bindings; each cell binding keeps the label beside it to recognise the layout later."""
    sheets = {s['name']: s for s in draft['sheets']}
    bindings = []
    for o in (draft.get('orders') or [])[:1]:
        for table, cols in (o.get('bind') or {}).items():
            for col, b in cols.items():
                if b.get('ai') and not b.get('confirmed'):
                    continue
                item = dict(table=table, column=col, binding={k: v for k, v in b.items() if k not in ('auto',)})
                if b.get('kind') == 'cell':
                    sheet = sheets.get(b.get('sheet'))
                    cell = next((c for c in sheet['cells'] if c['cell'] == b.get('cell')), None) if sheet else None
                    if cell is None:
                        continue
                    item['anchor'] = anchor(sheet, cell)
                if b.get('kind') == 'text':
                    continue   # a typed value belongs to one order, not to the layout
                bindings.append(item)
    targets = []
    for t in draft.get('targets') or default_targets():
        sec = t.get('section')
        targets.append(dict(table=t['table'], rows=t.get('rows'), link=t.get('link') or {},
                            section=dict(title=sec.get('title'), headers=[h['text'] if isinstance(h, dict) else h for h in sec.get('headers', [])]) if sec else None))
    return dict(kind='entry', targets=targets, bindings=bindings)


def apply_template(draft, template):
    """Reuse a saved mapping where the workbook still matches. A cell binding counts only when the label beside the
    cell reads the same; a linked Excel table is found again by its title and headers. Returns (share applied, missing)."""
    sheets = {s['name']: s for s in draft['sheets']}
    if not draft.get('orders'):
        return 0, []
    o = draft['orders'][0]
    found = draft.get('sections') or find_sections(draft['sheets'], draft['orders'])
    targets, missing, applied = [], [], 0
    for t in template.get('targets') or []:
        sec = t.get('section')
        if sec:
            want = [norm(h) for h in sec.get('headers') or []]
            match = next((s for s in found if norm(s['title']) == norm(sec.get('title')) and [norm(h['text']) for h in s['headers']][:len(want)] == want), None)
            if match is None:
                missing.append(dict(table=t['table'], what=f'Excel table "{sec.get("title")}"'))
                continue
            t = dict(t, section=dict(id=match['id']))
        targets.append(t)
    if targets:
        draft['targets'] = targets
    for item in template.get('bindings') or []:
        b = item['binding']
        if b.get('kind') == 'cell':
            sheet = sheets.get(b.get('sheet'))
            cell = next((c for c in sheet['cells'] if c['cell'] == b.get('cell')), None) if sheet else None
            if cell is None or (item.get('anchor') is not None and anchor(sheet, cell) != item['anchor']):
                missing.append(dict(table=item['table'], column=item['column'], what=f"{b.get('sheet')}!{b.get('cell')}"))
                continue
        if not any(x['table'] == item['table'] for x in draft.get('targets') or default_targets()):
            continue
        o.setdefault('bind', {}).setdefault(item['table'], {})[item['column']] = dict(b, via=f'Saved mapping "{template.get("name")}"')
        applied += 1
    return applied / (len(template.get('bindings') or []) or 1), missing
