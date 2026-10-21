"""Read-only workbook ingestion. Workbook text is data, never executable instructions."""
from __future__ import annotations
from copy import deepcopy
from datetime import date, datetime, time, timedelta
from io import BytesIO
import re
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook

ALIASES = {
 'buyer': ['buyer','buyer name','customer','customer name'],
 'style': ['style','style no','style number','item','item number','style #'],
 'description': ['description','garment description','item description'],
 'po_number': ['po','po number','po no','contract','contract no'],
 'order_ref': ['oc no','order reference','order ref','docid'],
 'ship_date': ['ship date','ship dt','delivery date','ex-factory date','fob date'],
 'size': ['size','sizes'], 'quantity': ['quantity','qty','sample qty','order qty'],
 'color': ['colour','color','combo','colour description'],
}

def norm(value):
    return re.sub(r'[^a-z0-9]+',' ',str(value or '').lower()).strip()

def scalar(cell):
    value = cell.value
    if isinstance(value, (datetime,date)):
        return value.isoformat()[:10]
    if isinstance(value, (time,timedelta)):
        return str(value)
    if isinstance(value, (int,float)) and re.fullmatch(r'0{2,}',cell.number_format):
        return str(int(value)).zfill(len(cell.number_format))
    return value.strip() if isinstance(value,str) else value

def text(value):
    return '' if value is None else str(value).strip()

def source(sheet, address, value, formula=None):
    return {'sheet':sheet,'cell':address,'value':value,'formula':formula}

def base_order():
    return dict(buyer='', style='', description='', po_number='', order_ref='', ship_date='',
        order_date='', season='', merch='', sample_type='',
        pieces_per_unit=None, quantity_basis='unconfirmed', lines=[], provenance={},
        warnings=[], order_lines=[], sample_lines=[])

def inspect(content):
    if len(content)>12*1024*1024:
        raise ValueError('Workbook exceeds 12 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist())>10000 or sum(z.file_size for z in archive.infolist())>80*1024*1024:
                raise ValueError('Workbook expanded size exceeds 80 MB.')
            if '[Content_Types].xml' not in archive.namelist():
                raise ValueError('Please upload an Excel .xlsx or .xlsm workbook.')
    except BadZipFile as exc:
        raise ValueError('Please upload an unencrypted .xlsx or .xlsm workbook.') from exc
    raw=load_workbook(BytesIO(content),data_only=False,keep_links=False)
    cached=load_workbook(BytesIO(content),data_only=True,keep_links=False)
    sheets=[]
    if len(raw.worksheets)>30:
        raise ValueError('Workbook exceeds 30 sheets.')
    for ws in raw:
        if ws.max_row*ws.max_column>250000:
            raise ValueError(f'{ws.title}: sheet exceeds 250,000 cells; trim unused formatting.')
        cells=[]
        for row in ws:
            for cell in row:
                if cell.value is not None:
                    formula=cell.value if cell.data_type=='f' else None
                    if formula is not None and not isinstance(formula,str):
                        # Array / data-table formulas: keep their formula text, never the object.
                        formula=str(getattr(formula,'text',None) or formula)
                    value=scalar(cached[ws.title][cell.coordinate]) if formula else scalar(cell)
                    cells.append(dict(cell=cell.coordinate,row=cell.row,col=cell.column,value=value,
                        formula=formula,format=cell.number_format if cell.number_format!='General' else None,hidden=bool(ws.row_dimensions[cell.row].hidden or ws.column_dimensions[cell.column_letter].hidden),
                        error=cell.data_type=='e'))
        sheets.append(dict(name=ws.title,hidden=ws.sheet_state!='visible',rows=ws.max_row,cols=ws.max_column,
                           merged=[str(r) for r in ws.merged_cells.ranges],cells=cells))
    raw.close(); cached.close()
    return sheets

def parse(content):
    sheets=inspect(content)
    orders=[]; formats=[]
    for sheet in sheets:
        if sheet['hidden']:
            continue
        cells={(c['row'],c['col']):c for c in sheet['cells']}
        def v(r,c): return cells.get((r,c),{}).get('value')
        def p(r,c):
            cell=cells.get((r,c),{})
            return source(sheet['name'],cell.get('cell',''),cell.get('value'),cell.get('formula'))
        def setval(o,k,r,c):
            o[k]=text(v(r,c)); o['provenance'][k]=p(r,c)
        # Structural signatures, never filenames or worksheet title (D70140 is not the style).
        if norm(v(2,2))=='buyer' and norm(v(4,2))=='item' and norm(v(11,2))=='garment':
            o=base_order()
            for key,r,c in [('buyer',2,3),('style',4,3),('description',5,3),('po_number',5,5),
                            ('ship_date',3,5),('season',3,3),('merch',6,3)]: setval(o,key,r,c)
            headers=[r for r in range(10,min(sheet['rows']+1,80)) if norm(v(r,2))=='garment' and norm(v(r,4))=='contract']
            blocks={}
            for header in headers:
                end=next((r for r in range(header+1,min(header+101,sheet['rows']+1)) if norm(v(r,2))=='total'),None)
                total_col=next((c for c in range(5,sheet['cols']+1) if norm(v(header,c))=='total'),None)
                if end is None or total_col is None:raise ValueError(f'{sheet["name"]}: incomplete garment/size block near row {header}.')
                role=norm(v(header+1,4))
                role=role if role in ('sample','summary') else 'order'
                if role in blocks:raise ValueError(f'{sheet["name"]}: multiple {role} blocks need explicit mapping; they were not combined.')
                blocks[role]=(header,end,total_col)
            def lines(role):
                result=[]
                if role not in blocks:return result
                header,end,total_col=blocks[role]
                for r in range(header+1,end):
                    if v(r,2) is None and v(r,3) is None:continue
                    result.append(dict(label=text(v(r,2)),color=text(v(r,3)),
                        declared_total=v(r,total_col),provenance={'color':p(r,3)},
                        sizes=[dict(size=text(v(header,c)),quantity=v(r,c),source=p(r,c),size_source=p(header,c)) for c in range(5,total_col)]))
                return result
            o['order_lines']=lines('order'); o['sample_lines']=lines('sample')
            # OCN (order confirmation number) sits under its label in the quantity header row.
            if norm(v(8,2))=='ocn' and v(9,2) is not None:setval(o,'order_ref',9,2)
            def total(ls):return sum(s['quantity'] for l in ls for s in l['sizes'] if isinstance(s['quantity'],(int,float)) and not isinstance(s['quantity'],bool))
            # An all-zero sample section holds no sample quantities; the order section is then the only real source.
            use_sample=bool(o['sample_lines']) and (total(o['sample_lines'])>0 or not o['order_lines'])
            o['quantity_basis']='sample' if use_sample else 'order'
            o['lines']=deepcopy(o['sample_lines'] if use_sample else o['order_lines']); o['pieces_per_unit']=1
            o['warnings']=[('Sample section selected. Switch to order quantities if this is a bulk order.' if use_sample else
                    'Order quantities selected (the sample section is empty). Switch to the sample section if needed.')+' Summary repeats are excluded.',
                'Fabric and accessory schedules remain reference-only in this import.',
                'BOM date is not an order confirmation date. Enter the order date if you have it.']
            if norm(v(8,4))=='qty in pieces' and isinstance(v(9,4),(int,float)) and o['order_lines'] and total(o['order_lines'])!=v(9,4):
                o['warnings'].append(f'The garment lines add up to {total(o["order_lines"])} pieces, but QTY(IN PIECES) in {p(9,4)["cell"]} says {v(9,4)}. Check the size quantities.')
            orders.append(o); formats.append('Merch BOM')
        elif any(norm(c['value'])=='contract confirmation report' for c in sheet['cells']):
            for cell in sheet['cells']:
                if norm(cell['value'])!='item number': continue
                r,c=cell['row'],cell['col']; o=base_order()
                setval(o,'style',r,c+1);setval(o,'description',r+1,c+1)
                header=next((x for x in range(r+2,min(r+15,sheet['rows']+1)) if norm(v(x,6))=='contract no'),None)
                if not header: continue
                data=header+1; setval(o,'po_number',data,6); setval(o,'ship_date',data,8)
                delivery=text(v(data,7))
                if delivery:
                    o['po_number']+='-'+delivery
                    o['provenance']['po_number']['derivation']='Contract number + delivery suffix'
                    o['provenance']['po_number']['delivery']=p(data,7)
                size_columns=[col for col in range(11,sheet['cols']+1) if v(header,col) is not None]
                if not size_columns:raise ValueError(f'{sheet["name"]}: no size headings found for {o["style"]}.')
                if any(norm(v(rr,10))=='quantity' for rr in range(data+1,min(data+5,sheet['rows']+1))):
                    raise ValueError(f'{sheet["name"]}: multiple deliveries for {o["style"]} require explicit separate orders.')
                line=dict(label=o['style'],color='',declared_total=v(data,9),provenance={},
                    sizes=[dict(size=text(v(header,col)),quantity=v(data,col),source=p(data,col),size_source=p(header,col)) for col in size_columns])
                o['order_lines']=[line];o['lines']=deepcopy(o['order_lines']);o['quantity_basis']='order'
                o['warnings']=['Quantities are the buyer contract (order) quantities, saved exactly as in Excel unless you enter pieces per unit.',
                    'Buyer, colour and order date are not in this layout; they stay blank unless you add them. Destination is not assumed to be buyer.']
                orders.append(o)
            formats.append('Buyer contract')
        else:
            mapping=None
            for row in range(1,min(sheet['rows'],80)+1):
                found={key:cell['col'] for cell in sheet['cells'] if cell['row']==row
                       for key,aliases in ALIASES.items() if norm(cell['value']) in [norm(a) for a in aliases]}
                if {'style','size','quantity'} <= found.keys():
                    mapping={'sheet':sheet['name'],'header_row':row,'columns':found};break
            if mapping:
                orders.extend(parse_mapping(sheets,mapping));formats.append('Tabular orders')
    if not orders:
        from .mapping import detect_generic
        found=detect_generic(sheets)
        if found:
            orders.append(found);formats.append('Detected labels')
    return dict(sheets=sheets,orders=orders,format=', '.join(dict.fromkeys(formats)) or 'Needs column mapping')

def parse_mapping(sheets,mapping):
    sheet=next((s for s in sheets if s['name']==mapping['sheet']),None)
    if not sheet: raise ValueError('Select a worksheet.')
    row=int(mapping['header_row']); columns=mapping['columns']
    if not {'style','size','quantity'} <= columns.keys():
        raise ValueError('Map style, size and quantity columns at minimum.')
    if any(k not in ALIASES or not isinstance(c,int) or c<1 or c>sheet['cols'] for k,c in columns.items()):
        raise ValueError('Invalid column mapping.')
    cells={(c['row'],c['col']):c for c in sheet['cells']}; groups={}
    for r in range(row+1,sheet['rows']+1):
        fields={k:cells.get((r,c),{}).get('value') for k,c in columns.items()}
        if not any(x is not None for x in fields.values()): continue
        if norm(fields.get('style')) in ('total','grand total'): continue
        # Blank styles cannot safely be forward-filled across unrelated sections.
        if not fields.get('style'): raise ValueError(f"{sheet['name']} row {r}: style is missing. Fill it or choose the correct header.")
        key=tuple(text(fields.get(k)) for k in ('buyer','style','po_number','order_ref','ship_date'))
        o=groups.setdefault(key,base_order())
        for k,value in fields.items():
            if k not in ('size','quantity','color'):
                if o.get(k) and o[k]!=text(value): raise ValueError(f'Conflicting {k} in row {r}.')
                o[k]=text(value)
                c=cells.get((r,columns[k]),{})
                o['provenance'][k]=source(sheet['name'],c.get('cell',''),value,c.get('formula'))
        color=text(fields.get('color'))
        line=next((x for x in o['lines'] if x['color']==color),None)
        if line is None:
            line=dict(label=o['style'],color=color,declared_total=None,sizes=[],provenance={});o['lines'].append(line)
        qc=cells.get((r,columns['quantity']),{})
        line['sizes'].append(dict(size=text(fields.get('size')),quantity=fields.get('quantity'),
                                 source=source(sheet['name'],qc.get('cell',''),fields.get('quantity'),qc.get('formula'))))
    for o in groups.values():
        o['warnings']=['Quantities are saved exactly as in Excel unless you enter pieces per unit.']
        o['order_lines']=deepcopy(o['lines']);o['quantity_basis']='order'
    return list(groups.values())
